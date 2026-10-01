"""Automated QA for the Growth OS workbooks.
Recalculates copies in LibreOffice and checks values against an independent Python model.
Usage: python qa.py <out_dir> <work_dir>
"""
import sys, os, json, shutil, subprocess, datetime as dt, re
from collections import defaultdict
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
import demo_data
import build_workbook as B

RECALC = "/root/.claude/skills/synced/537785f1-f0de-401f-bf47-0f2c32573f85_73cbe849-f847-4991-82d7-5f1ecd5f11a9/xlsx/scripts/recalc.py"
OUT, WORK = sys.argv[1], sys.argv[2]
os.makedirs(WORK, exist_ok=True)
B.build(os.path.join(WORK, "_rowmap.xlsx"), False)  # populates MONTHLY_ROWS / SETROW
DEMO = os.path.join(OUT, "CLEANING-BUSINESS-AI-GROWTH-OS-DEMO.xlsx")
CLEAN = os.path.join(OUT, "CLEANING-BUSINESS-AI-GROWTH-OS-CLEAN.xlsx")
results = []


def check(name, ok, detail=""):
    results.append({"test": name, "pass": bool(ok), "detail": str(detail)[:400]})
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {detail}"))


def recalc(path):
    r = subprocess.run([sys.executable, RECALC, path, "300"], capture_output=True, text=True)
    return json.loads(r.stdout)


def copy_calc(src, name):
    dst = os.path.join(WORK, name)
    shutil.copy(src, dst)
    return dst, recalc(dst)


def close(a, b, tol=0.02):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def col_values(ws, spec, key, last=None):
    c = spec.letters[key]
    return [ws[f"{c}{r}"].value for r in range(spec.first, (last or spec.last) + 1)]

# ------------------------------------------------------------------ 1. recalc / error scan
demo_calc, res = copy_calc(DEMO, "demo_calc.xlsx")
check("DEMO recalculates with zero formula errors", res.get("status") == "success" and res.get("total_errors") == 0, res)
DEMO_FORMULAS = res.get("total_formulas")
clean_calc, res2 = copy_calc(CLEAN, "clean_calc.xlsx")
check("CLEAN (blank state) recalculates with zero formula errors", res2.get("status") == "success" and res2.get("total_errors") == 0, res2)

wb = load_workbook(demo_calc, data_only=True)
wbf = load_workbook(DEMO)  # formulas
data = demo_data.generate()
jobs = data["_jobs_full"]
AS_OF = demo_data.AS_OF
L = B.LOGS

# ------------------------------------------------------------------ 2. structure
names = [n for n, _ in B.SHEETS]
check("All 22 modules present as tabs (+QC CHECKLISTS)", wbf.sheetnames == names, wbf.sheetnames)
# hyperlinks
bad = []
nlinks = 0
for ws in wbf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink is not None and c.hyperlink.location:
                nlinks += 1
                tgt = c.hyperlink.location.split("!")[0].strip("'")
                if tgt not in wbf.sheetnames:
                    bad.append((ws.title, c.coordinate, tgt))
check(f"All {nlinks} internal navigation links point to existing tabs", not bad and nlinks > 40, bad)
# data validations
dv_count = 0
dv_bad = []
for ws in wbf.worksheets:
    for dv in ws.data_validations.dataValidation:
        dv_count += 1
        f = dv.formula1 or ""
        m = re.findall(r"'([^']+)'!", f)
        for s in m:
            if s not in wbf.sheetnames:
                dv_bad.append((ws.title, f))
check(f"{dv_count} dropdown rules reference valid lists", dv_count > 50 and not dv_bad, dv_bad)
# required dropdown columns
req = {"LEADS": ["source", "service", "stage", "lost"], "JOBS": ["customer", "property", "service", "worker", "status", "plan"],
       "RECURRING": ["customer", "freq", "status"], "REVENUE": ["job", "method"], "EXPENSES": ["cat"],
       "REVIEWS": ["platform", "status"], "REFERRALS": ["status", "rstatus"]}
missing = []
for sh, keys in req.items():
    ws = wbf[sh]
    covered = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    for k in keys:
        col = L[sh].letters[k]
        if f"{col}{L[sh].first}:" not in covered:
            missing.append((sh, k))
check("Key columns have dropdowns (stage, status, customer, service, worker, etc.)", not missing, missing)
# filters
nofilter = [s for s in ["LEADS", "CUSTOMERS", "PROPERTIES", "JOBS", "RECURRING", "REVENUE", "EXPENSES", "MILEAGE", "SUPPLIES",
                        "REVIEWS", "REFERRALS", "QUALITY CONTROL", "FOLLOW-UPS", "STAFF & TASKS"] if not wbf[s].auto_filter.ref]
check("Filters enabled on every log table", not nofilter, nofilter)
# charts
charts = {ws.title: len(ws._charts) for ws in wbf.worksheets if ws._charts}
check("Charts present on DASHBOARD (2), MONTHLY (1), PROFITABILITY (1)", charts == {"DASHBOARD": 2, "MONTHLY": 1, "PROFITABILITY": 1}, charts)
# formula consistency down each auto column
incons = []
for key, spec in L.items():
    ws = wbf[spec.sheet]
    for col in spec.cols:
        if col.kind != "auto":
            continue
        c = spec.letters[col.key]
        f1 = ws[f"{c}{spec.first}"].value
        fl = ws[f"{c}{spec.last}"].value
        norm = lambda f, r: re.sub(rf"(?<![$\d]){r}(?!\d)", "{r}", f)
        if norm(f1, spec.first) != norm(fl, spec.last):
            incons.append((key, col.key))
check("Automatic-column formulas are identical (row-relative) from first to last row", not incons, incons)

# ------------------------------------------------------------------ 3. JOBS math vs independent model
J = L["JOBS"]; ws = wb["JOBS"]
pay = defaultdict(float); fee = defaultdict(float)
feepct = {"Card": 0.029, "Online Invoice": 0.029}
for p in data["payments"]:
    pay[p["job"]] += p["amount"]; fee[p["job"]] += round(p["amount"] * feepct.get(p["method"], 0), 2)
miles = defaultdict(float)
for m in data["mileage"]:
    if m.get("job"):
        miles[m["job"]] += (m["odo2"] - m["odo1"]) if m.get("odo1") is not None else m["miles"]
rates = {s["name"]: s["rate"] for s in data["staff"]}
props = {p["addr"]: p for p in data["properties"]}
supm = {t[0]: t[6] for t in B.PRICING_TYPES}
errs = []
for i, j in enumerate(jobs):
    r = J.first + i
    g = lambda k: ws[f"{J.letters[k]}{r}"].value
    if g("customer") != j["customer"] or g("price") != j["price"]:
        errs.append(("row mismatch", r)); continue
    p = props[j["property"]]
    model = round(demo_data.hours(j["service"], p["sqft"], p["beds"], p["baths"]), 2)
    if not close(g("model"), model):
        errs.append(("model", r, g("model"), model))
    hrs = j["act"] if j["act"] else (j["est"] if j["est"] else model)
    if not close(g("hrs"), hrs):
        errs.append(("hrs", r, g("hrs"), hrs))
    if j["status"] in ("Cancelled", "No-Show"):
        if g("gp") not in (None, ""):
            errs.append(("cancelled gp", r))
        continue
    labor = round(hrs * rates[j["worker"]] * 1.12, 2)
    sup = round(hrs * 2.5 * supm[j["service"]], 2)
    veh = round(miles[j["id"]] * 0.7, 2)
    gp = j["price"] - labor - sup - veh - fee[j["id"]]
    for k, v in (("labor", labor), ("sup", sup), ("vehicle", veh), ("fees", fee[j["id"]]), ("gp", gp), ("paid", pay[j["id"]])):
        if not close(g(k), v, 0.03):
            errs.append((k, r, g(k), v))
    if j["status"] == "Completed":
        bal = max(0, j["price"] - pay[j["id"]])
        if not close(g("balance"), bal):
            errs.append(("balance", r, g("balance"), bal))
        exp = "Paid" if bal <= 0.005 else ("Partial" if pay[j["id"]] > 0 else "Unpaid")
        if g("paystat") != exp:
            errs.append(("paystat", r, g("paystat"), exp))
        if not close(g("margin"), gp / j["price"], 0.0005):
            errs.append(("margin", r))
check(f"JOBS: hours, labor, supplies, vehicle, fees, gross profit, margin, paid, balance, payment status match independent model ({len(jobs)} jobs)", not errs, errs[:5])
blank_row = J.first + len(jobs) + 5
blank_vals = [ws[f"{c}{blank_row}"].value for c in [J.letters[k] for k in ("model", "hrs", "invoice", "gp", "margin", "key", "urank")]]
check("JOBS: empty rows show blanks (no zeros/errors)", all(v in (None, "") for v in blank_vals), blank_vals)

# ------------------------------------------------------------------ 4. MONTHLY & DASHBOARD totals
ms = wb["MONTHLY"]; MR = B.MONTHLY_ROWS
rev_m = defaultdict(float); jobs_m = defaultdict(int)
for j in jobs:
    if j["status"] == "Completed":
        rev_m[j["date"].month] += j["price"]; jobs_m[j["date"].month] += 1
errs = []
for m in range(1, 13):
    col = B.L(1 + m)
    if not close(ms[f"{col}{MR['rev']}"].value, rev_m[m]):
        errs.append(("rev", m, ms[f"{col}{MR['rev']}"].value, rev_m[m]))
    if ms[f"{col}{MR['jobs']}"].value != jobs_m[m]:
        errs.append(("jobs", m))
check("MONTHLY: revenue and jobs completed per month match independent totals", not errs, errs)
leads_m = defaultdict(int); booked_m = defaultdict(int)
for l in data["leads"]:
    leads_m[l["date"].month] += 1
    if l["stage"] in ("Booked", "Recurring Client"):
        booked_m[l["date"].month] += 1
errs = [m for m in range(1, 13) if ms[f"{B.L(1+m)}{MR['leads']}"].value != leads_m[m] or ms[f"{B.L(1+m)}{MR['booked']}"].value != booked_m[m]]
check("MONTHLY: leads and booked leads per month match", not errs, errs)
oh_cats = {k for k, v in zip([x for x in B.LISTS if x[0] == "expcat"][0][2], [x for x in B.LISTS if x[0] == "expoh"][0][2]) if v == "Yes"}
oh9 = sum(e["amount"] for e in data["expenses"] if e["date"].month == 9 and e["cat"] in oh_cats)
check("MONTHLY: overhead uses only categories flagged as overhead (Sep)", close(ms[f"J{MR['oh']}"].value, oh9), (ms[f"J{MR['oh']}"].value, oh9))
gp9 = ms[f"J{MR['rev']}"].value - sum(ms[f"J{MR[k]}"].value for k in ("labor", "sup", "veh", "fees"))
check("MONTHLY: gross profit = revenue − labor − supplies − vehicle − fees", close(ms[f"J{MR['gp']}"].value, gp9))
check("MONTHLY: operating profit = gross profit − overhead", close(ms[f"J{MR['op']}"].value, ms[f"J{MR['gp']}"].value - ms[f"J{MR['oh']}"].value))

ds = wb["DASHBOARD"]
check("DASHBOARD revenue tile = MONTHLY reporting-month revenue", close(ds["B6"].value, rev_m[9]), (ds["B6"].value, rev_m[9]))
check("DASHBOARD jobs-completed tile = September completed jobs", ds["B16"].value == jobs_m[9], (ds["B16"].value, jobs_m[9]))
n9 = leads_m[9]; b9 = booked_m[9]
check("DASHBOARD conversion rate = booked ÷ new leads (reporting month)", close(ds["H11"].value, b9 / n9, 0.0005), (ds["H11"].value, b9 / n9))
q9 = [l["quote"] for l in data["leads"] if l["date"].month == 9 and l.get("quote")]
check("DASHBOARD average quote", close(ds["J11"].value, sum(q9) / len(q9)), (ds["J11"].value, sum(q9) / len(q9)))
unpaid_total = sum(max(0, j["price"] - pay[j["id"]]) for j in jobs if j["status"] == "Completed")
check("DASHBOARD unpaid balance tile = sum of completed-job balances", close(ds["L6"].value, unpaid_total), (ds["L6"].value, unpaid_total))
nxt7 = sum(1 for j in jobs if j["status"] == "Scheduled" and AS_OF <= j["date"] <= AS_OF + dt.timedelta(days=7))
check("DASHBOARD jobs next 7 days", ds["H16"].value == nxt7, (ds["H16"].value, nxt7))
active = sum(1 for p in data["recurring"] if p["status"] == "Active")
check("DASHBOARD active recurring plans", ds["J16"].value == active, (ds["J16"].value, active))
mv = sum(p["price"] * {"Weekly": 4.33, "Biweekly": 2.17, "Every 4 weeks": 1.08, "Monthly": 1}[p["freq"]] for p in data["recurring"] if p["status"] == "Active")
check("DASHBOARD recurring monthly value", close(ds["L16"].value, mv, 0.05), (ds["L16"].value, mv))
fu_due = sum(1 for l in data["leads"] if l["stage"] not in ("Booked", "Recurring Client", "Lost/Declined") and l.get("next") and l["next"] <= AS_OF)
fu_due += sum(1 for f in data["followups"] if f["status"] == "Open" and f["due"] <= AS_OF)
check("DASHBOARD follow-ups due = overdue/due-today leads + open log items", ds["L11"].value == fu_due, (ds["L11"].value, fu_due))
stages = [x for x in B.LISTS if x[0] == "stages"][0][2]
pipe_ok = True
for row in range(25, 40):
    if ds[f"I{row}"].value in stages:
        s = ds[f"I{row}"].value
        if ds[f"K{row}"].value != sum(1 for l in data["leads"] if l["stage"] == s):
            pipe_ok = False
check("DASHBOARD pipeline-by-stage counts match leads", pipe_ok)

# ------------------------------------------------------------------ 5. CRM automation
LE = L["LEADS"]; wl = wb["LEADS"]
errs = []
for i, l in enumerate(data["leads"]):
    r = LE.first + i
    g = lambda k: wl[f"{LE.letters[k]}{r}"].value
    exp_out = "Booked" if l["stage"] in ("Booked", "Recurring Client") else ("Lost" if l["stage"] == "Lost/Declined" else "Open")
    if g("outcome") != exp_out:
        errs.append(("outcome", r))
    if exp_out == "Open":
        n = l.get("next")
        exp = "Set date" if not n else ("Overdue" if n < AS_OF else "Due today" if n == AS_OF else "Due soon" if n <= AS_OF + dt.timedelta(days=2) else "Scheduled")
        if g("fu") != exp:
            errs.append(("fu", r, g("fu"), exp))
check("LEADS: booking outcome and follow-up status correct for every lead", not errs, errs[:5])
flag = [wl[f"{LE.letters['check']}{LE.first+i}"].value for i in range(len(data["leads"]))]
check("LEADS: booked lead not in CUSTOMERS is flagged 'Add to CUSTOMERS'", flag.count("Add to CUSTOMERS") == 1, flag.count("Add to CUSTOMERS"))

# ------------------------------------------------------------------ 6. recurring
RC = L["RECURRING"]; wr = wb["RECURRING"]
errs = []
for i, p in enumerate(data["recurring"]):
    r = RC.first + i
    pid = f"R-{i+1:03d}"
    done = [j["date"] for j in jobs if j["plan"] == pid and j["status"] == "Completed"]
    fut = [j["date"] for j in jobs if j["plan"] == pid and j["status"] == "Scheduled" and j["date"] >= AS_OF]
    last = max(done) if done else None
    if fut:
        due = min(fut)
    elif last:
        if p["freq"] == "Monthly":
            mm = last.month % 12 + 1
            due = dt.date(last.year + (last.month == 12), mm, last.day)
        else:
            due = last + dt.timedelta(days={"Weekly": 7, "Biweekly": 14, "Every 4 weeks": 28}[p["freq"]])
    else:
        due = p["start"]
    g = lambda k: wr[f"{RC.letters[k]}{r}"].value
    if g("visits") != len(done):
        errs.append(("visits", pid))
    if due and (g("due") is None or g("due").date() != due):
        errs.append(("due", pid, g("due"), due))
    if p["status"] != "Active":
        exp = "—"
    elif fut:
        exp = "Booked"
    elif due < AS_OF:
        exp = "Overdue — book now"
    elif due <= AS_OF + dt.timedelta(days=7):
        exp = "Due this week — book"
    else:
        exp = "Upcoming"
    if g("book") != exp:
        errs.append(("book", pid, g("book"), exp))
check("RECURRING: visits completed, next due date and booking status correct for every plan", not errs, errs)
sc = wb["SCHEDULE"]
check("SCHEDULE: overdue recurring plan appears in 'needs booking' list", sc["B26"].value == "R-001" if False else any(sc[f"B{r}"].value == "R-001" for r in range(24, 40)), [sc[f"B{r}"].value for r in range(24, 40)])

# ------------------------------------------------------------------ 7. schedule
wk = dt.date(2026, 9, 28)
errs = []
for d in range(7):
    day = wk + dt.timedelta(days=d)
    day_jobs = sorted([j for j in jobs if j["date"] == day and j["status"] not in ("Cancelled", "Rescheduled")], key=lambda x: x["time"])
    col = B.L(2 + d)
    shown = [sc[f"{col}{r}"].value for r in range(9, 19)]
    shown_ids = [s.split("· ")[-1] for s in shown if s]
    if shown_ids != [j["id"] for j in day_jobs][:10]:
        errs.append((str(day), shown_ids, [j["id"] for j in day_jobs]))
    if sc[f"{col}19"].value != len(day_jobs):
        errs.append(("count", str(day)))
check("SCHEDULE: each day lists exactly that day's jobs in start-time order, with correct daily counts", not errs, errs)

# ------------------------------------------------------------------ 8. quote builder
qb = wb["QUOTE BUILDER"]
check("QUOTE BUILDER demo: estimated hours 7.50", close(qb["G7"].value, 7.5))
check("QUOTE BUILDER demo: price before tax $522.50", close(qb["G16"].value, 522.5))
check("QUOTE BUILDER demo: direct costs $283.37 / gross profit $239.13", close(qb["G26"].value, 283.37) and close(qb["G28"].value, 239.13))
check("QUOTE BUILDER demo: overhead share = 9.25 h × $5 = $46.25; profit after overhead $192.88", close(qb["G36"].value, 46.25) and close(qb["G37"].value, 192.88), (qb["G36"].value, qb["G37"].value))
check("QUOTE BUILDER demo: break-even $276.23 and target-margin price $569.47", close(qb["G32"].value, 276.23) and close(qb["G33"].value, 569.47))


def quote_model(t, sq, bd, ba, cond, freq, workers, ovh, ovr, pay_, supov, miles_, travel, addons, disc, tax, card):
    rate, sph, mb, mba, setup = demo_data.TYPES[t]
    cm = dict(B.CONDITIONS)[cond]
    fr = {f[0]: f for f in B.FREQS}[freq]
    ad = {a[0]: a for a in B.ADDONS}
    model = round((sq / sph + bd * mb / 60 + ba * mba / 60 + setup / 60) * cm + 1e-9, 2)
    addmin = sum(ad[a][2] * q for a, q in addons)
    hrs = (ovh if ovh else model) + addmin / 60
    r = ovr if ovr else rate
    base = max(120, round((ovh if ovh else model) * r, 2))
    fd = -round(base * fr[1], 2)
    adp = sum(ad[a][1] * q for a, q in addons)
    price = max(0, base + fd + adp - disc)
    taxv = round(price * tax, 2)
    total = price + taxv
    labor = round(hrs * pay_ * 1.12, 2)
    trav = round(travel / 60 * max(1, workers) * pay_ * 1.12, 2)
    veh = round(miles_ * 0.7, 2)
    sup = supov if supov is not None else round(hrs * 2.5 * supm[t], 2)
    adc = sum(ad[a][3] * q for a, q in addons)
    feev = round(total * 0.029, 2) if card else 0
    costs = labor + trav + veh + sup + adc + feev
    gp = price - costs
    return dict(model=model, price=price, total=total, costs=costs, gp=gp, margin=gp / price if price else 0)


scen = [
    dict(name="Standard weekly, 1,600 sq ft, light condition, 1 worker, 15% tax, cash",
         inp=dict(B9="Standard Clean", B10=1600, B11=3, B12=2, B13="Light", B14="Weekly", B15=1, B16=None, B17=None, B18=20, B19=None, B20=8, B21=15, B33=0, B34=0.15, B35="No"),
         args=("Standard Clean", 1600, 3, 2, "Light", "Weekly", 1, None, None, 20, None, 8, 15, [], 0, 0.15, False)),
    dict(name="Commercial, 6,000 sq ft, 4 restrooms, heavy, overrides (hours 5, rate $60), $25 discount",
         inp=dict(B9="Commercial Clean", B10=6000, B11=0, B12=4, B13="Heavy", B14="One-time", B15=3, B16=5, B17=60, B18=22, B19=30, B20=20, B21=10, B33=25, B34=0, B35="Yes"),
         args=("Commercial Clean", 6000, 0, 4, "Heavy", "One-time", 3, 5, 60, 22, 30, 20, 10, [], 25, 0, True)),
    dict(name="Move-out, very heavy, 2 add-ons x2, biweekly frequency (discount applies to base only)",
         inp=dict(B9="Move-In/Move-Out", B10=1800, B11=3, B12=2, B13="Very heavy", B14="Biweekly", B15=2, B16=None, B17=None, B18=19, B19=None, B20=14, B21=20, B33=0, B34=0, B35="Yes",
                  A25="Inside cabinets", B25=2, A26="Wall spot cleaning", B26=2, A27=None, B27=None),
         args=("Move-In/Move-Out", 1800, 3, 2, "Very heavy", "Biweekly", 2, None, None, 19, None, 14, 20, [("Inside cabinets", 2), ("Wall spot cleaning", 2)], 0, 0, True)),
    dict(name="Tiny job hits minimum price ($120 floor)",
         inp=dict(B9="Airbnb Turnover", B10=400, B11=0, B12=1, B13="Light", B14="One-time", B15=1, B16=None, B17=None, B18=20, B19=None, B20=4, B21=0, B33=0, B34=0, B35="No",
                  A25=None, B25=None, A26=None, B26=None, A27=None, B27=None),
         args=("Airbnb Turnover", 400, 0, 1, "Light", "One-time", 1, None, None, 20, None, 4, 0, [], 0, 0, False)),
]
for i, s in enumerate(scen):
    wq = load_workbook(DEMO)
    q = wq["QUOTE BUILDER"]
    for rr_ in range(25, 31):
        q[f"A{rr_}"] = None; q[f"B{rr_}"] = None
    for k, v in s["inp"].items():
        q[k] = v
    p = os.path.join(WORK, f"quote_{i}.xlsx")
    wq.save(p)
    rr = recalc(p)
    qv = load_workbook(p, data_only=True)["QUOTE BUILDER"]
    m = quote_model(*s["args"])
    ok = (rr.get("total_errors") == 0 and close(qv["G7"].value, m["model"]) and close(qv["G16"].value, m["price"], 0.05) and
          close(qv["G18"].value, m["total"], 0.05) and close(qv["G26"].value, m["costs"], 0.05) and close(qv["G28"].value, m["gp"], 0.05) and
          close(qv["G29"].value, m["margin"], 0.001))
    check(f"QUOTE scenario: {s['name']}", ok, dict(sheet=[qv[f'G{r}'].value for r in (7, 16, 18, 26, 28, 29)], model=m))

# ------------------------------------------------------------------ 9. scenario: CRM stage change + added rows
w2 = load_workbook(DEMO)
le = w2["LEADS"]
# move first open 'Quote Sent' lead to Booked
target = None
for i, l in enumerate(data["leads"]):
    if l["stage"] == "Quote Sent":
        target = (i, l); break
le[f"{LE.letters['stage']}{LE.first + target[0]}"] = "Booked"
# add new customer, property, job, payment in the next empty rows
CU, PR_, RV = L["CUSTOMERS"], L["PROPERTIES"], L["REVENUE"]
ncust = len(data["customers"]); nprop = len(data["properties"]); njob = len(jobs); npay = len(data["payments"])
w2["CUSTOMERS"][f"{CU.letters['name']}{CU.first + ncust}"] = "QA Test Customer"
w2["CUSTOMERS"][f"{CU.letters['ctype']}{CU.first + ncust}"] = "Residential"
w2["PROPERTIES"][f"{PR_.letters['addr']}{PR_.first + nprop}"] = "1 QA Street"
w2["PROPERTIES"][f"{PR_.letters['customer']}{PR_.first + nprop}"] = "QA Test Customer"
for k, v in (("sqft", 2000), ("beds", 3), ("baths", 2)):
    w2["PROPERTIES"][f"{PR_.letters[k]}{PR_.first + nprop}"] = v
jr = J.first + njob
for k, v in (("date", dt.date(2026, 9, 30)), ("time", dt.time(16, 0)), ("customer", "QA Test Customer"), ("property", "1 QA Street"),
             ("service", "Standard Clean"), ("worker", "Maria Lopez"), ("crew", 1), ("price", 200), ("status", "Completed")):
    w2["JOBS"][f"{J.letters[k]}{jr}"] = v
new_id = f"J-{njob+1:04d}"
w2["REVENUE"][f"{RV.letters['date']}{RV.first + npay}"] = dt.date(2026, 9, 30)
w2["REVENUE"][f"{RV.letters['job']}{RV.first + npay}"] = new_id
w2["REVENUE"][f"{RV.letters['amount']}{RV.first + npay}"] = 150
w2["REVENUE"][f"{RV.letters['method']}{RV.first + npay}"] = "Cash"
p2 = os.path.join(WORK, "scenario_rows.xlsx"); w2.save(p2)
rr = recalc(p2)
v2 = load_workbook(p2, data_only=True)
d2 = v2["DASHBOARD"]
check("Scenario recalculates with zero errors", rr.get("total_errors") == 0, rr)
lrow = LE.first + target[0]
check("CRM: changing a lead's Stage to Booked flips outcome to 'Booked' and stops follow-up",
      v2["LEADS"][f"{LE.letters['outcome']}{lrow}"].value == "Booked" and v2["LEADS"][f"{LE.letters['fu']}{lrow}"].value == "—")
check("CRM: dashboard pipeline/booked counts update after stage change",
      d2["F11"].value == ds["F11"].value + (1 if target[1]["date"].month == 9 else 0))
check("Added rows: new job flows to dashboard revenue (+$200) and jobs completed (+1)",
      close(d2["B6"].value, ds["B6"].value + 200) and d2["B16"].value == ds["B16"].value + 1, (d2["B6"].value, d2["B16"].value))
g2 = lambda k: v2["JOBS"][f"{J.letters[k]}{jr}"].value
check("Added rows: new job gets auto estimate from new property (2,000 sq ft 3/2 standard = 3.42 h), partial payment status, balance $50",
      close(g2("model"), 3.42) and g2("paystat") == "Partial" and close(g2("balance"), 50), (g2("model"), g2("paystat"), g2("balance")))
crow = CU.first + ncust
check("Added rows: new customer shows 1 completed job, $200 lifetime revenue, 'Collect balance' action",
      v2["CUSTOMERS"][f"{CU.letters['jobs']}{crow}"].value == 1 and close(v2["CUSTOMERS"][f"{CU.letters['ltv']}{crow}"].value, 200)
      and v2["CUSTOMERS"][f"{CU.letters['action']}{crow}"].value == "Collect balance")
check("Added rows: unpaid job appears in FOLLOW-UPS unpaid list",
      any(v2["FOLLOW-UPS"][f"H{r}"].value == new_id for r in range(6, 21)))
check("Added rows: new job appears on SCHEDULE for Wednesday Sep 30",
      any((v2["SCHEDULE"][f"D{r}"].value or "").endswith(new_id) for r in range(9, 19)))

# ------------------------------------------------------------------ 10. sample-data removal
w3 = load_workbook(DEMO)
for key, spec in L.items():
    ws3 = w3[spec.sheet]
    for col in spec.cols:
        if col.kind == "input":
            c = spec.letters[col.key]
            for r in range(spec.first, spec.last + 1):
                ws3[f"{c}{r}"].value = None
p3 = os.path.join(WORK, "demo_cleared.xlsx"); w3.save(p3)
rr = recalc(p3)
v3 = load_workbook(p3, data_only=True)
d3 = v3["DASHBOARD"]
tiles = [d3[f"{c}{r}"].value for r in (6, 11, 16, 21) for c in "BDFHJL"]
check("Sample-data removal: clearing all white input columns leaves zero formula errors", rr.get("total_errors") == 0, rr)
check("Sample-data removal: every dashboard tile shows 0 or blank", all(v in (None, "", 0) for v in tiles), tiles)
vc = load_workbook(clean_calc, data_only=True)["DASHBOARD"]
tiles_c = [vc[f"{c}{r}"].value for r in (6, 11, 16, 21) for c in "BDFHJL"]
check("CLEAN file: every dashboard tile shows 0 or blank", all(v in (None, "", 0) for v in tiles_c), tiles_c)
wc = load_workbook(CLEAN)
leftover = []
for key, spec in L.items():
    for col in spec.cols:
        if col.kind == "input":
            c = spec.letters[col.key]
            if any(wc[spec.sheet][f"{c}{r}"].value not in (None, "") for r in range(spec.first, spec.last + 1)):
                leftover.append((key, col.key))
check("CLEAN file contains no sample records in any log", not leftover, leftover)
check("CLEAN file: as-of date and reporting month blank (uses today)", wc["SETTINGS"][f"B{B.SETROW['AsOfInput']}"].value is None and wc["SETTINGS"][f"B{B.SETROW['RMInput']}"].value is None)

# ------------------------------------------------------------------ 10b. CLEAN date engine + first job
today = dt.date.today()
cs = load_workbook(clean_calc, data_only=True)["SETTINGS"]
asof = cs[f"B{B.SETROW['AsOf']}"].value; rm = cs[f"B{B.SETROW['ReportMonth']}"].value; ry = cs[f"B{B.SETROW['ReportYear']}"].value
check("CLEAN file: system date = today, reporting month = this month, reporting year = this year",
      asof is not None and asof.date() == today and rm is not None and rm.date() == today.replace(day=1) and ry == today.year, (asof, rm, ry))
for nm in ("AsOf", "ReportMonth", "ReportYear"):
    pass
wc4 = load_workbook(CLEAN)
check("CLEAN file: SETTINGS formula cells (AsOf, ReportMonth, ReportYear) contain formulas",
      all(str(wc4["SETTINGS"][f"B{B.SETROW[k]}"].value).startswith("=") for k in ("AsOf", "ReportMonth", "ReportYear")))
st, cu_, pr_, jb = L["STAFF & TASKS"], L["CUSTOMERS"], L["PROPERTIES"], L["JOBS"]
wc4["STAFF & TASKS"][f"{st.letters['name']}{st.first}"] = "Test Cleaner"
wc4["STAFF & TASKS"][f"{st.letters['rate']}{st.first}"] = 20
wc4["CUSTOMERS"][f"{cu_.letters['name']}{cu_.first}"] = "Test Customer"
wc4["CUSTOMERS"][f"{cu_.letters['ctype']}{cu_.first}"] = "Residential"
for k, v in (("addr", "1 Test Street"), ("customer", "Test Customer"), ("sqft", 2000), ("beds", 3), ("baths", 2)):
    wc4["PROPERTIES"][f"{pr_.letters[k]}{pr_.first}"] = v
for k, v in (("date", today), ("time", dt.time(9, 0)), ("customer", "Test Customer"), ("property", "1 Test Street"),
             ("service", "Standard Clean"), ("worker", "Test Cleaner"), ("crew", 1), ("price", 200), ("status", "Completed")):
    wc4["JOBS"][f"{jb.letters[k]}{jb.first}"] = v
p4 = os.path.join(WORK, "clean_first_job.xlsx"); wc4.save(p4)
rr4 = recalc(p4)
v4 = load_workbook(p4, data_only=True)
gj = lambda k: v4["JOBS"][f"{jb.letters[k]}{jb.first}"].value
check("CLEAN first job dated today: 0 errors, auto estimate 3.42 h, Unpaid, dashboard revenue $200 and jobs completed 1",
      rr4.get("total_errors") == 0 and close(gj("model"), 3.42) and gj("paystat") == "Unpaid"
      and close(v4["DASHBOARD"]["B6"].value, 200) and v4["DASHBOARD"]["B16"].value == 1,
      (rr4.get("total_errors"), gj("model"), gj("paystat"), v4["DASHBOARD"]["B6"].value, v4["DASHBOARD"]["B16"].value))

# ------------------------------------------------------------------ 11. text scan (spelling of fixed UI text)
words = set()
for ws in wbf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                for w in re.findall(r"[A-Za-z]{4,}", c.value):
                    words.add(w.lower())
json.dump(sorted(words), open(os.path.join(WORK, "ui_words.json"), "w"))

summary = {"passed": sum(r["pass"] for r in results), "failed": sum(not r["pass"] for r in results),
           "demo_formulas": DEMO_FORMULAS, "clean_formulas": res2.get("total_formulas"), "results": results}
json.dump(summary, open(os.path.join(WORK, "qa_results.json"), "w"), indent=1, default=str)
print(f"\n{summary['passed']} passed, {summary['failed']} failed")
