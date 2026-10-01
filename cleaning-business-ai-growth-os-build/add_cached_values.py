"""Stores each formula's calculated result as its cached value (<v>) so Excel Protected View,
file previews and Drive/Quick Look show numbers before recalculation. Formulas are not changed.
Results come from a full LibreOffice recalculation of a temporary copy.
Usage: python add_cached_values.py <workbook.xlsx> [...]   (rewrites in place)"""
import sys, os, re, json, shutil, subprocess, tempfile, zipfile, datetime as dt
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

RECALC = "/root/.claude/skills/synced/537785f1-f0de-401f-bf47-0f2c32573f85_73cbe849-f847-4991-82d7-5f1ecd5f11a9/xlsx/scripts/recalc.py"
CELL = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: [A-Za-z:]+="[^"]*")*)><f>([^<]*)</f><v></v></c>')


def sheet_parts(z):
    wbx = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid2t = {m.group(2): m.group(1).lstrip("/") for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)}
    out = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        name = m.group(1).replace("&amp;", "&")
        t = rid2t[m.group(2)]
        out[t if t.startswith("xl/") else "xl/" + t] = name
    return out


def encode(v):
    """Return (type_attr, value_text) for a cached value."""
    if v is None or v == "":
        return "str", ""
    if isinstance(v, bool):
        return "b", "1" if v else "0"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return "n", repr(float(to_excel(v)))
    if isinstance(v, (int, float)):
        return "n", repr(int(v)) if float(v).is_integer() else repr(float(v))
    s = str(v)
    if s.startswith("#"):
        raise ValueError(f"formula error value {s}")
    return "str", escape(s)


def process(path):
    tmpdir = tempfile.mkdtemp()
    calc = os.path.join(tmpdir, "calc.xlsx")
    shutil.copy(path, calc)
    r = json.loads(subprocess.run([sys.executable, RECALC, calc, "300"], capture_output=True, text=True).stdout)
    if r.get("status") != "success" or r.get("total_errors"):
        raise SystemExit(f"recalc failed: {r}")
    vals = load_workbook(calc, data_only=True)
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = {i.filename: z.read(i.filename) for i in infos}
        parts = sheet_parts(z)
    filled = 0
    for part, name in parts.items():
        ws = vals[name]
        xml = data[part].decode("utf-8")

        def rep(m):
            nonlocal filled
            ref, attrs, f = m.group(1), m.group(2), m.group(3)
            attrs = re.sub(r' t="[^"]*"', "", attrs)
            t, v = encode(ws[ref].value)
            filled += 1
            tattr = "" if t == "n" else f' t="{t}"'
            return f'<c r="{ref}"{attrs}{tattr}><f>{f}</f><v>{v}</v></c>'
        new = CELL.sub(rep, xml)
        data[part] = new.encode("utf-8")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=i.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = i.external_attr
            zo.writestr(zi, data[i.filename])
    os.replace(tmp, path)
    shutil.rmtree(tmpdir)
    return filled


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(os.path.basename(p), "cached values written:", process(p))
