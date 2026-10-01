"""Builds MARKETING-KIT: editable PPTX templates (custom sizes) + copy document."""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from docx.oxml.ns import qn

OUT = sys.argv[1]
os.makedirs(os.path.join(OUT, "editable-pptx"), exist_ok=True)
PRIMARY = RGBColor(0x0F, 0x4C, 0x5C); ACCENT = RGBColor(0x2B, 0xB3, 0xA3); LIGHT = RGBColor(0xF3, 0xF7, 0xF8)
WHITE = RGBColor(255, 255, 255); TEXT = RGBColor(0x1E, 0x2A, 0x30); AMBER = RGBColor(0xE0, 0x8E, 0x2B); MINT = RGBColor(0xD8, 0xF0, 0xEC)
FONT = "Arial"


def deck(w, h):
    p = Presentation(); p.slide_width = Inches(w); p.slide_height = Inches(h)
    return p


def slide(p):
    return p.slides.add_slide(p.slide_layouts[6])


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, t, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT; r.font.italic = italic
    return tb


def sparkle(s, x, y, size, color=ACCENT):
    sh = s.shapes.add_shape(MSO_SHAPE.STAR_4_POINT, Inches(x), Inches(y), Inches(size), Inches(size))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()


def photo_box(s, x, y, w, h, label="[ YOUR PHOTO ]"):
    b = rect(s, x, y, w, h, LIGHT, line=RGBColor(0xC9, 0xD3, 0xD7))
    text(s, x, y + h / 2 - 0.25, w, 0.5, label, 12, True, RGBColor(0x6B, 0x7B, 0x83), PP_ALIGN.CENTER)


def save(p, name):
    p.save(os.path.join(OUT, "editable-pptx", name))

# 1. Service flyer 8.5 x 11
p = deck(8.5, 11); s = slide(p)
rect(s, 0, 0, 8.5, 3.6, PRIMARY)
sparkle(s, 7.0, 0.5, 0.8); sparkle(s, 7.7, 1.3, 0.4, MINT)
text(s, 0.6, 0.5, 6.3, 0.5, "[YOUR BUSINESS NAME]", 14, True, MINT)
text(s, 0.6, 1.0, 7, 1.6, ["Come home to a", "clean house."], 40, True, WHITE)
text(s, 0.6, 2.75, 7, 0.6, "Reliable residential cleaning in [Service Area]", 16, False, MINT)
photo_box(s, 0.6, 3.95, 3.5, 2.6, "[ PHOTO: your team or a clean room ]")
text(s, 4.4, 3.95, 3.6, 0.4, "SERVICES", 12, True, ACCENT)
text(s, 4.4, 4.35, 3.6, 2.4, ["•  Standard cleaning", "•  Deep cleaning", "•  Move-in / move-out", "•  Short-term rental turnovers", "•  Weekly, biweekly or monthly plans"], 14, False, TEXT)
rect(s, 0.6, 6.85, 7.3, 1.45, LIGHT)
text(s, 0.8, 6.95, 2.3, 1.3, ["Same team,", "every visit"], 15, True, PRIMARY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 3.1, 6.95, 2.3, 1.3, ["Room-by-room", "checklist"], 15, True, PRIMARY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 5.4, 6.95, 2.3, 1.3, ["[Your", "differentiator]"], 15, True, PRIMARY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rect(s, 0.6, 8.6, 7.3, 1.6, ACCENT, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.9, 8.7, 4.6, 1.4, ["Get your free quote", "[Phone]  ·  [Website]"], 22, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 6.25, 8.75, 1.4, 1.3, WHITE)
text(s, 6.25, 9.15, 1.4, 0.5, "[QR]", 12, True, PRIMARY, PP_ALIGN.CENTER)
text(s, 0.6, 10.45, 7.3, 0.3, "[Optional offer line — terms and dates]   ·   [License/insurance statement only if true for you]", 8, False, RGBColor(0x6B, 0x7B, 0x83), PP_ALIGN.CENTER)
save(p, "01-Service-Flyer-8.5x11.pptx")

# 2. Door hanger 4.25 x 11 (front/back)
p = deck(4.25, 11)
s = slide(p)
rect(s, 0, 0, 4.25, 11, WHITE)
c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.25), Inches(0.45), Inches(1.75), Inches(1.75)); c.fill.solid(); c.fill.fore_color.rgb = LIGHT; c.line.color.rgb = PRIMARY
text(s, 0.2, 2.25, 3.85, 0.3, "(cut out circle for the door handle)", 8, False, RGBColor(0x6B, 0x7B, 0x83), PP_ALIGN.CENTER)
rect(s, 0, 2.8, 4.25, 3.6, PRIMARY)
sparkle(s, 3.3, 3.0, 0.5)
text(s, 0.3, 3.1, 3.7, 0.4, "HI, NEIGHBOR!", 13, True, MINT)
text(s, 0.3, 3.5, 3.7, 2.0, ["We're cleaning", "homes on your", "street this week."], 24, True, WHITE)
text(s, 0.3, 5.6, 3.7, 0.7, "[YOUR BUSINESS NAME]", 13, True, MINT)
text(s, 0.3, 6.65, 3.7, 2.2, ["✓  Standard & deep cleaning", "✓  Move-in / move-out", "✓  Recurring plans", "✓  [Your differentiator]"], 13, False, TEXT)
rect(s, 0.3, 8.9, 3.65, 1.6, ACCENT, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.4, 9.0, 3.45, 1.4, ["[New-client offer]", "Scan or call [Phone]"], 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
s = slide(p)
rect(s, 0, 0, 4.25, 11, LIGHT)
c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.25), Inches(0.45), Inches(1.75), Inches(1.75)); c.fill.solid(); c.fill.fore_color.rgb = WHITE; c.line.color.rgb = PRIMARY
text(s, 0.3, 2.8, 3.7, 0.5, "WHAT'S INCLUDED", 13, True, ACCENT)
text(s, 0.3, 3.3, 3.7, 3.4, ["Kitchen: counters, sink, appliance fronts, floors", "", "Bathrooms: toilets, tubs, showers, mirrors, floors", "",
                             "Living areas: dusting, surfaces, floors", "", "[Edit to match your checklist]"], 12, False, TEXT)
rect(s, 0.3, 7.0, 3.65, 1.8, WHITE)
text(s, 0.4, 7.05, 3.45, 1.7, ["\"[Real customer quote — only with permission]\"", "— [First name], [Area]"], 11, False, PRIMARY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, italic=True)
rect(s, 1.4, 9.0, 1.45, 1.45, WHITE, line=PRIMARY)
text(s, 1.4, 9.5, 1.45, 0.4, "[QR]", 12, True, PRIMARY, PP_ALIGN.CENTER)
text(s, 0.2, 10.5, 3.85, 0.4, "[Website]  ·  [Phone]", 11, True, PRIMARY, PP_ALIGN.CENTER)
save(p, "02-Door-Hanger-4.25x11.pptx")

# 3/4. Referral and review cards 3.5 x 2 (front/back)
for fname, front_title, front_sub, back_lines in [
    ("03-Referral-Card-3.5x2.pptx", "Share the clean.", "Know someone who'd love a cleaner home?",
     ["Give this card to a friend.", "When they book, tell us your name:", "", "Referred by: ____________________", "[Your thank-you — your terms]"]),
    ("04-Review-Card-3.5x2.pptx", "How did we do?", "Your home was cleaned today by ______",
     ["If you loved today's clean, an honest", "review helps our small team grow.", "", "Something not right? Tell us first:", "[Phone]  ·  we'll make it right."]),
]:
    p = deck(3.5, 2)
    s = slide(p); rect(s, 0, 0, 3.5, 2, PRIMARY); sparkle(s, 2.85, 0.2, 0.4)
    text(s, 0.2, 0.2, 2.6, 0.3, "[YOUR BUSINESS NAME]", 8, True, MINT)
    text(s, 0.2, 0.55, 3.1, 0.6, front_title, 22, True, WHITE)
    text(s, 0.2, 1.2, 3.1, 0.5, front_sub, 10, False, MINT)
    rect(s, 0, 1.85, 3.5, 0.15, ACCENT)
    s = slide(p); rect(s, 0, 0, 3.5, 2, WHITE)
    text(s, 0.15, 0.15, 2.25, 1.7, back_lines, 8.5, False, TEXT)
    rect(s, 2.45, 0.35, 0.9, 0.9, LIGHT, line=PRIMARY); text(s, 2.45, 0.6, 0.9, 0.3, "[QR]", 9, True, PRIMARY, PP_ALIGN.CENTER)
    text(s, 2.3, 1.35, 1.15, 0.5, "[Website]", 7, True, PRIMARY, PP_ALIGN.CENTER)
    rect(s, 0, 1.9, 3.5, 0.1, ACCENT)
    save(p, fname)

# 5-7, 9. Social squares (7.5in = 1080px at 144 dpi)
p = deck(7.5, 7.5)
# before/after
s = slide(p); rect(s, 0, 0, 7.5, 7.5, WHITE)
rect(s, 0, 0, 7.5, 1.1, PRIMARY)
text(s, 0.4, 0.2, 6.7, 0.7, "[Room] transformation", 26, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
photo_box(s, 0.3, 1.3, 3.35, 4.9, "[ BEFORE ]"); photo_box(s, 3.85, 1.3, 3.35, 4.9, "[ AFTER ]")
rect(s, 0.3, 1.45, 1.1, 0.4, AMBER, MSO_SHAPE.ROUNDED_RECTANGLE); text(s, 0.3, 1.45, 1.1, 0.4, "BEFORE", 11, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rect(s, 3.85, 1.45, 1.0, 0.4, ACCENT, MSO_SHAPE.ROUNDED_RECTANGLE); text(s, 3.85, 1.45, 1.0, 0.4, "AFTER", 11, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.4, 6.35, 5, 0.5, "[Service type]  ·  [time on site]", 15, True, PRIMARY)
text(s, 0.4, 6.8, 5, 0.4, "[YOUR BUSINESS NAME]  ·  [Website]", 11, False, TEXT)
sparkle(s, 6.6, 6.4, 0.6)
# recurring promo
s = slide(p); rect(s, 0, 0, 7.5, 7.5, PRIMARY); sparkle(s, 6.2, 0.5, 0.8); sparkle(s, 5.7, 1.4, 0.35, MINT)
text(s, 0.6, 0.6, 5.5, 0.4, "RECURRING CLEANING", 14, True, ACCENT)
text(s, 0.6, 1.1, 6.3, 2.6, ["Never spend", "Saturday", "cleaning again."], 44, True, WHITE)
rect(s, 0.6, 4.3, 6.3, 1.6, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.8, 4.4, 5.9, 1.4, ["Weekly · Every 2 weeks · Monthly", "[Your recurring benefit, e.g., same team, priority scheduling]"], 15, True, PRIMARY, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.25, 6.3, 0.5, "Book at [Website]  ·  [Phone]", 16, True, MINT, PP_ALIGN.CENTER)
text(s, 0.6, 6.8, 6.3, 0.4, "[Offer terms if any]", 9, False, MINT, PP_ALIGN.CENTER)
# new client promo
s = slide(p); rect(s, 0, 0, 7.5, 7.5, LIGHT)
rect(s, 0.45, 0.45, 6.6, 6.6, WHITE, line=ACCENT)
text(s, 0.8, 0.9, 5.9, 0.4, "NEW CLIENT OFFER", 14, True, ACCENT, PP_ALIGN.CENTER)
text(s, 0.8, 1.5, 5.9, 1.8, ["[Your offer]", "on your first clean"], 38, True, PRIMARY, PP_ALIGN.CENTER)
text(s, 0.8, 3.6, 5.9, 1.2, ["Standard · Deep · Move-in/out", "Serving [Service Area]"], 16, False, TEXT, PP_ALIGN.CENTER)
rect(s, 1.9, 5.0, 3.7, 0.9, ACCENT, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.9, 5.0, 3.7, 0.9, "Book: [Website]", 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.8, 6.15, 5.9, 0.5, "[Offer terms, expiry date, eligibility — your wording]", 9, False, RGBColor(0x6B, 0x7B, 0x83), PP_ALIGN.CENTER)
# tip post
s = slide(p); rect(s, 0, 0, 7.5, 7.5, WHITE); rect(s, 0, 0, 0.35, 7.5, ACCENT)
text(s, 0.8, 0.7, 6, 0.4, "CLEANING TIP", 14, True, ACCENT)
text(s, 0.8, 1.2, 6.2, 2.2, "[Hook: a specific problem, e.g., 'Why your shower glass still looks cloudy']", 28, True, PRIMARY)
text(s, 0.8, 3.7, 6.2, 2.4, ["1.  [Step one]", "2.  [Step two]", "3.  [Step three]"], 18, False, TEXT)
text(s, 0.8, 6.5, 6.2, 0.5, "Save this for later  ·  [YOUR BUSINESS NAME]", 12, True, PRIMARY)
# testimonial / review spotlight
s = slide(p); rect(s, 0, 0, 7.5, 7.5, PRIMARY)
text(s, 0.7, 0.6, 6, 1.2, "“", 96, True, ACCENT)
text(s, 0.8, 1.9, 5.9, 3.0, "[Paste a REAL review, word for word, with the customer's permission.]", 24, True, WHITE)
text(s, 0.8, 5.1, 5.9, 0.5, "★★★★★  — [First name], [Area]", 15, True, MINT)
text(s, 0.8, 6.5, 5.9, 0.4, "[YOUR BUSINESS NAME]  ·  [Website]", 11, False, MINT)
save(p, "05-Social-Squares-1080x1080.pptx")

# Story 9:16
p = deck(4.5, 8)
s = slide(p); rect(s, 0, 0, 4.5, 8, PRIMARY); sparkle(s, 3.6, 0.5, 0.5)
text(s, 0.4, 0.9, 3.7, 0.4, "THIS WEEK", 12, True, ACCENT)
text(s, 0.4, 1.3, 3.7, 1.8, ["[2] openings for", "recurring clients"], 26, True, WHITE)
photo_box(s, 0.4, 3.3, 3.7, 2.6, "[ PHOTO ]")
rect(s, 0.6, 6.3, 3.3, 0.8, ACCENT, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.6, 6.3, 3.3, 0.8, "DM \"CLEAN\" to book", 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
save(p, "06-Story-1080x1920.pptx")

# GBP image 4:3
p = deck(8, 6)
for title, sub in [("Now booking [month]", "Recurring and one-time cleans in [Service Area]"),
                   ("Moving soon?", "Move-in / move-out cleaning — book 1–2 weeks ahead"),
                   ("Deep clean season", "[Seasonal tip or service spotlight]")]:
    s = slide(p); rect(s, 0, 0, 8, 6, WHITE); rect(s, 0, 0, 3.4, 6, PRIMARY)
    sparkle(s, 2.5, 0.4, 0.6)
    text(s, 0.35, 4.6, 2.8, 1.0, "[YOUR BUSINESS NAME]", 13, True, MINT)
    photo_box(s, 3.7, 0.4, 3.9, 2.9, "[ YOUR PHOTO ]")
    text(s, 3.7, 3.5, 4.0, 1.0, title, 26, True, PRIMARY)
    text(s, 3.7, 4.5, 4.0, 1.0, sub, 14, False, TEXT)
save(p, "07-Google-Business-Post-Images-1200x900.pptx")

# copy document
d = Document()
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = DPt(10.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def H(t, lvl=1):
    h = d.add_heading(t, lvl)
    for r in h.runs:
        r.font.color.rgb = DRGB(0x0F, 0x4C, 0x5C); r.font.name = "Arial"


def P_(t, it=False):
    p = d.add_paragraph(); r = p.add_run(t); r.italic = it


H("Marketing Copy Templates", 0)
P_("Replace every [bracket]. Only make offers you will honor, only use real reviews (with permission), and follow each platform's rules. "
   "Track results by setting the Lead Source in your LEADS sheet.", True)
H("Google Business Profile posts")
gbp = [
    ("Service spotlight — recurring", "Busy week ahead? Our recurring cleaning plans keep your home reset every [week / two weeks] with the same team and the same checklist each visit. "
     "Now booking new recurring clients in [Service Area]. Tap 'Book' to request a quote."),
    ("Move-in / move-out", "Moving is stressful enough. Our move-in / move-out clean covers [inside cabinets, appliances, baseboards — edit to your scope] so you can hand over the keys or settle in with confidence. "
     "Book 1–2 weeks ahead for best availability in [Service Area]."),
    ("Seasonal tip", "Quick tip: [one practical, accurate cleaning tip]. Want us to handle it? We're booking [deep cleans / spring cleans] in [Service Area] for [month]."),
    ("Team highlight", "Meet [first name], our [role]. [One genuine detail, shared with permission.] When you book with [Business Name], you get [your consistency promise — only if true]."),
    ("Short-term rental hosts", "Hosting in [Service Area]? We offer turnover cleans with [linen change, restock checks, photo reports — edit]. Message us your calendar and we'll quote your turnovers."),
    ("Review thank-you", "Thank you to everyone who has shared a review this month — we read every one. [Optional: a real quote with permission.] Ready for a cleaner home? Request a quote today."),
]
for t, c in gbp:
    H(t, 2); P_(c)
H("Facebook / Instagram captions")
caps = [
    ("Before & after", "This [room] took [time] and a lot of [problem] 👀 Swipe to see the after. [One sentence on what was done.] Want yours done? Link in bio / message us. #[YourTown]Cleaning #DeepClean"),
    ("Recurring", "The best part of a recurring clean? You stop thinking about it. Same team, same checklist, every [week / two weeks]. [N] spots open this month in [Area]. Comment 'CLEAN' and we'll send details."),
    ("New client offer", "New here? [Your offer] on your first clean in [Service Area] — [terms, expiry]. Book at [link]."),
    ("Tip post", "Save this 👇 [Hook]. 1) [Step] 2) [Step] 3) [Step]. Questions? Ask below — we answer every one."),
    ("Referral", "Our favorite compliment is a referral. Know someone who needs a hand at home? Send them our way and mention your name. [Your thank-you, terms.]"),
    ("Behind the scenes", "What's in our cleaning caddy? [3 items + why]. (Always follow product labels!) Which one surprised you?"),
]
for t, c in caps:
    H(t, 2); P_(c)
H("Promotions")
H("Recurring-cleaning promotion", 2)
P_("Headline: Never spend Saturday cleaning again.\nBody: Weekly, every-2-week and monthly plans with [your benefit]. [Offer, if any — e.g., your recurring discount from PRICING].\nCall to action: Book a walkthrough at [link] / [phone].\nTerms: [Your wording.]")
H("New-client promotion", 2)
P_("Headline: [Your offer] on your first clean.\nBody: Standard, deep and move-in/out cleaning in [Service Area].\nCall to action: Book at [link].\nTerms: [New clients only, expiry date, one per household — your wording.]")
H("Print notes")
for t in ["Flyer: 8.5 x 11 in. Door hanger: 4.25 x 11 in (ask your printer for a die-cut hole template). Cards: 3.5 x 2 in.",
          "Add 0.125 in bleed if your printer requires it, and export to PDF at high quality.",
          "Replace [QR] boxes with a QR code linking to your booking page (many free generators exist). Test it before printing.",
          "Use only photos you took or have rights to, and get customer permission before showing their home."]:
    d.add_paragraph(t, style="List Bullet")
d.save(os.path.join(OUT, "MARKETING-COPY-TEMPLATES.docx"))
print("marketing built")
