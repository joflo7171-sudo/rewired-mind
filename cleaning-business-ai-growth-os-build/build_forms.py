"""Builds the 14 editable CLIENT-FORMS (.docx) — PDFs are made afterwards with LibreOffice."""
import sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
PRIMARY = RGBColor(0x0F, 0x4C, 0x5C)
ACCENT = RGBColor(0x2B, 0xB3, 0xA3)
MUTED = RGBColor(0x6B, 0x7B, 0x83)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def borders(table, color="C9D3D7"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def set_widths(t, widths):
    t.autofit = False
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


class Form:
    def __init__(self, title, subtitle, agreement=False):
        self.d = Document()
        sec = self.d.sections[0]
        sec.left_margin = sec.right_margin = Inches(0.75)
        sec.top_margin = Inches(0.6); sec.bottom_margin = Inches(0.6)
        st = self.d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.paragraph_format.space_after = Pt(4)
        # header band
        t = self.d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        shade(c0, "0F4C5C"); shade(c1, "0F4C5C")
        p = c0.paragraphs[0]; r = p.add_run("[YOUR BUSINESS NAME]"); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(255, 255, 255)
        p2 = c0.add_paragraph(); r2 = p2.add_run("[Phone]  ·  [Email]  ·  [Website]"); r2.font.size = Pt(8.5); r2.font.color.rgb = RGBColor(0xD8, 0xF0, 0xEC)
        p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("[LOGO]"); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xD8, 0xF0, 0xEC)
        c0.width = Inches(5.2); c1.width = Inches(1.8)
        self.d.add_paragraph()
        h = self.d.add_paragraph(); r = h.add_run(title.upper()); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = PRIMARY
        h.paragraph_format.space_after = Pt(0)
        s = self.d.add_paragraph(); r = s.add_run(subtitle); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = MUTED
        if agreement:
            self.notice("CUSTOMIZABLE TEMPLATE — Replace the bracketed text with your own policies. This is not legal advice. "
                        "Have it reviewed by a qualified professional where you operate before you rely on it.")

    def notice(self, text, fill="FDF1DC"):
        t = self.d.add_table(rows=1, cols=1); c = t.rows[0].cells[0]; shade(c, fill)
        r = c.paragraphs[0].add_run(text); r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = RGBColor(0x7A, 0x4E, 0x10)
        self.d.add_paragraph()

    def section(self, text):
        p = self.d.add_paragraph(); p.paragraph_format.space_before = Pt(8)
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = ACCENT
        pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "8"); bt.set(qn("w:color"), "2BB3A3"); bdr.append(bt); pPr.append(bdr)

    def fields(self, labels, cols=2):
        rows = (len(labels) + cols - 1) // cols
        t = self.d.add_table(rows=rows, cols=cols * 2); borders(t)
        for i, lab in enumerate(labels):
            r, c = divmod(i, cols)
            lc = t.rows[r].cells[c * 2]; vc = t.rows[r].cells[c * 2 + 1]
            shade(lc, "F3F7F8")
            run = lc.paragraphs[0].add_run(lab); run.font.size = Pt(8.5); run.bold = True; run.font.color.rgb = PRIMARY
            vc.paragraphs[0].add_run(" ")
            lc.width = Inches(1.4); vc.width = Inches(2.1 if cols == 2 else 5.6)
        set_widths(t, ([1.4, 2.1] * cols) if cols == 2 else [1.6, 5.4])
        for row in t.rows:
            row.height = Cm(0.75)
        self.d.add_paragraph()

    def grid(self, headers, nrows, widths=None, prefill=None):
        t = self.d.add_table(rows=nrows + 1, cols=len(headers)); borders(t)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; shade(c, "0F4C5C")
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(255, 255, 255)
        for ri in range(nrows):
            for ci in range(len(headers)):
                val = prefill[ri][ci] if prefill and ri < len(prefill) and ci < len(prefill[ri]) else ""
                run = t.rows[ri + 1].cells[ci].paragraphs[0].add_run(val); run.font.size = Pt(9)
            t.rows[ri + 1].height = Cm(0.7)
        if widths:
            set_widths(t, widths)
        self.d.add_paragraph()

    def checks(self, items, cols=2):
        rows = (len(items) + cols - 1) // cols
        t = self.d.add_table(rows=rows, cols=cols)
        for i, it in enumerate(items):
            r, c = divmod(i, cols)
            run = t.rows[r].cells[c].paragraphs[0].add_run("☐  " + it); run.font.size = Pt(9.5)
        self.d.add_paragraph()

    def para(self, text, size=10, bold=False, italic=False, color=None):
        p = self.d.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = color
        return p

    def bullets(self, items):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet"); r = p.add_run(it); r.font.size = Pt(9.5)

    def lines(self, label, n=3):
        self.para(label, 9, True, color=PRIMARY)
        for _ in range(n):
            self.para("_" * 98, 9, color=MUTED)

    def signatures(self, who=("Customer", "Business representative")):
        t = self.d.add_table(rows=2, cols=len(who))
        for i, w in enumerate(who):
            t.rows[0].cells[i].paragraphs[0].add_run("\n______________________________").font.size = Pt(10)
            r = t.rows[1].cells[i].paragraphs[0].add_run(f"{w} signature        Date"); r.font.size = Pt(8); r.font.color.rgb = MUTED

    def footer(self, text="Template — edit any wording to match your business."):
        p = self.d.sections[0].footer.paragraphs[0]
        r = p.add_run(text); r.font.size = Pt(7.5); r.font.color.rgb = MUTED

    def save(self, name):
        self.footer()
        self.d.save(os.path.join(OUT, name))


# 01 Client intake
f = Form("Client Intake Form", "Completed by phone, email or at the first walkthrough. Used to build the quote and the CUSTOMERS / PROPERTIES records.")
f.section("Contact")
f.fields(["Full name", "Preferred name", "Mobile", "Email", "Best way to reach you", "Best time to reach you", "How did you hear about us?", "Referred by (if any)"])
f.section("Property")
f.fields(["Service address", "City / area", "Property type", "Square feet (approx.)", "Bedrooms", "Bathrooms", "Floors / stairs", "Flooring types"])
f.section("Service requested")
f.checks(["Standard clean", "Deep clean", "Move-in / move-out", "Short-term rental turnover", "Office / commercial", "Other: ____________"])
f.para("Frequency:", 9, True, color=PRIMARY)
f.checks(["One-time", "Weekly", "Every 2 weeks", "Every 4 weeks", "Monthly", "Not sure yet"], cols=3)
f.para("Add-ons of interest:", 9, True, color=PRIMARY)
f.checks(["Inside oven", "Inside refrigerator", "Interior windows", "Inside cabinets", "Laundry / linens", "Baseboards detail"], cols=3)
f.section("Home details")
f.fields(["Pets (type / temperament)", "Areas to skip", "Product preferences / sensitivities", "Parking", "Entry method (no codes here)", "Preferred day / time"])
f.lines("Priorities — what matters most to you in a clean home?", 3)
f.para("Privacy: we use this information only to quote, schedule and deliver your service. [Add your own privacy statement.]", 8, italic=True, color=MUTED)
f.save("01-Client-Intake-Form.docx")

# 02 Walkthrough
f = Form("Property Walkthrough", "Room-by-room notes from an in-person or video walkthrough. Transfer totals to PROPERTIES and QUOTE BUILDER.")
f.fields(["Customer", "Date", "Address", "Walkthrough by", "Square feet", "Condition (Light / Average / Heavy / Very heavy)"])
f.grid(["Room / area", "Size / notes", "Condition", "Special tasks", "Est. minutes"], 14, [1.6, 1.6, 1.0, 1.9, 0.9],
       [["Kitchen"], ["Living room"], ["Dining"], ["Primary bedroom"], ["Bedroom 2"], ["Bedroom 3"], ["Primary bath"], ["Bath 2"], ["Laundry"], ["Hallways / stairs"], ["Office"], ["Entry"]])
f.fields(["Total est. labor hours", "Recommended crew size", "Add-ons agreed", "Photos taken? (with permission)"])
f.lines("Access, parking, pets and safety notes (store codes separately and securely)", 3)
f.save("02-Property-Walkthrough.docx")

# 03 Quote / estimate
f = Form("Cleaning Estimate", "Estimate prepared from the details you provided. Final price may change if the scope or condition is different on arrival — we'll always ask first.")
f.fields(["Estimate #", "Date", "Prepared for", "Valid until", "Service address", "Prepared by"])
f.grid(["Description", "Qty", "Unit price", "Amount"], 8, [3.9, 0.7, 1.1, 1.3],
       [["[Service type] — [sq ft], [bd]/[ba], [condition]", "1"], ["Add-on: [name]"], ["Add-on: [name]"], ["Recurring discount ([frequency])"], ["Discount"]])
f.grid(["", "Amount"], 3, [5.7, 1.3], [["Subtotal (before tax)"], ["Sales tax (only if applicable) [__%]"], ["ESTIMATED TOTAL"]])
f.section("What's included")
f.bullets(["[List the tasks from your checklist for this service]", "[e.g., kitchen surfaces, bathrooms, dusting, floors]", "[Add-ons listed above]"])
f.section("Not included unless listed")
f.bullets(["[e.g., exterior windows, heavy mold, biohazards, moving furniture over a set weight]"])
f.para("To book: [reply to this email / call / booking link]. Payment: [when and how you accept payment].", 9.5)
f.save("03-Quote-Estimate.docx")

# 04 Proposal
f = Form("Cleaning Services Proposal", "For recurring residential or commercial clients. Pair with AI Workflow 15 to draft the scope wording.", agreement=True)
f.fields(["Prepared for", "Contact", "Address", "Date", "Proposal #", "Valid until"])
f.section("1. Summary")
f.lines("[Two or three sentences: the client's goal and how your service meets it.]", 3)
f.section("2. Scope of work")
f.grid(["Area", "Tasks", "Frequency"], 7, [1.5, 4.3, 1.2], [["Entry / reception"], ["Work areas"], ["Restrooms"], ["Break room / kitchen"], ["Floors"], ["Trash & recycling"]])
f.section("3. Schedule")
f.fields(["Visit days", "Time window", "Start date", "Key / access arrangement"])
f.section("4. Investment")
f.grid(["Item", "Per visit", "Visits / month", "Monthly estimate"], 3, [3.1, 1.3, 1.3, 1.3])
f.section("5. Supplies, quality checks and communication")
f.bullets(["Supplies & equipment: [who provides what]", "Quality checks: [e.g., monthly inspection using our checklist]", "Point of contact: [name, phone, email]"])
f.section("6. Terms to complete")
f.bullets(["[Payment terms — your wording]", "[Notice period to change or end service — your wording]", "[Insurance / liability details — complete with your professional advisors]"])
f.signatures(("Client", "Provider"))
f.save("04-Proposal.docx")

# 05 Welcome guide
f = Form("Welcome Guide", "Send after the first booking. Edit every section so it matches how you actually work.")
f.para("Welcome, [first name]! Thank you for choosing [Business Name]. Here's everything you need to know before your first clean.", 11)
f.section("Before your first visit")
f.bullets(["Pick up clutter so we can clean surfaces, not tidy them.", "Secure pets or let us know how they do with visitors.",
           "Tell us about any areas to skip or products to avoid.", "Confirm how we'll get in (we keep entry details private)."])
f.section("What to expect")
f.bullets(["Arrival window: [e.g., 30-minute window]. We'll text when we're on the way.", "Your team: [names] — the same team whenever possible.",
           "We follow a room-by-room checklist and do a final walk-through.", "Supplies: [we bring everything / let us know if you prefer your products]."])
f.section("After each visit")
f.bullets(["Something missed? Tell us within [24 hours] and we'll make it right [your policy].", "Payment: [method and timing].",
           "Recurring clients keep their preferred slot — let us know at least [__] hours ahead to change it."])
f.section("Contact us")
f.fields(["Phone / text", "Email", "Office hours", "Booking link"])
f.save("05-Welcome-Guide.docx")

# 06 Job checklist
f = Form("Job Checklist", "Bring on every job. Mark each item as you finish. Matches the QC CHECKLISTS tab.")
f.fields(["Job ID", "Date", "Customer", "Service", "Cleaner(s)", "Time in / out"])
for area, items in [("Kitchen", ["Counters & backsplash", "Sink & faucet", "Appliance fronts", "Stovetop", "Microwave", "Cabinet fronts (spot)", "Floor", "Trash out & new liner"]),
                    ("Bathrooms", ["Toilet incl. base", "Shower / tub", "Sink, counter, faucet", "Mirrors", "Floor", "Towels straightened"]),
                    ("Living & bedrooms", ["Dust reachable surfaces", "Switches & handles", "Floors", "Beds made / linens", "Glass & mirrors"]),
                    ("Wrap-up", ["Special requests done", "Nothing left behind", "Final walk-through", "Locked / secured as instructed"])]:
    f.section(area); f.checks(items)
f.lines("Notes for the office (issues, damage, supplies low)", 2)
f.save("06-Job-Checklist.docx")

# 07 Invoice
f = Form("Invoice", "Match the Invoice # to the Ref column in REVENUE and the Job ID in JOBS.")
f.fields(["Invoice #", "Invoice date", "Bill to", "Due date", "Service address", "Job ID"])
f.grid(["Date", "Description", "Qty", "Rate", "Amount"], 7, [0.9, 3.4, 0.6, 1.0, 1.1])
f.grid(["", "Amount"], 5, [5.9, 1.1], [["Subtotal"], ["Discount"], ["Sales tax (only if applicable) [__%]"], ["Payments received"], ["BALANCE DUE"]])
f.section("How to pay")
f.bullets(["[Payment methods you accept]", "[Payment terms in your own words]"])
f.para("Thank you for your business!", 10, True, color=PRIMARY)
f.save("07-Invoice.docx")

# 08 Receipt
f = Form("Payment Receipt", "Give or send when payment is received. Log the same payment in REVENUE.")
f.fields(["Receipt #", "Date received", "Received from", "Invoice # / Job ID", "Payment method", "Reference #"])
f.grid(["Description", "Amount"], 4, [5.9, 1.1], [["Cleaning service — [date]"], ["Tip (optional)"], ["TOTAL RECEIVED"]])
f.fields(["Remaining balance", "Received by"])
f.para("Thank you! Questions about this receipt: [phone / email].", 9.5)
f.save("08-Receipt.docx")

# 09 Review request
f = Form("Review Request", "Printable leave-behind card text plus message versions. Never offer rewards for reviews; ask for honest feedback.")
f.section("Leave-behind card (print and leave on the counter)")
f.para("Your home was cleaned today by [cleaner first name].", 11, True)
f.para("If you were happy with today's clean, a short review helps our small team more than you know. If anything wasn't right, please tell us first — we'll make it right.", 10)
f.fields(["Review link / QR code", "Questions or concerns"], cols=1)
f.section("Text message")
f.para("Hi [first name], thanks for having us today! If you have a minute, we'd really appreciate an honest review: [link]. — [Your name]", 10, italic=True)
f.section("Email")
f.para("Subject: How did we do today?\n\nHi [first name],\n\nThank you for trusting [Business Name] with your home. We'd love to hear how today's clean went. "
       "If you have two minutes, you can leave a review here: [link].\n\nIf anything wasn't perfect, just reply to this email and we'll make it right.\n\n[Your name]", 10)
f.save("09-Review-Request.docx")

# 10 Referral request
f = Form("Referral Request", "Use after a compliment or 5-star review. Track every referral in REFERRALS.")
f.para("Know someone who'd love a cleaner home?", 13, True, color=PRIMARY)
f.para("Our business grows through happy clients like you. If a friend, neighbor or coworker could use help, share our details below.", 10)
f.fields(["Your name", "Friend's name", "Friend's phone / email", "Best time to contact them", "Service they might need", "Okay to mention your name?"])
f.section("Our thank-you (optional — edit or delete)")
f.para("[Describe your thank-you in your own words, e.g., a service credit after your friend's first completed clean. Set your own terms.]", 9.5, italic=True)
f.section("Share-ready message")
f.para("\"I use [Business Name] for cleaning and they've been great — reliable and detailed. Here's their number/link: [link]. Tell them [your name] sent you.\"", 10, italic=True)
f.save("10-Referral-Request.docx")

# 11 Policy template
f = Form("Rescheduling & Cancellation Policy", "Customizable policy template. Choose the options that fit your business and delete the rest.", agreement=True)
f.section("Rescheduling")
f.bullets(["Please give at least [__ hours] notice to reschedule. We'll offer the next available opening.",
           "If we need to reschedule (illness, weather, emergencies), we'll contact you as early as possible and offer [priority rebooking / other]."])
f.section("Cancellations")
f.bullets(["Cancellations with [__ hours] or more notice: [no charge].", "Cancellations with less than [__ hours] notice: [your policy — e.g., a fee of $__ or __% of the visit].",
           "Lockouts (we can't access the property at the scheduled time): [your policy].", "Recurring clients who skip [__] visits in a row may lose their reserved time slot [your policy]."])
f.section("Weather and safety")
f.bullets(["We may reschedule when travel or working conditions are unsafe. [Your wording.]"])
f.section("Changes to this policy")
f.para("[How and when you will notify customers of changes.]", 9.5, italic=True)
f.para("Acknowledged by:", 9.5, True, color=PRIMARY)
f.signatures(("Customer",))
f.save("11-Rescheduling-Cancellation-Policy-TEMPLATE.docx")

# 12 Complaint resolution
f = Form("Complaint Resolution Form", "Internal form. Log the same issue in QUALITY CONTROL → Issue & Rework Log. Use AI Workflow 17 for the reply.")
f.fields(["Date reported", "Job ID", "Customer", "Contact method", "Job date", "Cleaner(s)", "Reported by", "Issue type"])
f.lines("Customer's description (their words)", 3)
f.lines("What we found (facts only)", 3)
f.section("Resolution")
f.checks(["Apology / acknowledgement sent", "Re-clean scheduled: ________", "Coaching with cleaner", "Credit approved by owner: ________",
          "Escalated to owner (damage, injury, claim or legal matter)", "Customer confirmed resolved"])
f.fields(["Action owner", "Target date", "Cost of fix", "Resolved date"])
f.notice("Damage, injury, refund demands tied to claims, or legal questions: do not admit fault or promise payment in writing. The owner handles these, with professional advice where needed.")
f.save("12-Complaint-Resolution-Form.docx")

# 13 Quality inspection
f = Form("Quality Inspection Form", "Score = items passed ÷ items checked. Enter both numbers in the QUALITY CONTROL log.")
f.fields(["Job ID", "Date", "Customer", "Cleaner(s)", "Inspector", "Checklist used"])
f.grid(["Area", "Item checked", "Pass", "Fail", "N/A", "Notes"], 18, [1.1, 2.6, 0.5, 0.5, 0.5, 1.8],
       [["Kitchen", "Counters"], ["Kitchen", "Sink"], ["Kitchen", "Floor"], ["Bath", "Toilet"], ["Bath", "Shower / tub"], ["Bath", "Mirror"],
        ["Bedrooms", "Dusting"], ["Bedrooms", "Floors"], ["Living", "Surfaces"], ["Living", "Glass"], ["All", "Trash"], ["All", "Special requests"]])
f.fields(["Items checked", "Items passed", "Score %", "Pass / Fail", "Rework needed?", "Rework date"])
f.save("13-Quality-Inspection-Form.docx")

# 14 Employee job sheet
f = Form("Employee Job Sheet", "Give one per job. Keep entry codes off paper — share them securely, separately.")
f.fields(["Job ID", "Date", "Arrival time", "Est. hours", "Customer (first name)", "Service", "Street / area", "Crew"])
f.section("Customer preferences")
f.checks(["Pets on site: ________", "Unscented products", "Shoes off / shoe covers", "Skip areas: ________", "Product to avoid: ________", "Special request: ________"])
f.section("Priorities if time runs short")
f.grid(["#", "Priority task"], 3, [0.4, 6.6], [["1"], ["2"], ["3"]])
f.section("Before you leave")
f.checks(["Checklist complete", "Final walk-through", "Photos (if required)", "Nothing left behind", "Secured as instructed", "Issues reported to office"])
f.fields(["Actual time in", "Actual time out", "Supplies running low", "Notes for office"])
f.save("14-Employee-Job-Sheet.docx")
print("forms built")
