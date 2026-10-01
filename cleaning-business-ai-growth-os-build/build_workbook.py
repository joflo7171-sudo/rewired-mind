"""Builds CLEANING BUSINESS AI GROWTH OS workbooks (DEMO and CLEAN).

Usage: python build_workbook.py <out_dir>
"""
import sys, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

import demo_data

# ---------------------------------------------------------------- style
FONT = "Arial"
C_PRIMARY = "0F4C5C"   # deep teal
C_ACCENT = "2BB3A3"    # mint
C_AUTOHDR = "5B8A95"   # muted teal for automatic columns
C_AUTOFILL = "EEF3F5"
C_LIGHT = "F5F8F9"
C_INPUT = "FFFDF2"     # warm off-white for key inputs
C_TEXT = "1E2A30"
C_MUTED = "6B7B83"
C_RED = "C0392B"; C_REDBG = "FBE3E0"
C_AMBER = "B7791F"; C_AMBERBG = "FDF1DC"
C_GREEN = "1E7F4F"; C_GREENBG = "E2F4EA"
C_BLUEBG = "E4EEF7"; C_BLUE = "2A5D8F"

thin = Side(style="thin", color="D5DEE2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(size=10, bold=False, color=C_TEXT, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)

def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)

FMT_MONEY = '$#,##0.00;[Red]-$#,##0.00;"–"'
FMT_MONEY0 = '$#,##0;[Red]-$#,##0;"–"'
FMT_PCT = '0.0%;[Red]-0.0%;"–"'
FMT_DATE = 'mmm d, yyyy'
FMT_TIME = 'h:mm AM/PM'
FMT_NUM1 = '0.0;-0.0;"–"'
FMT_NUM2 = '0.00;-0.00;"–"'
FMT_INT = '#,##0;-#,##0;"–"'

FIRST = 5  # first data row on log sheets

# tab groups
TAB = {"start": "0F4C5C", "sales": "2A5D8F", "ops": "2BB3A3", "money": "1E7F4F",
       "growth": "E08E2B", "admin": "7F8C8D"}

# ---------------------------------------------------------------- registries
LOGS = {}  # sheet -> LogSpec


class Col:
    def __init__(self, key, header, width=14, kind="input", fmt=None, formula=None,
                 dv=None, note=None, align=None):
        self.key, self.header, self.width, self.kind = key, header, width, kind
        self.fmt, self.formula, self.dv, self.note, self.align = fmt, formula, dv, note, align


class LogSpec:
    def __init__(self, sheet, cols, n, first=FIRST, id_prefix=None, id_width=4, header_row=None):
        n = max(5, int(n * float(os.environ.get("CAPSCALE", "1"))))
        self.sheet, self.cols, self.n, self.first = sheet, cols, n, first
        self.header_row = header_row or first - 1
        self.last = first + n - 1
        self.letters = {c.key: L(i + 1) for i, c in enumerate(cols)}
        self.id_prefix, self.id_width = id_prefix, id_width

    def rng(self, key):
        c = self.letters[key]
        return f"'{self.sheet}'!${c}${self.first}:${c}${self.last}"


def R(sheet, key):
    return LOGS[sheet].rng(key)


def q(sheet):
    return f"'{sheet}'!"

# ---------------------------------------------------------------- settings layout
SET = "SETTINGS"
SETTINGS_ROWS = [
    # (name, label, default_demo, default_clean, fmt, note)
    ("section", "BUSINESS PROFILE"),
    ("BizName", "Business name", "Sparkle & Co. Cleaning (Demo)", "[Your Business Name]", None, "Shown on the dashboard. Fictional in the demo."),
    ("OwnerName", "Owner / manager", "Jordan Rivera (fictional)", "", None, ""),
    ("BizPhone", "Business phone", "(555) 010-0100", "", None, "555-01xx numbers are reserved for fiction."),
    ("BizEmail", "Business email", "hello@example.com", "", None, ""),
    ("BizWeb", "Website / booking link", "www.example.com", "", None, "Used in your AI prompts and marketing kit."),
    ("BizArea", "Service area", "Sampletown & nearby", "", None, ""),
    ("section", "DATES"),
    ("AsOfInput", "As-of date (leave blank = today)", dt.date(2026, 9, 30), None, FMT_DATE, "Leave BLANK in normal use. The demo fixes it so the sample data always looks current."),
    ("AsOf", "Date the system uses (automatic)", "=IF(B{AsOfInput}=\"\",TODAY(),B{AsOfInput})", None, FMT_DATE, "Automatic. Do not type here."),
    ("RMInput", "Reporting month (blank = month of the date above)", dt.date(2026, 9, 1), None, 'mmmm yyyy', "Type any date in the month you want to review."),
    ("ReportMonth", "Reporting month used (automatic)", "=IF(B{RMInput}=\"\",DATE(YEAR(AsOf),MONTH(AsOf),1),DATE(YEAR(B{RMInput}),MONTH(B{RMInput}),1))", None, 'mmmm yyyy', "Automatic. Drives the dashboard and profitability views."),
    ("ReportYear", "Reporting year (automatic)", "=YEAR(ReportMonth)", None, '0', "Automatic. Drives MONTHLY."),
    ("section", "LABOR & COST ASSUMPTIONS  (your estimates — edit freely)"),
    ("PayRate", "Default worker pay rate ($ per hour)", 20, 20, FMT_MONEY, "Used when a worker has no rate on STAFF & TASKS."),
    ("Burden", "Extra labor cost % (your estimate)", 0.12, 0.12, FMT_PCT, "Your own estimate of costs on top of wages (for example payroll costs or insurance). Confirm with your payroll provider or accountant."),
    ("SupHr", "Supplies cost per labor hour ($)", 2.5, 2.5, FMT_MONEY, "Average chemicals/consumables used per labor hour. Adjusted per service type on PRICING."),
    ("MileRate", "Vehicle cost per mile ($)", 0.7, 0.7, FMT_MONEY, "Your own estimate of fuel + wear. Some owners use an official standard mileage rate; check the current rate for your country."),
    ("TravelMin", "Default paid travel time per worker (minutes)", 20, 20, '0', "Pre-fills the QUOTE BUILDER."),
    ("FeePct", "Card / online payment fee %", 0.029, 0.029, FMT_PCT, "Applied to payments whose method is marked 'Yes' in the payment-method list."),
    ("TaxDefault", "Default sales tax % (only if it applies to you)", 0, 0, FMT_PCT, "Leave 0 unless you have confirmed sales tax applies to your services where you work."),
    ("OverheadHr", "Overhead share per labor hour ($, optional)", 5, 0, FMT_MONEY, "Used only by QUOTE BUILDER to estimate profit after overhead. Tip: monthly overhead ÷ monthly labor hours (both on MONTHLY)."),
    ("section", "TARGETS & RULES"),
    ("TargetMargin", "Target gross margin %", 0.5, 0.5, FMT_PCT, "Jobs and quotes below this are flagged. Your choice, not an industry standard."),
    ("MinJob", "Minimum job price ($)", 120, 120, FMT_MONEY, "Quotes never go below this base price."),
    ("SoonDays", "Follow-up 'due soon' window (days)", 2, 2, '0', ""),
    ("ReactDays", "Suggest reactivation after (days since last job)", 90, 90, '0', ""),
    ("QCPass", "Quality-control passing score", 0.9, 0.9, FMT_PCT, "Inspections scoring below this are marked Fail."),
    ("ReviewWait", "Flag review requests with no reply after (days)", 7, 7, '0', ""),
]

# dropdown lists on SETTINGS: (key, header, values, locked)
LISTS = [
    ("sources", "Lead Sources", ["Google Search", "Google Business Profile", "Facebook", "Instagram", "Nextdoor",
                                 "Referral", "Website Form", "Repeat Customer", "Flyer / Door Hanger",
                                 "Yard Sign / Vehicle", "Other"], False),
    ("stages", "Pipeline Stages", ["New Lead", "Contacted", "Quote Requested", "Quote Sent", "Follow-Up",
                                   "Booked", "Recurring Client", "Lost/Declined"], True),
    ("lost", "Lost Reasons", ["Price", "Chose another company", "No response", "Timing / availability",
                              "Outside service area", "Not a fit", "Other"], False),
    ("jobstatus", "Job Statuses", ["Scheduled", "Completed", "Cancelled", "Rescheduled", "No-Show"], True),
    ("paymeth", "Payment Methods", ["Cash", "Check", "Card", "Online Invoice", "Bank Transfer", "Payment App", "Other"], False),
    ("payfee", "Fee Applies?", ["No", "No", "Yes", "Yes", "No", "No", "No"], False),
    ("expcat", "Expense Categories", ["Advertising & Marketing", "Software & Apps", "Insurance", "Phone & Internet",
                                      "Equipment", "Uniforms", "Office & Admin", "Professional Fees",
                                      "Licenses & Permits", "Training", "Bank Fees", "Supplies (job use)",
                                      "Fuel & Vehicle", "Staff Pay", "Merchant Fees", "Other"], False),
    ("expoh", "Counted as Overhead?", ["Yes", "Yes", "Yes", "Yes", "Yes", "Yes", "Yes", "Yes", "Yes", "Yes", "Yes",
                                       "No", "No", "No", "No", "Yes"], False),
    ("ctype", "Customer Types", ["Residential", "Commercial", "Short-Term Rental Host", "Property Manager", "Other"], False),
    ("cstatus", "Customer Statuses", ["Active", "Inactive", "Do not service"], True),
    ("contact", "Contact Preferences", ["Text", "Call", "Email"], False),
    ("ptype", "Property Types", ["House", "Apartment", "Condo", "Townhome", "Office", "Retail", "Short-Term Rental",
                                 "Medical/Dental", "Other"], False),
    ("planstatus", "Plan Statuses", ["Active", "Paused", "Ended"], True),
    ("staffstatus", "Staff Statuses", ["Active", "Inactive"], True),
    ("roles", "Staff Roles", ["Owner", "Team Lead", "Cleaner", "Inspector", "Office"], False),
    ("priority", "Priorities", ["High", "Medium", "Low"], True),
    ("taskstatus", "Task Statuses", ["Not started", "In progress", "Done"], True),
    ("supcat", "Supply Categories", ["Chemicals", "Tools", "Paper & Liners", "Laundry", "Protective Gear",
                                     "Equipment", "Other"], False),
    ("futype", "Follow-Up Types", ["Lead", "Quote", "Invoice", "Review", "Referral", "Reactivation",
                                   "Complaint", "Other"], False),
    ("fustatus", "Follow-Up Statuses", ["Open", "Done", "Skipped"], True),
    ("channel", "Channels", ["Text", "Call", "Email", "In person", "Social DM"], False),
    ("platform", "Review Platforms", ["Google", "Facebook", "Yelp", "Nextdoor", "Other"], False),
    ("revstatus", "Review Statuses", ["Requested", "Reminder sent", "Received", "Declined", "No response"], True),
    ("refstatus", "Referral Statuses", ["Referred", "Contacted", "Quoted", "Booked", "Not booked"], True),
    ("rewstatus", "Reward Statuses", ["Not due", "Owed", "Given"], True),
    ("checklists", "QC Checklists", ["Residential Standard", "Deep Clean", "Move-In/Move-Out", "Airbnb Turnover",
                                     "Commercial", "Final Inspection"], False),
    ("reporter", "Reported By", ["Customer", "Staff", "Inspection"], False),
    ("issuetype", "Issue Types", ["Missed area", "Quality complaint", "Damage report", "Late / no-show",
                                  "Access problem", "Billing question", "Other"], False),
    ("issuestatus", "Issue Statuses", ["Open", "In progress", "Resolved"], True),
    ("yesno", "Yes / No", ["Yes", "No"], True),
    ("period", "Periods", ["Reporting month", "Year to date", "All time"], True),
    ("workflows", "AI Workflows", ["01 New lead reply", "02 Quote draft", "03 Quote follow-up", "04 No-response follow-up",
                                   "05 Appointment confirmation", "06 Rescheduling message", "07 Review request",
                                   "08 Review response", "09 Referral request", "10 Reactivation",
                                   "11 Google Business post", "12 Facebook post", "13 Instagram caption",
                                   "14 FAQ answer", "15 Commercial proposal", "16 Employee task sheet",
                                   "17 Complaint response", "18 Weekly marketing ideas"], False),
]
LIST_FIRST, LIST_SLOTS = 5, 20
LISTCOL = {}  # key -> column letter on SETTINGS (lists start at column F)
for i, (k, *_rest) in enumerate(LISTS):
    LISTCOL[k] = L(6 + i)


def LST(key, exact=False):
    c = LISTCOL[key]
    vals = [x for x in LISTS if x[0] == key][0][2]
    last = LIST_FIRST + (len(vals) if exact else LIST_SLOTS) - 1
    return f"'{SET}'!${c}${LIST_FIRST}:${c}${last}"

# ---------------------------------------------------------------- pricing layout
PR = "PRICING"
PT_FIRST, PT_LAST = 7, 16
COND_FIRST, COND_LAST = 20, 24
FREQ_FIRST, FREQ_LAST = 28, 33
ADD_FIRST, ADD_LAST = 37, 51

def prng(col, a, b):
    return f"'{PR}'!${col}${a}:${col}${b}"

PT_TYPE, PT_RATE, PT_SQFT, PT_BED, PT_BATH, PT_SETUP, PT_SUP = [prng(c, PT_FIRST, PT_LAST) for c in "ABCDEFG"]
COND_NAME, COND_MULT = prng("A", COND_FIRST, COND_LAST), prng("B", COND_FIRST, COND_LAST)
FREQ_NAME, FREQ_DISC, FREQ_DAYS, FREQ_VPM = [prng(c, FREQ_FIRST, FREQ_LAST) for c in "ABCD"]
FREQ_RECUR = prng("A", FREQ_FIRST + 1, FREQ_LAST)
ADD_NAME, ADD_PRICE, ADD_MIN, ADD_COST = [prng(c, ADD_FIRST, ADD_LAST) for c in "ABCD"]

PRICING_TYPES = [
    # type, rate, sqft/hr, min/bed, min/bath, setup, supplies mult, note
    ("Standard Clean", 50, 800, 5, 15, 10, 1.0, "Routine maintenance clean"),
    ("Deep Clean", 55, 400, 10, 30, 15, 1.5, "Detail work: baseboards, buildup, fixtures"),
    ("Move-In/Move-Out", 55, 350, 15, 40, 20, 1.6, "Empty home, inside cabinets & appliances as add-ons"),
    ("Airbnb Turnover", 50, 900, 10, 20, 15, 1.2, "Reset + staging; laundry/restock as add-ons"),
    ("Commercial Clean", 45, 2500, 0, 20, 15, 0.8, "Use 'bathrooms' for restrooms; bedrooms = 0"),
]
CONDITIONS = [("Light", 0.9), ("Average", 1.0), ("Heavy", 1.25), ("Very heavy", 1.5)]
FREQS = [("One-time", 0, 0, 0), ("Weekly", 0.15, 7, 4.33), ("Biweekly", 0.10, 14, 2.17),
         ("Every 3 weeks", 0.07, 21, 1.44), ("Every 4 weeks", 0.05, 28, 1.08), ("Monthly", 0.05, 30, 1)]
ADDONS = [("Inside oven", 35, 30, 1.0), ("Inside refrigerator", 35, 30, 1.0), ("Interior windows (per 10)", 40, 45, 1.0),
          ("Baseboards detail", 30, 30, 0.5), ("Inside cabinets", 40, 45, 1.0), ("Laundry (per load)", 15, 10, 0.75),
          ("Bed linen change (per bed)", 10, 10, 0), ("Balcony / patio", 25, 20, 0.5), ("Garage sweep", 30, 25, 0.5),
          ("Wall spot cleaning", 25, 20, 0.5), ("Restock supplies (STR)", 15, 10, 0)]


def model_hours(t, sq, bd, ba, cond=None):
    m = f"MATCH({t},{PT_TYPE},0)"
    e = (f"({sq}/INDEX({PT_SQFT},{m})+{bd}*INDEX({PT_BED},{m})/60+{ba}*INDEX({PT_BATH},{m})/60"
         f"+INDEX({PT_SETUP},{m})/60)")
    if cond:
        e += f"*IFERROR(INDEX({COND_MULT},MATCH({cond},{COND_NAME},0)),1)"
    return e

# ---------------------------------------------------------------- sheet order
SHEETS = [
    ("START HERE", "start"), ("DASHBOARD", "start"), ("LEADS", "sales"), ("CUSTOMERS", "sales"),
    ("PROPERTIES", "sales"), ("QUOTE BUILDER", "sales"), ("PRICING", "sales"), ("JOBS", "ops"),
    ("RECURRING", "ops"), ("SCHEDULE", "ops"), ("STAFF & TASKS", "ops"), ("REVENUE", "money"),
    ("EXPENSES", "money"), ("PROFITABILITY", "money"), ("MILEAGE", "money"), ("SUPPLIES", "ops"),
    ("FOLLOW-UPS", "growth"), ("REVIEWS", "growth"), ("REFERRALS", "growth"), ("QUALITY CONTROL", "ops"),
    ("QC CHECKLISTS", "ops"), ("MONTHLY", "money"), ("SETTINGS", "admin"),
]

# ---------------------------------------------------------------- log specs
def blank_guard(c, key, r, expr):
    return f'=IF({c[key]}{r}="","",{expr})'


def define_logs():
    J, CU, PRP, LE, RV, EX, MI, QC, IS, RC, RF, RW, ST, TK, SU, FU = (
        "JOBS", "CUSTOMERS", "PROPERTIES", "LEADS", "REVENUE", "EXPENSES", "MILEAGE", "QUALITY CONTROL",
        "QUALITY CONTROL", "RECURRING", "REFERRALS", "REVIEWS", "STAFF & TASKS", "STAFF & TASKS",
        "SUPPLIES", "FOLLOW-UPS")

    # LEADS ------------------------------------------------------------
    LOGS["LEADS"] = LogSpec("LEADS", [
        Col("id", "Lead ID", 9, "id"),
        Col("date", "Date Added", 12, fmt=FMT_DATE),
        Col("name", "Name", 20),
        Col("phone", "Phone", 15),
        Col("email", "Email", 24),
        Col("source", "Lead Source", 18, dv=LST("sources")),
        Col("service", "Requested Service", 18, dv=PT_TYPE),
        Col("area", "Area / Property Notes", 22),
        Col("quote", "Quote Value ($)", 12, fmt=FMT_MONEY),
        Col("stage", "Stage", 16, dv=LST("stages", True)),
        Col("last", "Last Contact", 12, fmt=FMT_DATE),
        Col("next", "Next Follow-Up", 13, fmt=FMT_DATE),
        Col("lost", "Lost Reason", 17, dv=LST("lost")),
        Col("notes", "Notes", 28),
        Col("outcome", "Booking Outcome", 12, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF(OR({c["stage"]}{r}="Booked",{c["stage"]}{r}="Recurring Client"),"Booked",IF({c["stage"]}{r}="Lost/Declined","Lost","Open"))')),
        Col("fu", "Follow-Up Status", 13, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF({c["open"]}{r}=0,"—",IF({c["next"]}{r}="","Set date",IF({c["next"]}{r}<AsOf,"Overdue",IF({c["next"]}{r}=AsOf,"Due today",IF({c["next"]}{r}<=AsOf+SoonDays,"Due soon","Scheduled")))))')),
        Col("days", "Days Since Added", 10, "auto", fmt='0', formula=lambda r, c: blank_guard(c, "name", r,
            f'IF({c["date"]}{r}="","",AsOf-{c["date"]}{r})')),
        Col("check", "Customer Record", 16, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF({c["outcome"]}{r}<>"Booked","",IF(COUNTIF({R("CUSTOMERS","name")},{c["name"]}{r})>0,"✓ In CUSTOMERS","Add to CUSTOMERS"))')),
        Col("open", "Open? (1/0)", 8, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF(OR({c["stage"]}{r}="Booked",{c["stage"]}{r}="Recurring Client",{c["stage"]}{r}="Lost/Declined"),0,1)')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
        Col("rank", "Follow-Up Order", 9, "auto", formula=lambda r, c:
            f'=IF(OR({c["open"]}{r}<>1,{c["next"]}{r}=""),"",COUNTIFS({R("LEADS","open")},1,{R("LEADS","next")},"<"&{c["next"]}{r})'
            f'+COUNTIFS(${c["open"]}${FIRST}:{c["open"]}{r},1,${c["next"]}${FIRST}:{c["next"]}{r},{c["next"]}{r}))'),
    ], 500, id_prefix="L-", id_width=4)

    # CUSTOMERS --------------------------------------------------------
    LOGS["CUSTOMERS"] = LogSpec("CUSTOMERS", [
        Col("id", "Customer ID", 10, "id"),
        Col("name", "Customer Name (unique)", 22),
        Col("phone", "Phone", 15),
        Col("email", "Email", 24),
        Col("ctype", "Customer Type", 18, dv=LST("ctype")),
        Col("source", "Lead Source", 18, dv=LST("sources")),
        Col("contact", "Preferred Contact", 11, dv=LST("contact")),
        Col("status", "Status", 12, dv=LST("cstatus", True)),
        Col("added", "Date Added", 12, fmt=FMT_DATE),
        Col("notes", "Notes", 28),
        Col("first", "First Job", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "name", r,
            f'IF(_xlfn.MINIFS({R("JOBS","date")},{R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed")=0,"",_xlfn.MINIFS({R("JOBS","date")},{R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed"))')),
        Col("lastjob", "Last Job", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "name", r,
            f'IF(_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed")=0,"",_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed"))')),
        Col("jobs", "Completed Jobs", 10, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "name", r,
            f'COUNTIFS({R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed")')),
        Col("ltv", "Lifetime Revenue", 13, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "name", r,
            f'SUMIFS({R("JOBS","price")},{R("JOBS","customer")},{c["name"]}{r},{R("JOBS","status")},"Completed")')),
        Col("bal", "Balance Due", 12, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "name", r,
            f'SUMIFS({R("JOBS","balance")},{R("JOBS","customer")},{c["name"]}{r})')),
        Col("recur", "Recurring Plan?", 10, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF(COUNTIFS({R("RECURRING","customer")},{c["name"]}{r},{R("RECURRING","status")},"Active")>0,"Yes","No")')),
        Col("since", "Days Since Last Job", 10, "auto", fmt='0', formula=lambda r, c: blank_guard(c, "lastjob", r,
            f'AsOf-{c["lastjob"]}{r}')),
        Col("action", "Suggested Next Action", 18, "auto", formula=lambda r, c: blank_guard(c, "name", r,
            f'IF({c["status"]}{r}="Do not service","—",IF({c["bal"]}{r}>0.005,"Collect balance",IF({c["jobs"]}{r}=0,"Book first job",'
            f'IF(AND({c["recur"]}{r}="No",{c["since"]}{r}>=ReactDays),"Reactivate",IF(COUNTIF({R("REVIEWS","customer")},{c["name"]}{r})=0,"Ask for review",'
            f'IF(AND({c["recur"]}{r}="No",{c["jobs"]}{r}>=1),"Offer recurring plan","—"))))))')),
        Col("refs", "Referrals Given", 9, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "name", r,
            f'COUNTIF({R("REFERRALS","by")},{c["name"]}{r})')),
    ], 400, id_prefix="C-", id_width=3)

    # PROPERTIES ------------------------------------------------------
    LOGS["PROPERTIES"] = LogSpec("PROPERTIES", [
        Col("id", "Property ID", 10, "id"),
        Col("addr", "Address / Label (unique)", 26),
        Col("customer", "Customer", 20, dv=lambda: R("CUSTOMERS", "name")),
        Col("area", "City / Area", 14),
        Col("ptype", "Property Type", 15, dv=LST("ptype")),
        Col("sqft", "Sq Ft", 9, fmt=FMT_INT),
        Col("beds", "Bedrooms", 9, fmt='0'),
        Col("baths", "Bathrooms", 9, fmt='0.0'),
        Col("pets", "Pets", 12),
        Col("access", "Access / Entry Instructions", 26),
        Col("parking", "Parking", 16),
        Col("special", "Special Requests / Notes", 28),
        Col("jobs", "Completed Jobs", 10, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "addr", r,
            f'COUNTIFS({R("JOBS","property")},{c["addr"]}{r},{R("JOBS","status")},"Completed")')),
        Col("lastclean", "Last Clean", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "addr", r,
            f'IF(_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","property")},{c["addr"]}{r},{R("JOBS","status")},"Completed")=0,"",_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","property")},{c["addr"]}{r},{R("JOBS","status")},"Completed"))')),
        Col("avg", "Avg Price", 11, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "addr", r,
            f'IFERROR(AVERAGEIFS({R("JOBS","price")},{R("JOBS","property")},{c["addr"]}{r},{R("JOBS","status")},"Completed"),"")')),
    ], 400, id_prefix="P-", id_width=3)

    # JOBS --------------------------------------------------------------
    def prop(field, r, c):
        return f'INDEX({R("PROPERTIES",field)},MATCH({c["property"]}{r},{R("PROPERTIES","addr")},0))'

    LOGS["JOBS"] = LogSpec("JOBS", [
        Col("id", "Job ID", 9, "id"),
        Col("date", "Date", 12, fmt=FMT_DATE),
        Col("time", "Start Time", 10, fmt=FMT_TIME),
        Col("customer", "Customer", 20, dv=lambda: R("CUSTOMERS", "name")),
        Col("property", "Property", 24, dv=lambda: R("PROPERTIES", "addr")),
        Col("service", "Service", 17, dv=PT_TYPE),
        Col("plan", "Recurring Plan ID", 10, dv=lambda: R("RECURRING", "id")),
        Col("worker", "Assigned Worker (lead)", 16, dv=lambda: R("STAFF & TASKS", "name")),
        Col("crew", "Crew Size", 8, fmt='0'),
        Col("est", "Est. Labor Hrs (optional)", 10, fmt=FMT_NUM2, note="Total labor hours for all workers. Leave blank to use the automatic estimate from the property details."),
        Col("act", "Actual Labor Hrs", 10, fmt=FMT_NUM2, note="Total labor hours actually worked (all workers combined)."),
        Col("price", "Price (before tax)", 12, fmt=FMT_MONEY),
        Col("tax", "Tax %", 7, fmt=FMT_PCT),
        Col("supin", "Supplies $ (blank = auto)", 10, fmt=FMT_MONEY),
        Col("status", "Status", 12, dv=LST("jobstatus", True)),
        Col("notes", "Notes", 26),
        # automatic
        Col("model", "Auto Est. Hrs", 9, "auto", fmt=FMT_NUM2, formula=lambda r, c:
            f'=IF(OR({c["service"]}{r}="",{c["property"]}{r}=""),"",IFERROR(IF({prop("sqft", r, c)}=0,"",ROUND('
            + model_hours(f'{c["service"]}{r}', prop("sqft", r, c), prop("beds", r, c), prop("baths", r, c)) + ',2)),""))'),
        Col("hrs", "Labor Hrs Used", 9, "auto", fmt=FMT_NUM2, formula=lambda r, c: blank_guard(c, "date", r,
            f'IF(N({c["act"]}{r})>0,{c["act"]}{r},IF(N({c["est"]}{r})>0,{c["est"]}{r},N({c["model"]}{r})))')),
        Col("invoice", "Invoice Total", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "price", r,
            f'ROUND({c["price"]}{r}*(1+N({c["tax"]}{r})),2)')),
        Col("paid", "Paid", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'SUMIFS({R("REVENUE","amount")},{R("REVENUE","job")},{c["id"]}{r})')),
        Col("balance", "Balance Due", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'IF({c["status"]}{r}="Completed",MAX(0,N({c["invoice"]}{r})-{c["paid"]}{r}),"")')),
        Col("paystat", "Payment Status", 11, "auto", formula=lambda r, c: blank_guard(c, "date", r,
            f'IF({c["status"]}{r}<>"Completed",IF({c["paid"]}{r}>0,"Deposit","—"),IF({c["balance"]}{r}<=0.005,"Paid",IF({c["paid"]}{r}>0,"Partial","Unpaid")))')),
        Col("rate", "Pay Rate Used", 9, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'IFERROR(1/(1/INDEX({R("STAFF & TASKS","rate")},MATCH({c["worker"]}{r},{R("STAFF & TASKS","name")},0))),PayRate)')),
        Col("labor", "Labor Cost", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'IF(OR({c["status"]}{r}="Cancelled",{c["status"]}{r}="No-Show"),"",ROUND({c["hrs"]}{r}*{c["rate"]}{r}*(1+Burden),2))')),
        Col("sup", "Supplies Cost", 10, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'IF(OR({c["status"]}{r}="Cancelled",{c["status"]}{r}="No-Show"),"",IF({c["supin"]}{r}<>"",{c["supin"]}{r},ROUND({c["hrs"]}{r}*SupHr*IFERROR(INDEX({PT_SUP},MATCH({c["service"]}{r},{PT_TYPE},0)),1),2)))')),
        Col("miles", "Miles (from MILEAGE)", 9, "auto", fmt=FMT_NUM1, formula=lambda r, c: blank_guard(c, "date", r,
            f'SUMIFS({R("MILEAGE","used")},{R("MILEAGE","job")},{c["id"]}{r})')),
        Col("vehicle", "Vehicle Cost", 10, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'ROUND({c["miles"]}{r}*MileRate,2)')),
        Col("fees", "Payment Fees", 9, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'SUMIFS({R("REVENUE","fee")},{R("REVENUE","job")},{c["id"]}{r})')),
        Col("gp", "Gross Profit (est.)", 12, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "price", r,
            f'IF(OR({c["status"]}{r}="Cancelled",{c["status"]}{r}="No-Show"),"",{c["price"]}{r}-N({c["labor"]}{r})-N({c["sup"]}{r})-{c["vehicle"]}{r}-{c["fees"]}{r})')),
        Col("margin", "Gross Margin", 9, "auto", fmt=FMT_PCT, formula=lambda r, c:
            f'=IF(OR({c["gp"]}{r}="",N({c["price"]}{r})=0),"",{c["gp"]}{r}/{c["price"]}{r})'),
        Col("revhr", "Revenue / Labor Hr", 10, "auto", fmt=FMT_MONEY, formula=lambda r, c:
            f'=IF(OR({c["gp"]}{r}="",N({c["hrs"]}{r})=0),"",{c["price"]}{r}/{c["hrs"]}{r})'),
        Col("var", "Hours Over (+) / Under (−)", 10, "auto", fmt='+0.00;-0.00;"–"', formula=lambda r, c:
            f'=IF(OR(N({c["act"]}{r})=0,N({c["est"]}{r})+N({c["model"]}{r})=0),"",{c["act"]}{r}-IF(N({c["est"]}{r})>0,{c["est"]}{r},{c["model"]}{r}))'),
        Col("qc", "QC Score", 8, "auto", fmt=FMT_PCT, formula=lambda r, c: blank_guard(c, "date", r,
            f'IFERROR(AVERAGEIFS({R("QUALITY CONTROL","score")},{R("QUALITY CONTROL","job")},{c["id"]}{r}),"")')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
        Col("ctype", "Customer Type", 14, "auto", formula=lambda r, c: blank_guard(c, "customer", r,
            f'IFERROR(INDEX({R("CUSTOMERS","ctype")},MATCH({c["customer"]}{r},{R("CUSTOMERS","name")},0)),"")')),
        Col("stime", "Sort Time", 7, "auto", formula=lambda r, c: blank_guard(c, "date", r, f'N({c["time"]}{r})')),
        Col("key", "Schedule Key", 11, "auto", formula=lambda r, c:
            f'=IF(OR({c["date"]}{r}="",{c["status"]}{r}="Cancelled",{c["status"]}{r}="Rescheduled"),"",{c["date"]}{r}&"|"&('
            f'COUNTIFS({R("JOBS","date")},{c["date"]}{r},{R("JOBS","stime")},"<"&{c["stime"]}{r},{R("JOBS","status")},"<>Cancelled",{R("JOBS","status")},"<>Rescheduled")'
            f'+COUNTIFS(${c["date"]}${FIRST}:{c["date"]}{r},{c["date"]}{r},${c["stime"]}${FIRST}:{c["stime"]}{r},{c["stime"]}{r},${c["status"]}${FIRST}:{c["status"]}{r},"<>Cancelled",${c["status"]}${FIRST}:{c["status"]}{r},"<>Rescheduled")))'),
        Col("urank", "Unpaid Order", 8, "auto", formula=lambda r, c:
            f'=IF(N({c["balance"]}{r})<=0.005,"",COUNTIFS({R("JOBS","balance")},">0.005",{R("JOBS","date")},"<"&{c["date"]}{r})'
            f'+COUNTIFS(${c["balance"]}${FIRST}:{c["balance"]}{r},">0.005",${c["date"]}${FIRST}:{c["date"]}{r},{c["date"]}{r}))'),
    ], 1500, id_prefix="J-", id_width=4)

    # RECURRING ---------------------------------------------------------
    def fq(col, r, c):
        return f'INDEX({col},MATCH({c["freq"]}{r},{FREQ_NAME},0))'
    LOGS["RECURRING"] = LogSpec("RECURRING", [
        Col("id", "Plan ID", 9, "id"),
        Col("customer", "Customer", 20, dv=lambda: R("CUSTOMERS", "name")),
        Col("property", "Property", 24, dv=lambda: R("PROPERTIES", "addr")),
        Col("service", "Service", 17, dv=PT_TYPE),
        Col("freq", "Frequency", 13, dv=FREQ_RECUR),
        Col("price", "Price per Visit", 11, fmt=FMT_MONEY),
        Col("worker", "Usual Worker", 15, dv=lambda: R("STAFF & TASKS", "name")),
        Col("start", "Start Date", 12, fmt=FMT_DATE),
        Col("pref", "Preferred Day / Time", 16),
        Col("status", "Plan Status", 10, dv=LST("planstatus", True)),
        Col("notes", "Notes", 24),
        Col("visits", "Visits Completed", 9, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "customer", r,
            f'COUNTIFS({R("JOBS","plan")},{c["id"]}{r},{R("JOBS","status")},"Completed")')),
        Col("lastv", "Last Visit", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF(_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","plan")},{c["id"]}{r},{R("JOBS","status")},"Completed")=0,"",_xlfn.MAXIFS({R("JOBS","date")},{R("JOBS","plan")},{c["id"]}{r},{R("JOBS","status")},"Completed"))')),
        Col("nextb", "Next Booked Visit", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF(_xlfn.MINIFS({R("JOBS","date")},{R("JOBS","plan")},{c["id"]}{r},{R("JOBS","status")},"Scheduled",{R("JOBS","date")},">="&AsOf)=0,"",_xlfn.MINIFS({R("JOBS","date")},{R("JOBS","plan")},{c["id"]}{r},{R("JOBS","status")},"Scheduled",{R("JOBS","date")},">="&AsOf))')),
        Col("due", "Next Visit Due", 12, "auto", fmt=FMT_DATE, formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF({c["nextb"]}{r}<>"",{c["nextb"]}{r},IF({c["lastv"]}{r}="",IF({c["start"]}{r}="","",{c["start"]}{r}),'
            f'IF({c["freq"]}{r}="Monthly",EDATE({c["lastv"]}{r},1),{c["lastv"]}{r}+IFERROR({fq(FREQ_DAYS, r, c)},14))))')),
        Col("daysto", "Days Until Due", 9, "auto", fmt='0;[Red]-0;0', formula=lambda r, c: blank_guard(c, "due", r,
            f'{c["due"]}{r}-AsOf')),
        Col("book", "Booking Status", 17, "auto", formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF({c["status"]}{r}<>"Active","—",IF({c["nextb"]}{r}<>"","Booked",IF({c["due"]}{r}="","Set start date",'
            f'IF({c["due"]}{r}<AsOf,"Overdue — book now",IF({c["due"]}{r}<=AsOf+7,"Due this week — book","Upcoming")))))')),
        Col("mval", "Monthly Value", 11, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF({c["status"]}{r}="Active",N({c["price"]}{r})*IFERROR({fq(FREQ_VPM, r, c)},0),0)')),
        Col("aval", "Annual Value", 11, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "customer", r,
            f'{c["mval"]}{r}*12')),
        Col("needs", "Needs Booking (1/0)", 8, "auto", formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF(OR({c["book"]}{r}="Overdue — book now",{c["book"]}{r}="Due this week — book"),1,0)')),
        Col("rank", "Booking Order", 8, "auto", formula=lambda r, c:
            f'=IF(N({c["needs"]}{r})<>1,"",COUNTIFS({R("RECURRING","needs")},1,{R("RECURRING","due")},"<"&{c["due"]}{r})'
            f'+COUNTIFS(${c["needs"]}${FIRST}:{c["needs"]}{r},1,${c["due"]}${FIRST}:{c["due"]}{r},{c["due"]}{r}))'),
    ], 150, id_prefix="R-", id_width=3)

    # STAFF (roster) ------------------------------------------------------
    LOGS["STAFF & TASKS"] = LogSpec("STAFF & TASKS", [
        Col("name", "Name (unique)", 18),
        Col("role", "Role", 12, dv=LST("roles")),
        Col("rate", "Pay Rate ($/hr)", 11, fmt=FMT_MONEY),
        Col("phone", "Phone", 15),
        Col("email", "Email", 22),
        Col("status", "Status", 10, dv=LST("staffstatus", True)),
        Col("start", "Start Date", 12, fmt=FMT_DATE),
        Col("notes", "Notes", 22),
        Col("mjobs", "Jobs (Rpt Month)", 10, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "name", r,
            f'COUNTIFS({R("JOBS","worker")},{c["name"]}{r},{R("JOBS","status")},"Completed",{R("JOBS","month")},ReportMonth)')),
        Col("mhrs", "Labor Hrs (Rpt Month)", 10, "auto", fmt=FMT_NUM1, formula=lambda r, c: blank_guard(c, "name", r,
            f'SUMIFS({R("JOBS","hrs")},{R("JOBS","worker")},{c["name"]}{r},{R("JOBS","status")},"Completed",{R("JOBS","month")},ReportMonth)')),
        Col("mrev", "Revenue (Rpt Month)", 11, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "name", r,
            f'SUMIFS({R("JOBS","price")},{R("JOBS","worker")},{c["name"]}{r},{R("JOBS","status")},"Completed",{R("JOBS","month")},ReportMonth)')),
        Col("qc", "Avg QC Score (all)", 10, "auto", fmt=FMT_PCT, formula=lambda r, c: blank_guard(c, "name", r,
            f'IFERROR(AVERAGEIFS({R("QUALITY CONTROL","score")},{R("QUALITY CONTROL","worker")},{c["name"]}{r}),"")')),
        Col("issues", "Issues Logged (all)", 9, "auto", fmt=FMT_INT, formula=lambda r, c: blank_guard(c, "name", r,
            f'COUNTIF({R("ISSUES","worker")},{c["name"]}{r})')),
    ], 20)

    LOGS["TASKS"] = LogSpec("STAFF & TASKS", [
        Col("task", "Task / Instruction", 34),
        Col("to", "Assigned To", 16, dv=lambda: R("STAFF & TASKS", "name")),
        Col("job", "Job ID (optional)", 10, dv=lambda: R("JOBS", "id")),
        Col("due", "Due Date", 12, fmt=FMT_DATE),
        Col("prio", "Priority", 9, dv=LST("priority", True)),
        Col("status", "Status", 12, dv=LST("taskstatus", True)),
        Col("notes", "Notes", 26),
        Col("flag", "Due Flag", 11, "auto", formula=lambda r, c: blank_guard(c, "task", r,
            f'IF({c["status"]}{r}="Done","✓ Done",IF({c["due"]}{r}="","",IF({c["due"]}{r}<AsOf,"Overdue",IF({c["due"]}{r}=AsOf,"Due today","Upcoming"))))')),
    ], 200, first=31, header_row=30)

    # REVENUE (payments) ---------------------------------------------------
    LOGS["REVENUE"] = LogSpec("REVENUE", [
        Col("date", "Payment Date", 12, fmt=FMT_DATE),
        Col("job", "Job ID", 10, dv=lambda: R("JOBS", "id")),
        Col("amount", "Amount Received ($)", 13, fmt=FMT_MONEY),
        Col("tip", "Tip ($, optional)", 10, fmt=FMT_MONEY),
        Col("method", "Payment Method", 15, dv=LST("paymeth")),
        Col("ref", "Invoice / Ref #", 13),
        Col("notes", "Notes", 26),
        Col("customer", "Customer", 20, "auto", formula=lambda r, c: blank_guard(c, "job", r,
            f'IFERROR(INDEX({R("JOBS","customer")},MATCH({c["job"]}{r},{R("JOBS","id")},0)),"Job ID not found")')),
        Col("service", "Service", 16, "auto", formula=lambda r, c: blank_guard(c, "job", r,
            f'IFERROR(INDEX({R("JOBS","service")},MATCH({c["job"]}{r},{R("JOBS","id")},0)),"")')),
        Col("fee", "Est. Processing Fee", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "amount", r,
            f'ROUND({c["amount"]}{r}*IF(IFERROR(INDEX({LST("payfee")},MATCH({c["method"]}{r},{LST("paymeth")},0)),"No")="Yes",FeePct,0),2)')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 1500)

    # EXPENSES ----------------------------------------------------------
    LOGS["EXPENSES"] = LogSpec("EXPENSES", [
        Col("date", "Date", 12, fmt=FMT_DATE),
        Col("cat", "Category", 22, dv=LST("expcat")),
        Col("vendor", "Vendor", 18),
        Col("desc", "Description", 26),
        Col("amount", "Amount ($)", 11, fmt=FMT_MONEY),
        Col("paidwith", "Paid With", 14, dv=LST("paymeth")),
        Col("job", "Job ID (optional)", 10, dv=lambda: R("JOBS", "id")),
        Col("receipt", "Receipt Saved?", 9, dv=LST("yesno", True)),
        Col("notes", "Notes", 22),
        Col("oh", "Counted as Overhead?", 11, "auto", formula=lambda r, c: blank_guard(c, "cat", r,
            f'IFERROR(INDEX({LST("expoh")},MATCH({c["cat"]}{r},{LST("expcat")},0)),"Yes")')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 1000)

    # MILEAGE -------------------------------------------------------------
    LOGS["MILEAGE"] = LogSpec("MILEAGE", [
        Col("date", "Date", 12, fmt=FMT_DATE),
        Col("job", "Job ID (optional)", 10, dv=lambda: R("JOBS", "id")),
        Col("purpose", "Purpose", 22),
        Col("from", "From", 16),
        Col("to", "To", 16),
        Col("odo1", "Odometer Start", 11, fmt='#,##0'),
        Col("odo2", "Odometer End", 11, fmt='#,##0'),
        Col("miles", "Miles (if no odometer)", 10, fmt=FMT_NUM1),
        Col("driver", "Driver / Vehicle", 14),
        Col("used", "Miles Used", 9, "auto", fmt=FMT_NUM1, formula=lambda r, c: blank_guard(c, "date", r,
            f'IF(AND({c["odo1"]}{r}<>"",{c["odo2"]}{r}<>""),MAX(0,{c["odo2"]}{r}-{c["odo1"]}{r}),N({c["miles"]}{r}))')),
        Col("value", "Cost at Your Rate", 11, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "date", r,
            f'ROUND({c["used"]}{r}*MileRate,2)')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 1500)

    # SUPPLIES -------------------------------------------------------------
    LOGS["SUPPLIES"] = LogSpec("SUPPLIES", [
        Col("item", "Item", 26),
        Col("cat", "Category", 15, dv=LST("supcat")),
        Col("unit", "Unit", 12),
        Col("cost", "Unit Cost ($)", 10, fmt=FMT_MONEY),
        Col("onhand", "On Hand", 9, fmt='0'),
        Col("reorder", "Reorder Level", 9, fmt='0'),
        Col("supplier", "Preferred Supplier", 18),
        Col("lastbuy", "Last Purchased", 12, fmt=FMT_DATE),
        Col("notes", "Notes (dilution, storage, SDS location)", 30),
        Col("value", "Stock Value", 10, "auto", fmt=FMT_MONEY, formula=lambda r, c: blank_guard(c, "item", r,
            f'N({c["cost"]}{r})*N({c["onhand"]}{r})')),
        Col("flag", "Reorder?", 10, "auto", formula=lambda r, c: blank_guard(c, "item", r,
            f'IF({c["reorder"]}{r}="","",IF(N({c["onhand"]}{r})<={c["reorder"]}{r},"Reorder","OK"))')),
    ], 100)

    # FOLLOW-UPS (manual log) ----------------------------------------------
    LOGS["FOLLOW-UPS"] = LogSpec("FOLLOW-UPS", [
        Col("due", "Due Date", 12, fmt=FMT_DATE),
        Col("contact", "Contact Name", 20),
        Col("type", "Type", 12, dv=LST("futype")),
        Col("related", "Related ID", 10),
        Col("channel", "Channel", 10, dv=LST("channel")),
        Col("workflow", "AI Workflow to Use", 24, dv=LST("workflows")),
        Col("status", "Status", 9, dv=LST("fustatus", True)),
        Col("notes", "Notes / Outcome", 30),
        Col("flag", "Due Flag", 11, "auto", formula=lambda r, c: blank_guard(c, "due", r,
            f'IF({c["status"]}{r}<>"Open","—",IF({c["due"]}{r}<AsOf,"Overdue",IF({c["due"]}{r}=AsOf,"Due today","Upcoming")))')),
    ], 300, first=25, header_row=24)

    # REVIEWS -------------------------------------------------------------
    LOGS["REVIEWS"] = LogSpec("REVIEWS", [
        Col("customer", "Customer", 20, dv=lambda: R("CUSTOMERS", "name")),
        Col("job", "Job ID", 10, dv=lambda: R("JOBS", "id")),
        Col("req", "Date Requested", 12, fmt=FMT_DATE),
        Col("platform", "Platform", 12, dv=LST("platform")),
        Col("status", "Status", 13, dv=LST("revstatus", True)),
        Col("rec", "Date Received", 12, fmt=FMT_DATE),
        Col("rating", "Rating (1–5)", 9, fmt='0'),
        Col("replied", "You Replied?", 9, dv=LST("yesno", True)),
        Col("notes", "Notes / Quote from Review", 34),
        Col("wait", "Days Waiting", 9, "auto", fmt='0', formula=lambda r, c: blank_guard(c, "req", r,
            f'IF(OR({c["status"]}{r}="Received",{c["status"]}{r}="Declined"),"",AsOf-{c["req"]}{r})')),
        Col("chase", "Action", 16, "auto", formula=lambda r, c: blank_guard(c, "customer", r,
            f'IF(AND({c["status"]}{r}="Received",{c["replied"]}{r}<>"Yes"),"Reply to review",IF(N({c["wait"]}{r})>=ReviewWait,IF({c["status"]}{r}="Reminder sent","Close out","Send reminder"),"—"))')),
        Col("recmonth", "Received Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "rec", r,
            f'DATE(YEAR({c["rec"]}{r}),MONTH({c["rec"]}{r}),1)')),
        Col("reqmonth", "Requested Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "req", r,
            f'DATE(YEAR({c["req"]}{r}),MONTH({c["req"]}{r}),1)')),
    ], 500)

    # REFERRALS -----------------------------------------------------------
    LOGS["REFERRALS"] = LogSpec("REFERRALS", [
        Col("date", "Date Referred", 12, fmt=FMT_DATE),
        Col("by", "Referred By (customer)", 20, dv=lambda: R("CUSTOMERS", "name")),
        Col("new", "New Contact Name", 20),
        Col("contact", "Phone / Email", 20),
        Col("status", "Status", 11, dv=LST("refstatus", True)),
        Col("reward", "Reward / Thank-You", 20),
        Col("rstatus", "Reward Status", 10, dv=LST("rewstatus", True)),
        Col("notes", "Notes", 24),
        Col("rev", "Revenue from Referral", 12, "auto", fmt=FMT_MONEY0, formula=lambda r, c: blank_guard(c, "new", r,
            f'SUMIFS({R("JOBS","price")},{R("JOBS","customer")},{c["new"]}{r},{R("JOBS","status")},"Completed")')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 200)

    # QUALITY CONTROL: inspections + issues --------------------------------
    def jl(field, r, c):
        return f'IFERROR(INDEX({R("JOBS",field)},MATCH({c["job"]}{r},{R("JOBS","id")},0)),"")'
    LOGS["QUALITY CONTROL"] = LogSpec("QUALITY CONTROL", [
        Col("date", "Inspection Date", 12, fmt=FMT_DATE),
        Col("job", "Job ID", 10, dv=lambda: R("JOBS", "id")),
        Col("list", "Checklist Used", 18, dv=LST("checklists")),
        Col("inspector", "Inspector", 15),
        Col("checked", "Items Checked", 9, fmt='0'),
        Col("passed", "Items Passed", 9, fmt='0'),
        Col("found", "Issues Found", 28),
        Col("rework", "Rework Needed?", 9, dv=LST("yesno", True)),
        Col("rdate", "Rework Date", 12, fmt=FMT_DATE),
        Col("resolved", "Resolved?", 9, dv=LST("yesno", True)),
        Col("customer", "Customer", 18, "auto", formula=lambda r, c: blank_guard(c, "job", r, jl("customer", r, c))),
        Col("worker", "Worker", 14, "auto", formula=lambda r, c: blank_guard(c, "job", r, jl("worker", r, c))),
        Col("score", "Score", 8, "auto", fmt=FMT_PCT, formula=lambda r, c:
            f'=IF(N({c["checked"]}{r})=0,"",MIN(1,N({c["passed"]}{r})/{c["checked"]}{r}))'),
        Col("result", "Result", 8, "auto", formula=lambda r, c: blank_guard(c, "score", r,
            f'IF({c["score"]}{r}>=QCPass,"Pass","Fail")')),
        Col("open", "Rework Status", 13, "auto", formula=lambda r, c: blank_guard(c, "date", r,
            f'IF({c["rework"]}{r}<>"Yes","—",IF({c["resolved"]}{r}="Yes","Resolved","Rework open"))')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 300)

    LOGS["ISSUES"] = LogSpec("QUALITY CONTROL", [
        Col("date", "Date Reported", 12, fmt=FMT_DATE),
        Col("job", "Job ID", 10, dv=lambda: R("JOBS", "id")),
        Col("by", "Reported By", 18, dv=LST("reporter")),
        Col("type", "Issue Type", 15, dv=LST("issuetype")),
        Col("desc", "What Happened", 28),
        Col("action", "Action Taken / Planned", 28),
        Col("rework", "Rework / Fix Date", 12, fmt=FMT_DATE),
        Col("status", "Status", 11, dv=LST("issuestatus", True)),
        Col("cost", "Cost of Fix ($)", 10, fmt=FMT_MONEY),
        Col("closed", "Resolved Date", 12, fmt=FMT_DATE),
        Col("customer", "Customer", 18, "auto", formula=lambda r, c: blank_guard(c, "job", r, jl("customer", r, c))),
        Col("worker", "Worker", 14, "auto", formula=lambda r, c: blank_guard(c, "job", r, jl("worker", r, c))),
        Col("days", "Days Open", 8, "auto", fmt='0', formula=lambda r, c: blank_guard(c, "date", r,
            f'IF({c["status"]}{r}="Resolved",IF({c["closed"]}{r}="","",{c["closed"]}{r}-{c["date"]}{r}),AsOf-{c["date"]}{r})')),
        Col("month", "Month", 10, "auto", fmt='mmm yyyy', formula=lambda r, c: blank_guard(c, "date", r,
            f'DATE(YEAR({c["date"]}{r}),MONTH({c["date"]}{r}),1)')),
    ], 150, first=312, header_row=311)


define_logs()

# ---------------------------------------------------------------- helpers to write
def title_block(ws, title, subtitle, width_cols, tab):
    ws.sheet_properties.tabColor = TAB[tab]
    ws.sheet_view.showGridLines = False
    last = L(max(width_cols, 6))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = title
    ws["A1"].font = font(18, True, "FFFFFF")
    ws["A1"].fill = fill(C_PRIMARY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws["A2"] = subtitle
    ws["A2"].font = font(10, False, C_MUTED, italic=True)
    ws["A2"].alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30
    for col in range(1, max(width_cols, 6) + 1):
        ws.cell(1, col).fill = fill(C_PRIMARY)
    link(ws["A3"], "START HERE", "← Start Here")
    link(ws["B3"], "DASHBOARD", "Dashboard →")
    ws.row_dimensions[3].height = 18


def link(cell, sheet, text, anchor="A1"):
    cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{anchor}", display=text)
    cell.font = font(9, True, C_BLUE, underline="single")


def write_header_row(ws, spec):
    hr = spec.header_row
    for i, col in enumerate(spec.cols, start=1):
        cell = ws.cell(hr, i, col.header + (" ⚙" if col.kind == "auto" else ""))
        cell.font = font(9, True, "FFFFFF")
        cell.fill = fill(C_AUTOHDR if col.kind == "auto" else C_PRIMARY)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        if col.note:
            cell.comment = Comment(col.note, "Growth OS")
        ws.column_dimensions[L(i)].width = max(ws.column_dimensions[L(i)].width or 0, col.width)
    ws.row_dimensions[hr].height = 42


def write_log(ws, spec, rows=None):
    write_header_row(ws, spec)
    c = spec.letters
    rows = rows or []
    for idx in range(spec.n):
        r = spec.first + idx
        data = rows[idx] if idx < len(rows) else {}
        for i, col in enumerate(spec.cols, start=1):
            cell = ws.cell(r, i)
            if col.kind == "id":
                cell.value = f"{spec.id_prefix}{idx + 1:0{spec.id_width}d}"
                cell.font = font(9, False, C_MUTED)
                cell.fill = fill(C_AUTOFILL)
            elif col.kind == "auto":
                cell.value = col.formula(r, c)
                cell.font = font(9, False, "33434A")
                cell.fill = fill(C_AUTOFILL)
            else:
                v = data.get(col.key)
                if v is not None:
                    cell.value = v
                cell.font = font(9)
            if col.fmt:
                cell.number_format = col.fmt
            cell.border = BORDER
    # data validation
    for i, col in enumerate(spec.cols, start=1):
        if col.dv:
            dvf = col.dv() if callable(col.dv) else col.dv
            dv = DataValidation(type="list", formula1=dvf,
                                allow_blank=True, showErrorMessage=False)
            dv.add(f"{L(i)}{spec.first}:{L(i)}{spec.last}")
            ws.add_data_validation(dv)


def cf_text(ws, rng, col_letter, first, text, fg, bg, op="="):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col_letter}{first}{op}"{text}"'],
                                                   font=Font(name=FONT, color=fg, bold=True), fill=fill(bg)))


def cf_cell(ws, spec, key, text, fg, bg, whole_row=False):
    c = spec.letters[key]
    rng = f"{c}{spec.first}:{c}{spec.last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{c}{spec.first}="{text}"'],
                                                   font=Font(name=FONT, color=fg, bold=True), fill=fill(bg)))


def cf_formula(ws, rng, formula, fg, bg, bold=True):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=Font(name=FONT, color=fg, bold=bold),
                                                   fill=fill(bg)))


def setup_log_sheet(ws, spec, title, subtitle, tab, rows, freeze="C5"):
    title_block(ws, title, subtitle, len(spec.cols), tab)
    write_log(ws, spec, rows)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{spec.header_row}:{L(len(spec.cols))}{spec.last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{spec.header_row}:{spec.header_row}"


def kv_cell(ws, ref, value, fmt=None, bold=False, size=10, color=C_TEXT, fillc=None, border=True, align=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font(size, bold, color)
    if fmt:
        cell.number_format = fmt
    if fillc:
        cell.fill = fill(fillc)
    if border:
        cell.border = BORDER
    if align:
        cell.alignment = align
    return cell


def section(ws, ref, text, span=None, color=C_PRIMARY):
    cell = ws[ref]
    cell.value = text
    cell.font = font(11, True, color)
    cell.border = Border(bottom=Side(style="medium", color=C_ACCENT))
    if span:
        for col in range(cell.column + 1, cell.column + span):
            ws.cell(cell.row, col).border = Border(bottom=Side(style="medium", color=C_ACCENT))

# ---------------------------------------------------------------- SETTINGS
SETROW = {}


def build_settings(wb, demo):
    ws = wb[SET]
    title_block(ws, "SETTINGS", "Your business details, cost assumptions, targets and dropdown lists. "
                "Edit the WHITE cells. Lists marked LOCKED drive formulas — you can add to editable lists, "
                "but do not rename locked values.", 8, "admin")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 3
    r = 5
    for item in SETTINGS_ROWS:
        if item[0] == "section":
            section(ws, f"A{r}", item[1], 3)
            r += 1
            continue
        SETROW[item[0]] = r
        r += 1
    # write values (after all rows known so formulas can reference)
    for item in SETTINGS_ROWS:
        if item[0] == "section":
            continue
        name, label, vdemo, vclean, fmt, note = item
        row = SETROW[name]
        # formula rows (AsOf, ReportMonth, ReportYear) must exist in BOTH files
        v = vdemo if (demo or (isinstance(vdemo, str) and vdemo.startswith("="))) else vclean
        if isinstance(v, str) and v.startswith("="):
            v = v.format(**SETROW)
            auto = True
        else:
            auto = False
        kv_cell(ws, f"A{row}", label, bold=False)
        cell = kv_cell(ws, f"B{row}", v, fmt, bold=True, fillc=C_AUTOFILL if auto else C_INPUT)
        cell.alignment = Alignment(horizontal="left")
        ncell = kv_cell(ws, f"C{row}", note, size=9, color=C_MUTED)
        ncell.alignment = Alignment(wrap_text=True, vertical="center")
        wb.defined_names[name] = DefinedName(name, attr_text=f"'{SET}'!$B${row}")
    # lists
    ws["F3"] = "DROPDOWN LISTS  —  add items in the blank rows of editable lists."
    ws["F3"].font = font(10, True, C_PRIMARY)
    for key, header, vals, locked in LISTS:
        c = LISTCOL[key]
        ws.column_dimensions[c].width = 22
        h = ws[f"{c}4"]
        h.value = header + ("  (LOCKED)" if locked else "")
        h.font = font(9, True, "FFFFFF")
        h.fill = fill(C_AUTOHDR if locked else C_PRIMARY)
        h.alignment = Alignment(wrap_text=True, vertical="center")
        n = len(vals) if locked else LIST_SLOTS
        for i in range(n):
            cell = ws[f"{c}{LIST_FIRST + i}"]
            cell.value = vals[i] if i < len(vals) else None
            cell.font = font(9)
            cell.border = BORDER
            cell.fill = fill(C_AUTOFILL if locked else "FFFFFF")
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"

# ---------------------------------------------------------------- PRICING
def build_pricing(wb):
    ws = wb[PR]
    title_block(ws, "PRICING CALCULATOR", "Your pricing assumptions. Every number is an EXAMPLE starting point — "
                "replace them with your own. These are not market rates or recommendations.", 9, "sales")
    widths = [30, 16, 16, 14, 14, 13, 14, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w

    def hdr(row, labels, auto_from=None):
        for i, t in enumerate(labels, start=1):
            cell = ws.cell(row, i, t)
            cell.font = font(9, True, "FFFFFF")
            cell.fill = fill(C_AUTOHDR if (auto_from and i >= auto_from) else C_PRIMARY)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = BORDER
        ws.row_dimensions[row].height = 36

    section(ws, "A5", "1. CLEANING TYPES  —  how long each type takes and what you bill per labor hour", 9)
    hdr(6, ["Cleaning Type", "Billing Rate per Labor Hr ($)", "Sq Ft Cleaned per Labor Hr", "Minutes per Bedroom",
            "Minutes per Bathroom", "Setup Minutes per Job", "Supplies Multiplier", "Example: 2,000 sq ft, 3 bd / 2 ba (hrs)", "Notes"], auto_from=8)
    for i in range(PT_LAST - PT_FIRST + 1):
        r = PT_FIRST + i
        vals = PRICING_TYPES[i] if i < len(PRICING_TYPES) else (None,) * 8
        for j, v in enumerate(vals[:7]):
            cell = ws.cell(r, j + 1, v)
            cell.font = font(10, j == 0)
            cell.border = BORDER
            cell.number_format = [None, FMT_MONEY, '#,##0', '0', '0', '0', '0.00'][j] or "General"
        ex = ws.cell(r, 8, f'=IF(A{r}="","",IFERROR(ROUND(' + model_hours(f"A{r}", 2000, 3, 2) + ',2),""))')
        ex.fill = fill(C_AUTOFILL); ex.border = BORDER; ex.number_format = FMT_NUM2; ex.font = font(10)
        nc = ws.cell(r, 9, vals[7] if i < len(PRICING_TYPES) else None)
        nc.font = font(9, color=C_MUTED); nc.border = BORDER
    ws.cell(PT_LAST + 1, 1, "Add your own service types in the blank rows (e.g., Post-Construction). They appear in every Service dropdown.").font = font(8, italic=True, color=C_MUTED)

    section(ws, "A18", "2. PROPERTY CONDITION  —  multiplies estimated time", 3)
    hdr(19, ["Condition", "Time Multiplier"])
    for i in range(COND_LAST - COND_FIRST + 1):
        r = COND_FIRST + i
        v = CONDITIONS[i] if i < len(CONDITIONS) else (None, None)
        a = ws.cell(r, 1, v[0]); a.font = font(10, True); a.border = BORDER
        b = ws.cell(r, 2, v[1]); b.number_format = '0.00"x"'; b.border = BORDER; b.font = font(10)

    section(ws, "A26", "3. FREQUENCY  —  recurring discount and visit spacing", 4)
    hdr(27, ["Frequency", "Discount %", "Days Between Visits", "Visits per Month"])
    for i, (n, d, days, vpm) in enumerate(FREQS):
        r = FREQ_FIRST + i
        for j, (v, f) in enumerate([(n, None), (d, FMT_PCT), (days, '0'), (vpm, '0.00')]):
            cell = ws.cell(r, j + 1, v); cell.border = BORDER; cell.font = font(10, j == 0)
            if f: cell.number_format = f
    ws.cell(FREQ_LAST + 1, 1, "Keep 'One-time' in the first row. 'Monthly' visits are scheduled by calendar month.").font = font(8, italic=True, color=C_MUTED)

    section(ws, "A35", "4. ADD-ONS  —  price, time and direct cost per unit", 4)
    hdr(36, ["Add-On", "Price per Unit ($)", "Minutes per Unit", "Direct Cost per Unit ($)"])
    for i in range(ADD_LAST - ADD_FIRST + 1):
        r = ADD_FIRST + i
        v = ADDONS[i] if i < len(ADDONS) else (None,) * 4
        for j, f in enumerate([None, FMT_MONEY, '0', FMT_MONEY]):
            cell = ws.cell(r, j + 1, v[j]); cell.border = BORDER; cell.font = font(10, j == 0)
            if f: cell.number_format = f

    # Rate card
    section(ws, "A53", "5. RATE CARD  —  one-time base prices from your assumptions (before add-ons, discounts, tax)", 9)
    kv_cell(ws, "A54", "Condition for rate card", bold=True)
    kv_cell(ws, "B54", "Average", fillc=C_INPUT, bold=True)
    dv = DataValidation(type="list", formula1=COND_NAME, allow_blank=False); dv.add("B54"); ws.add_data_validation(dv)
    kv_cell(ws, "C54", "Recurring frequency to compare", bold=True)
    kv_cell(ws, "D54", "Biweekly", fillc=C_INPUT, bold=True)
    dv2 = DataValidation(type="list", formula1=FREQ_RECUR, allow_blank=False); dv2.add("D54"); ws.add_data_validation(dv2)
    heads = ["Sq Ft", "Bedrooms", "Bathrooms"] + [f"=A{PT_FIRST + i}" for i in range(4)] + ['="Standard at "&D54', "Standard Monthly Value"]
    for i, t in enumerate(heads, start=1):
        cell = ws.cell(55, i, t)
        cell.font = font(9, True, "FFFFFF"); cell.fill = fill(C_PRIMARY if i <= 3 else C_AUTOHDR)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); cell.border = BORDER
    ws.row_dimensions[55].height = 36
    sizes = [(800, 1, 1), (1000, 2, 1), (1500, 3, 2), (2000, 3, 2), (2500, 4, 3), (3000, 4, 3), (3500, 5, 4), (4000, 5, 4)]
    for i, (sq, bd, ba) in enumerate(sizes):
        r = 56 + i
        for j, v in enumerate([sq, bd, ba]):
            cell = ws.cell(r, j + 1, v); cell.border = BORDER; cell.font = font(10, j == 0)
            cell.number_format = '#,##0' if j == 0 else '0'
        for k in range(4):
            col = 4 + k
            t = f"{L(col)}$55"
            hrs = model_hours(t, f"$A{r}", f"$B{r}", f"$C{r}", "$B$54")
            f = f'=IFERROR(MAX(MinJob,ROUND({hrs}*INDEX({PT_RATE},MATCH({t},{PT_TYPE},0)),0)),"")'
            cell = ws.cell(r, col, f); cell.number_format = FMT_MONEY0; cell.border = BORDER; cell.fill = fill(C_AUTOFILL); cell.font = font(10)
        c8 = ws.cell(r, 8, f'=IFERROR(ROUND(D{r}*(1-INDEX({FREQ_DISC},MATCH($D$54,{FREQ_NAME},0))),0),"")')
        c9 = ws.cell(r, 9, f'=IFERROR(H{r}*INDEX({FREQ_VPM},MATCH($D$54,{FREQ_NAME},0)),"")')
        for cc in (c8, c9):
            cc.number_format = FMT_MONEY0; cc.border = BORDER; cc.fill = fill(C_AUTOFILL); cc.font = font(10)
    ws.cell(64, 1, "Rate card uses: time = (sq ft ÷ sq ft per hr + bedrooms × min + bathrooms × min + setup) × condition; "
            "price = MAX(minimum job price, time × billing rate). Edit rows 56–63 to match the homes you usually quote.").font = font(8, italic=True, color=C_MUTED)
    ws.freeze_panes = "A5"

# ---------------------------------------------------------------- QUOTE BUILDER
def build_quote(wb, demo):
    ws = wb["QUOTE BUILDER"]
    title_block(ws, "QUOTE BUILDER  +  PROFIT ENGINE", "Enter the job details on the left. The right side shows the price, "
                "your costs, gross profit and margin. All assumptions come from PRICING and SETTINGS and can be overridden here. "
                "Results are estimates from YOUR assumptions — not market rates.", 8, "sales")
    for col, w in zip("ABCDEFGH", [34, 22, 12, 12, 3, 36, 18, 30]):
        ws.column_dimensions[col].width = w
    section(ws, "A5", "JOB DETAILS", 4)
    inputs = [
        (6, "Quote date", dt.date(2026, 9, 30) if demo else None, FMT_DATE, None),
        (7, "Customer name", "Avery Collins" if demo else None, None, None),
        (8, "Property address", "14 Birch Lane, Sampletown" if demo else None, None, None),
        (9, "Cleaning type", "Deep Clean", None, PT_TYPE),
        (10, "Square footage", 2200 if demo else 2000, '#,##0', None),
        (11, "Bedrooms", 3, '0', None),
        (12, "Bathrooms", 2.5 if demo else 2, '0.0', None),
        (13, "Property condition", "Average", None, COND_NAME),
        (14, "Frequency", "One-time", None, FREQ_NAME),
        (15, "Number of workers", 2, '0', None),
        (16, "Labor hours override (optional)", None, FMT_NUM2, None),
        (17, "Billing rate per labor hour override (optional)", None, FMT_MONEY, None),
        (18, "Worker pay rate ($/hr)", "=PayRate", FMT_MONEY, None),
        (19, "Supplies cost override (optional)", None, FMT_MONEY, None),
        (20, "Round-trip distance (miles)", 12 if demo else 10, FMT_NUM1, None),
        (21, "Paid travel time per worker (minutes)", "=TravelMin", '0', None),
    ]
    for r, label, v, fmt, dvf in inputs:
        kv_cell(ws, f"A{r}", label)
        ws.merge_cells(f"B{r}:D{r}")
        cell = kv_cell(ws, f"B{r}", v, fmt, bold=True, fillc=C_INPUT)
        cell.alignment = Alignment(horizontal="left")
        for cc in ("C", "D"):
            ws[f"{cc}{r}"].border = BORDER
        if dvf:
            dv = DataValidation(type="list", formula1=dvf, allow_blank=True)
            dv.add(f"B{r}"); ws.add_data_validation(dv)
    ws["A22"] = "Formulas in B18 and B21 pull your SETTINGS defaults — type over them to override for this quote."
    ws["A22"].font = font(8, italic=True, color=C_MUTED)

    section(ws, "A23", "ADD-ONS", 4)
    for i, t in enumerate(["Add-on", "Quantity", "Price", "Minutes"], start=1):
        cell = ws.cell(24, i, t); cell.font = font(9, True, "FFFFFF")
        cell.fill = fill(C_PRIMARY if i <= 2 else C_AUTOHDR); cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    demo_addons = [("Inside oven", 1), ("Inside refrigerator", 1), ("Interior windows (per 10)", 1)] if demo else []
    dva = DataValidation(type="list", formula1=ADD_NAME, allow_blank=True)
    for i in range(6):
        r = 25 + i
        a = demo_addons[i] if i < len(demo_addons) else (None, None)
        kv_cell(ws, f"A{r}", a[0], fillc=C_INPUT); dva.add(f"A{r}")
        kv_cell(ws, f"B{r}", a[1], '0', fillc=C_INPUT)
        kv_cell(ws, f"C{r}", f'=IF(OR(A{r}="",N(B{r})=0),0,B{r}*IFERROR(INDEX({ADD_PRICE},MATCH(A{r},{ADD_NAME},0)),0))', FMT_MONEY, fillc=C_AUTOFILL)
        kv_cell(ws, f"D{r}", f'=IF(OR(A{r}="",N(B{r})=0),0,B{r}*IFERROR(INDEX({ADD_MIN},MATCH(A{r},{ADD_NAME},0)),0))', '0', fillc=C_AUTOFILL)
    ws.add_data_validation(dva)

    section(ws, "A32", "DISCOUNT & TAX", 4)
    extra = [(33, "Manual discount ($)", 0, FMT_MONEY, None),
             (34, "Sales tax % (only if it applies to you)", "=TaxDefault", FMT_PCT, None),
             (35, "Customer pays by card/online (adds fee to costs)?", "Yes", None, '"Yes,No"')]
    for r, label, v, fmt, dvf in extra:
        kv_cell(ws, f"A{r}", label)
        ws.merge_cells(f"B{r}:D{r}")
        cell = kv_cell(ws, f"B{r}", v, fmt, bold=True, fillc=C_INPUT); cell.alignment = Alignment(horizontal="left")
        if dvf:
            dv = DataValidation(type="list", formula1=dvf, allow_blank=False); dv.add(f"B{r}"); ws.add_data_validation(dv)

    # hidden-ish helper cells in column H (visible, labelled)
    m = f"MATCH($B$9,{PT_TYPE},0)"
    # RESULTS
    section(ws, "F5", "RESULTS  (automatic)", 3)
    res = [
        ("hdr", "TIME"),
        (7, "Estimated labor hours (model, before add-ons)", f'=IFERROR(ROUND({model_hours("$B$9", "$B$10", "$B$11", "$B$12", "$B$13")},2),0)', FMT_NUM2),
        (8, "Labor hours used (incl. add-ons)", '=IF(N(B16)>0,B16,G7)+SUM(D25:D30)/60', FMT_NUM2),
        (9, "Time on site with your crew (hours)", '=IFERROR(G8/MAX(1,N(B15)),0)', FMT_NUM2),
        ("hdr", "PRICE"),
        (11, "Billing rate used ($ per labor hour)", f'=IF(N(B17)>0,B17,IFERROR(INDEX({PT_RATE},{m}),0))', FMT_MONEY),
        (12, "Base service price", '=IF(G7=0,0,MAX(MinJob,ROUND(IF(N(B16)>0,B16,G7)*G11,2)))', FMT_MONEY),
        (13, "Frequency discount", f'=-ROUND(G12*IFERROR(INDEX({FREQ_DISC},MATCH(B14,{FREQ_NAME},0)),0),2)', FMT_MONEY),
        (14, "Add-ons", '=SUM(C25:C30)', FMT_MONEY),
        (15, "Manual discount", '=-N(B33)', FMT_MONEY),
        (16, "PRICE BEFORE TAX", '=MAX(0,G12+G13+G14+G15)', FMT_MONEY),
        (17, "Sales tax", '=ROUND(G16*N(B34),2)', FMT_MONEY),
        (18, "CUSTOMER TOTAL", '=G16+G17', FMT_MONEY),
        ("hdr", "DIRECT COSTS"),
        (20, "Labor cost (hours × pay × (1 + extra labor %))", '=ROUND(G8*N(B18)*(1+Burden),2)', FMT_MONEY),
        (21, "Paid travel time", '=ROUND(N(B21)/60*MAX(1,N(B15))*N(B18)*(1+Burden),2)', FMT_MONEY),
        (22, "Vehicle cost (miles × your rate)", '=ROUND(N(B20)*MileRate,2)', FMT_MONEY),
        (23, "Supplies", f'=IF(B19<>"",N(B19),ROUND(G8*SupHr*IFERROR(INDEX({PT_SUP},{m}),1),2))', FMT_MONEY),
        (24, "Add-on direct costs", f'=SUMPRODUCT(N(+B25:B30),IFERROR(INDEX({ADD_COST},N(IF(1,MATCH(A25:A30,{ADD_NAME},0)))),0))', FMT_MONEY),
        (25, "Payment processing fee", '=IF(B35="Yes",ROUND(G18*FeePct,2),0)', FMT_MONEY),
        (26, "TOTAL DIRECT COSTS", '=SUM(G20:G25)', FMT_MONEY),
        ("hdr", "PROFIT"),
        (28, "GROSS PROFIT", '=G16-G26', FMT_MONEY),
        (29, "GROSS MARGIN %", '=IF(G16=0,0,G28/G16)', FMT_PCT),
        (30, "Revenue per labor hour", '=IF(G8=0,0,G16/G8)', FMT_MONEY),
        (31, "Gross profit per labor hour", '=IF(G8=0,0,G28/G8)', FMT_MONEY),
        (32, "Break-even price (before tax)", '=IFERROR(ROUND((G26-G25)/(1-IF(B35="Yes",FeePct*(1+N(B34)),0)),2),0)', FMT_MONEY),
        (33, "Price needed for your target margin", '=IFERROR(ROUND((G26-G25)/(1-TargetMargin-IF(B35="Yes",FeePct*(1+N(B34)),0)),2),0)', FMT_MONEY),
        (34, "Monthly value if recurring", f'=IFERROR(G16*INDEX({FREQ_VPM},MATCH(B14,{FREQ_NAME},0)),0)', FMT_MONEY),
        (36, "Overhead share (labor hours × SETTINGS rate)", '=ROUND(G8*OverheadHr,2)', FMT_MONEY),
        (37, "Est. profit after overhead share", '=G28-G36', FMT_MONEY),
        (35, "Margin check", '=IF(G16=0,"Enter job details",IF(G28<0,"Below break-even — raise price or reduce hours",IF(G29<TargetMargin,"Below your target margin","Meets your target margin")))', None),
    ]
    rr = 6
    for item in res:
        if item[0] == "hdr":
            cell = ws[f"F{rr}"]; cell.value = item[1]; cell.font = font(9, True, C_ACCENT)
            rr += 1
            continue
        r, label, f, fmt = item
        rr = r + 1
        big = label.isupper() or label.startswith("GROSS")
        kv_cell(ws, f"F{r}", label, bold=big)
        cell = kv_cell(ws, f"G{r}", f, fmt, bold=True, size=12 if big else 10,
                       fillc=C_GREENBG if big else C_AUTOFILL, color=C_PRIMARY if big else C_TEXT)
    # Add-on direct cost: simpler robust formula (avoid array semantics differences)
    ws["G24"] = "=" + "+".join(
        f'IF(OR(A{r}="",N(B{r})=0),0,B{r}*IFERROR(INDEX({ADD_COST},MATCH(A{r},{ADD_NAME},0)),0))' for r in range(25, 31))
    cf_formula(ws, "G35", 'LEFT(G35,5)="Below"', C_RED, C_REDBG)
    cf_formula(ws, "G35", 'G35="Meets your target margin"', C_GREEN, C_GREENBG)
    cf_formula(ws, "G28", 'G28<0', C_RED, C_REDBG)
    cf_formula(ws, "G29", 'AND(G16>0,G29<TargetMargin)', C_AMBER, C_AMBERBG)

    section(ws, "A38", "QUOTE SUMMARY  —  copy into your quote, the LEADS sheet, or AI Workflow 02", 8)
    ws.merge_cells("A39:H41")
    s = ('=IF(G16=0,"Enter job details above.","Quote for "&IF(B7="","[customer]",B7)&" at "&IF(B8="","[address]",B8)&": "&B9&", "'
         '&TEXT(B10,"#,##0")&" sq ft, "&B11&" bd / "&B12&" ba, "&LOWER(B13)&" condition, "&LOWER(B14)&". "'
         '&"Crew of "&B15&", about "&TEXT(G9,"0.0")&" hours on site. Add-ons: "&IF(G14=0,"none",TEXT(G14,"$#,##0.00"))&". "'
         '&"Price before tax "&TEXT(G16,"$#,##0.00")&IF(G17>0,", tax "&TEXT(G17,"$#,##0.00"),"")&", total "&TEXT(G18,"$#,##0.00")&".")')
    c = ws["A39"]; c.value = s; c.font = font(10); c.alignment = Alignment(wrap_text=True, vertical="top"); c.fill = fill(C_LIGHT)
    ws.row_dimensions[39].height = 22
    ws["A43"] = ("How to use: price shown is what YOUR assumptions produce. Compare it to what you know about your local market, "
                 "then decide. Copy the price-before-tax into the LEADS Quote Value column when you send the quote.")
    ws["A43"].font = font(8, italic=True, color=C_MUTED)
    ws.freeze_panes = "A5"

# ---------------------------------------------------------------- SCHEDULE
SLOTS = 10


def build_schedule(wb):
    ws = wb["SCHEDULE"]
    title_block(ws, "WEEKLY SCHEDULE", "Shows every scheduled/completed job from JOBS for the chosen week, sorted by start time. "
                "Type any date in the week you want to see (blank = the current week).", 8, "ops")
    ws.column_dimensions["A"].width = 16
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 27
    kv_cell(ws, "A4", "Week containing:", bold=True)
    kv_cell(ws, "B4", None, FMT_DATE, bold=True, fillc=C_INPUT)
    kv_cell(ws, "C4", '="Week of "&TEXT(B5,"mmm d, yyyy")', bold=True, color=C_PRIMARY, border=False)
    kv_cell(ws, "A5", "Monday", color=C_MUTED, border=False)
    ws["B5"] = '=IF(B4="",AsOf,B4)-WEEKDAY(IF(B4="",AsOf,B4),3)'
    ws["B5"].number_format = FMT_DATE; ws["B5"].font = font(9, color=C_MUTED)
    J = LOGS["JOBS"]
    hr = 7
    kv_cell(ws, f"A{hr}", "Slot", bold=True, color="FFFFFF", fillc=C_PRIMARY)
    kv_cell(ws, f"A{hr+1}", "", fillc=C_PRIMARY)
    for d in range(7):
        col = L(2 + d)
        c1 = kv_cell(ws, f"{col}{hr}", f'=TEXT($B$5+{d},"dddd")', bold=True, color="FFFFFF", fillc=C_PRIMARY,
                     align=Alignment(horizontal="center"))
        c2 = kv_cell(ws, f"{col}{hr+1}", f'=$B$5+{d}', 'mmm d', bold=True, color="FFFFFF", fillc=C_PRIMARY,
                     align=Alignment(horizontal="center"))
        for s in range(1, SLOTS + 1):
            r = hr + 1 + s
            f = (f'=IFERROR(INDEX({J.rng("id")},MATCH({col}${hr+1}&"|"&{s},{J.rng("key")},0)),"")')
            # Display: time · customer · worker
            disp = (f'=IFERROR(TEXT(INDEX({J.rng("time")},MATCH({col}${hr+1}&"|"&$A{r},{J.rng("key")},0)),"h:mm AM/PM")&"  "'
                    f'&INDEX({J.rng("customer")},MATCH({col}${hr+1}&"|"&$A{r},{J.rng("key")},0))&CHAR(10)'
                    f'&INDEX({J.rng("service")},MATCH({col}${hr+1}&"|"&$A{r},{J.rng("key")},0))&" · "'
                    f'&INDEX({J.rng("worker")},MATCH({col}${hr+1}&"|"&$A{r},{J.rng("key")},0))&" · "'
                    f'&INDEX({J.rng("id")},MATCH({col}${hr+1}&"|"&$A{r},{J.rng("key")},0)),"")')
            cell = kv_cell(ws, f"{col}{r}", disp, size=9, align=Alignment(wrap_text=True, vertical="top"))
    for s in range(1, SLOTS + 1):
        r = hr + 1 + s
        kv_cell(ws, f"A{r}", s, '0', bold=True, color=C_MUTED, fillc=C_LIGHT, align=Alignment(horizontal="center", vertical="top"))
        ws.row_dimensions[r].height = 30
    tr = hr + SLOTS + 2
    labels = [("Jobs", lambda col: f'=COUNTIFS({J.rng("date")},{col}${hr+1},{J.rng("status")},"<>Cancelled",{J.rng("status")},"<>Rescheduled")', '0'),
              ("Labor hours", lambda col: f'=SUMIFS({J.rng("hrs")},{J.rng("date")},{col}${hr+1},{J.rng("status")},"<>Cancelled",{J.rng("status")},"<>Rescheduled")', FMT_NUM1),
              ("Revenue", lambda col: f'=SUMIFS({J.rng("price")},{J.rng("date")},{col}${hr+1},{J.rng("status")},"<>Cancelled",{J.rng("status")},"<>Rescheduled")', FMT_MONEY0)]
    for i, (lab, fn, fmt) in enumerate(labels):
        r = tr + i
        kv_cell(ws, f"A{r}", lab, bold=True, fillc=C_LIGHT)
        for d in range(7):
            col = L(2 + d)
            kv_cell(ws, f"{col}{r}", fn(col), fmt, bold=True, fillc=C_AUTOFILL, align=Alignment(horizontal="center"))
    ws[f"A{tr+3}"] = f"Shows up to {SLOTS} jobs per day. Cancelled and rescheduled jobs are hidden. Slots fill in start-time order."
    ws[f"A{tr+3}"].font = font(8, italic=True, color=C_MUTED)
    # recurring needing booking
    rr = tr + 5
    section(ws, f"A{rr}", "RECURRING VISITS THAT NEED BOOKING  (overdue or due within 7 days — add them to JOBS)", 8)
    RC = LOGS["RECURRING"]
    heads = ["#", "Plan ID", "Customer", "Service", "Frequency", "Next Due", "Usual Worker", "Status"]
    keys = [None, "id", "customer", "service", "freq", "due", "worker", "book"]
    for i, h in enumerate(heads):
        kv_cell(ws, f"{L(1+i)}{rr+1}", h, bold=True, color="FFFFFF", fillc=C_AUTOHDR)
    for k in range(1, 13):
        r = rr + 1 + k
        kv_cell(ws, f"A{r}", k, '0', color=C_MUTED, fillc=C_LIGHT)
        for i, key in enumerate(keys[1:], start=1):
            f = f'=IFERROR(INDEX({RC.rng(key)},MATCH($A{r},{RC.rng("rank")},0)),"")'
            kv_cell(ws, f"{L(1+i)}{r}", f, FMT_DATE if key == "due" else None, size=9)
    ws.freeze_panes = "B9"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

# ---------------------------------------------------------------- FOLLOW-UPS top section
def build_followups(wb, rows):
    ws = wb["FOLLOW-UPS"]
    spec = LOGS["FOLLOW-UPS"]
    setup_log_sheet(ws, spec, "FOLLOW-UPS", "Top: automatic lists of leads to chase and unpaid jobs. Bottom: your own follow-up log "
                    "for anything else (reviews, reactivation, complaints). Pick an AI workflow to draft the message.", "growth", rows, freeze="A5")
    LE, J = LOGS["LEADS"], LOGS["JOBS"]
    section(ws, "A4", "LEADS TO FOLLOW UP  (soonest first, open leads only)", 5)
    heads = ["#", "Lead", "Stage", "Next Follow-Up", "Status"]
    keys = ["name", "stage", "next", "fu"]
    for i, h in enumerate(heads):
        kv_cell(ws, f"{L(1+i)}5", h, bold=True, color="FFFFFF", fillc=C_AUTOHDR)
    for k in range(1, 16):
        r = 5 + k
        kv_cell(ws, f"A{r}", k, '0', color=C_MUTED, fillc=C_LIGHT)
        for i, key in enumerate(keys, start=1):
            f = f'=IFERROR(INDEX({LE.rng(key)},MATCH($A{r},{LE.rng("rank")},0)),"")'
            kv_cell(ws, f"{L(1+i)}{r}", f, FMT_DATE if key == "next" else None, size=9)
    cf_formula(ws, "E6:E20", 'E6="Overdue"', C_RED, C_REDBG)
    cf_formula(ws, "E6:E20", 'E6="Due today"', C_AMBER, C_AMBERBG)
    section(ws, "G4", "UNPAID COMPLETED JOBS  (oldest first)", 5)
    heads = ["#", "Job ID", "Customer", "Job Date", "Balance Due"]
    keys = ["id", "customer", "date", "balance"]
    for i, h in enumerate(heads):
        kv_cell(ws, f"{L(7+i)}5", h, bold=True, color="FFFFFF", fillc=C_AUTOHDR)
    for k in range(1, 16):
        r = 5 + k
        kv_cell(ws, f"G{r}", k, '0', color=C_MUTED, fillc=C_LIGHT)
        for i, key in enumerate(keys, start=1):
            f = f'=IFERROR(INDEX({J.rng(key)},MATCH($G{r},{J.rng("urank")},0)),"")'
            kv_cell(ws, f"{L(7+i)}{r}", f, FMT_DATE if key == "date" else (FMT_MONEY if key == "balance" else None), size=9)
    ws["A22"] = "YOUR FOLLOW-UP LOG"
    ws["A22"].font = font(11, True, C_PRIMARY)
    ws["A23"] = "Log any follow-up here. The Due Flag turns red when it is overdue. Mark Status 'Done' when finished."
    ws["A23"].font = font(8, italic=True, color=C_MUTED)
    ws.freeze_panes = None
    cf_cell(ws, spec, "flag", "Overdue", C_RED, C_REDBG)
    cf_cell(ws, spec, "flag", "Due today", C_AMBER, C_AMBERBG)

# ---------------------------------------------------------------- QC CHECKLISTS
CHECKLISTS = {
    "Residential Standard": [
        ("Kitchen", ["Countertops and backsplash wiped", "Sink and faucet cleaned and shined", "Outside of appliances wiped",
                     "Stovetop cleaned", "Microwave inside and out", "Cabinet fronts spot-cleaned", "Floor vacuumed and mopped",
                     "Trash emptied and liner replaced"]),
        ("Bathrooms", ["Toilet cleaned inside, outside and base", "Shower/tub scrubbed", "Sink, counter and faucet cleaned",
                       "Mirrors streak-free", "Floor cleaned", "Towels straightened / replaced if requested"]),
        ("Living areas & bedrooms", ["Surfaces dusted (reachable)", "Light switches and door handles wiped",
                                     "Floors vacuumed / mopped", "Beds made (linens changed if requested)", "Mirrors and glass spot-cleaned"]),
        ("Wrap-up", ["Customer special requests completed", "Supplies and equipment removed", "Doors locked / alarm set as instructed"]),
    ],
    "Deep Clean": [
        ("Detail work", ["Baseboards hand-wiped", "Door frames and doors wiped", "Light fixtures and ceiling fans dusted",
                         "Window sills and tracks cleaned", "Vents and returns dusted", "Switch plates and outlets wiped"]),
        ("Kitchen", ["Cabinet fronts degreased", "Backsplash grout scrubbed", "Appliance exteriors detailed",
                     "Range hood wiped", "Inside microwave detailed", "Behind/under movable items cleaned"]),
        ("Bathrooms", ["Soap scum and hard-water buildup removed", "Grout scrubbed", "Fixtures descaled and polished",
                       "Exhaust fan cover dusted", "Behind toilet cleaned"]),
        ("Wrap-up", ["Add-ons completed (oven, fridge, windows)", "Final walk-through with room-by-room check", "Before/after photos taken (with permission)"]),
    ],
    "Move-In/Move-Out": [
        ("Every room", ["Closets: shelves and rods wiped", "Walls spot-cleaned", "Baseboards and trim wiped", "Windows sills/tracks cleaned",
                        "Floors vacuumed and mopped", "Light fixtures wiped"]),
        ("Kitchen", ["Inside all cabinets and drawers", "Inside oven", "Inside refrigerator/freezer", "Inside dishwasher wiped",
                     "Countertops and sink detailed"]),
        ("Bathrooms", ["Inside vanities and drawers", "Tub/shower detailed", "Toilet detailed including base", "Mirrors and fixtures polished"]),
        ("Wrap-up", ["All trash removed from property", "Photos of each room taken", "Keys/access returned as instructed"]),
    ],
    "Airbnb Turnover": [
        ("Reset", ["Linens stripped and fresh set made", "Towels replaced and staged", "Dishes washed and put away",
                   "Trash removed from all rooms", "Floors vacuumed and mopped"]),
        ("Restock", ["Toilet paper restocked", "Soap / shampoo restocked", "Coffee / kitchen basics restocked", "Paper towels restocked"]),
        ("Guest-ready check", ["Thermostat set per host instructions", "Lights and TV remote checked", "Damage or missing items photographed and reported",
                               "Lost-and-found items logged", "Welcome items staged as host requested", "Door locked / code reset as instructed"]),
    ],
    "Commercial": [
        ("Work areas", ["Desks and surfaces wiped (clear areas only)", "Trash and recycling emptied", "High-touch points disinfected (as agreed)",
                        "Floors vacuumed / mopped", "Glass doors and partitions spot-cleaned"]),
        ("Restrooms", ["Toilets and urinals cleaned", "Sinks and counters cleaned", "Mirrors cleaned", "Dispensers restocked", "Floors mopped"]),
        ("Break room", ["Counters and sink cleaned", "Microwave wiped", "Tables and chairs wiped", "Floor cleaned"]),
        ("Close-out", ["Lights off / doors locked per client instructions", "Issues noted for client (maintenance, supply needs)"]),
    ],
    "Final Inspection": [
        ("Overall", ["Every room on the job sheet was completed", "No streaks on mirrors or glass", "No dust on reachable surfaces",
                     "Floors free of debris and streaks", "Trash removed", "Customer special requests completed",
                     "Nothing left behind (tools, cloths, supplies)", "Property secured as instructed", "Photos taken if required",
                     "Any damage or issue reported to office"]),
    ],
}


def build_checklists(wb):
    ws = wb["QC CHECKLISTS"]
    title_block(ws, "QUALITY-CONTROL CHECKLISTS", "Printable checklists for each service. Mark each item Yes / No / N/A. "
                "Enter the Yes count and total checked into the QUALITY CONTROL log. Edit tasks to match your service standards.", 5, "ops")
    for col, w in zip("ABCDE", [24, 52, 10, 3, 30]):
        ws.column_dimensions[col].width = w
    r = 5
    dv = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    ws["E5"] = "JUMP TO"; ws["E5"].font = font(9, True, C_PRIMARY)
    jump = 6
    for name, groups in CHECKLISTS.items():
        link(ws[f"E{jump}"], "QC CHECKLISTS", name, f"A{r}")
        jump += 1
        start = r
        ws.merge_cells(f"A{r}:C{r}")
        kv_cell(ws, f"A{r}", name.upper() + " CHECKLIST", bold=True, size=12, color="FFFFFF", fillc=C_PRIMARY)
        r += 1
        kv_cell(ws, f"A{r}", "Job ID / Date / Cleaner:", size=9, color=C_MUTED)
        ws.merge_cells(f"B{r}:C{r}")
        kv_cell(ws, f"B{r}", None, fillc=C_INPUT)
        r += 1
        first_item = r
        for area, tasks in groups:
            for t in tasks:
                kv_cell(ws, f"A{r}", area, size=9, color=C_MUTED)
                kv_cell(ws, f"B{r}", t, size=10)
                kv_cell(ws, f"C{r}", None, fillc=C_INPUT, align=Alignment(horizontal="center"))
                dv.add(f"C{r}")
                r += 1
        last_item = r - 1
        kv_cell(ws, f"A{r}", "Score", bold=True, fillc=C_LIGHT)
        kv_cell(ws, f"B{r}", f'=IF(COUNTIF(C{first_item}:C{last_item},"Yes")+COUNTIF(C{first_item}:C{last_item},"No")=0,"Mark items to see score",'
                f'COUNTIF(C{first_item}:C{last_item},"Yes")&" of "&(COUNTIF(C{first_item}:C{last_item},"Yes")+COUNTIF(C{first_item}:C{last_item},"No"))&" passed")',
                bold=True, fillc=C_LIGHT)
        kv_cell(ws, f"C{r}", f'=IFERROR(COUNTIF(C{first_item}:C{last_item},"Yes")/(COUNTIF(C{first_item}:C{last_item},"Yes")+COUNTIF(C{first_item}:C{last_item},"No")),"")',
                FMT_PCT, bold=True, fillc=C_LIGHT)
        cf_formula(ws, f"C{first_item}:C{last_item}", f'C{first_item}="No"', C_RED, C_REDBG)
        cf_formula(ws, f"C{first_item}:C{last_item}", f'C{first_item}="Yes"', C_GREEN, C_GREENBG)
        r += 3
    ws.add_data_validation(dv)
    ws.freeze_panes = "A5"

# ---------------------------------------------------------------- PROFITABILITY
def build_profitability(wb):
    ws = wb["PROFITABILITY"]
    title_block(ws, "JOB PROFITABILITY", "Where you make (and lose) money: by service, by worker and by customer type. "
                "Completed jobs only. Profit figures are estimates based on your cost assumptions.", 16, "money")
    J = LOGS["JOBS"]
    kv_cell(ws, "A4", "Period:", bold=True)
    kv_cell(ws, "B4", "Reporting month", bold=True, fillc=C_INPUT)
    dv = DataValidation(type="list", formula1=LST("period", True), allow_blank=False); dv.add("B4"); ws.add_data_validation(dv)
    kv_cell(ws, "C4", "From", color=C_MUTED)
    kv_cell(ws, "D4", '=IF(B4="All time",DATE(1900,1,1),IF(B4="Year to date",DATE(ReportYear,1,1),ReportMonth))', FMT_DATE, fillc=C_AUTOFILL)
    kv_cell(ws, "E4", "To", color=C_MUTED)
    kv_cell(ws, "F4", '=IF(B4="All time",DATE(2999,12,31),IF(B4="Year to date",MAX(AsOf,EOMONTH(ReportMonth,0)),EOMONTH(ReportMonth,0)))', FMT_DATE, fillc=C_AUTOFILL)
    ws.column_dimensions["A"].width = 24
    for i in range(2, 17):
        ws.column_dimensions[L(i)].width = 12
    heads = ["Jobs", "Revenue", "Labor Cost", "Supplies", "Vehicle", "Fees", "Gross Profit", "Margin", "Avg Price",
             "Labor Hrs", "Revenue / Labor Hr", "Profit / Labor Hr", "Jobs Below Target", "Jobs Losing Money", "Avg Hours Over/Under"]
    fmts = [FMT_INT, FMT_MONEY0, FMT_MONEY0, FMT_MONEY0, FMT_MONEY0, FMT_MONEY0, FMT_MONEY0, FMT_PCT, FMT_MONEY0,
            FMT_NUM1, FMT_MONEY, FMT_MONEY, FMT_INT, FMT_INT, '+0.00;-0.00;"–"']
    crit_base = f'{J.rng("status")},"Completed",{J.rng("date")},">="&$D$4,{J.rng("date")},"<="&$F$4'

    def block(top, title, dim_key, items):
        section(ws, f"A{top}", title, 16)
        kv_cell(ws, f"A{top+1}", "", fillc=C_PRIMARY)
        for i, h in enumerate(heads, start=2):
            c = kv_cell(ws, f"{L(i)}{top+1}", h, bold=True, size=9, color="FFFFFF", fillc=C_PRIMARY,
                        align=Alignment(wrap_text=True, horizontal="center", vertical="center"))
        ws.row_dimensions[top + 1].height = 30
        rows = []
        for k, item in enumerate(items):
            r = top + 2 + k
            rows.append(r)
            kv_cell(ws, f"A{r}", item, bold=True)
            crit = f'{crit_base},{J.rng(dim_key)},$A{r}'
            fs = [
                f'=IF($A{r}="","",COUNTIFS({crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("price")},{crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("labor")},{crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("sup")},{crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("vehicle")},{crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("fees")},{crit}))',
                f'=IF($A{r}="","",SUMIFS({J.rng("gp")},{crit}))',
                f'=IF(N(C{r})=0,"",H{r}/C{r})',
                f'=IF(N(B{r})=0,"",C{r}/B{r})',
                f'=IF($A{r}="","",SUMIFS({J.rng("hrs")},{crit}))',
                f'=IF(N(K{r})=0,"",C{r}/K{r})',
                f'=IF(N(K{r})=0,"",H{r}/K{r})',
                f'=IF($A{r}="","",COUNTIFS({crit},{J.rng("margin")},"<"&TargetMargin))',
                f'=IF($A{r}="","",COUNTIFS({crit},{J.rng("gp")},"<0"))',
                f'=IF($A{r}="","",IFERROR(AVERAGEIFS({J.rng("var")},{crit}),""))',
            ]
            for i, (f, fm) in enumerate(zip(fs, fmts), start=2):
                kv_cell(ws, f"{L(i)}{r}", f, fm, fillc=C_AUTOFILL)
        cf_formula(ws, f"I{rows[0]}:I{rows[-1]}", f'AND(I{rows[0]}<>"",I{rows[0]}<TargetMargin)', C_AMBER, C_AMBERBG)
        return rows

    svc_items = [f"='{PR}'!A{PT_FIRST + i}&\"\"" for i in range(PT_LAST - PT_FIRST + 1)]
    svc_rows = block(6, "BY SERVICE TYPE", "service", svc_items)
    tr = svc_rows[-1] + 1
    kv_cell(ws, f"A{tr}", "TOTAL", bold=True, fillc=C_LIGHT)
    for i in range(2, 17):
        col = L(i)
        if col in "BCDEFGHKNO":
            f = f"=SUM({col}{svc_rows[0]}:{col}{svc_rows[-1]})"
        elif col == "I":
            f = f'=IF(N(C{tr})=0,"",H{tr}/C{tr})'
        elif col == "J":
            f = f'=IF(N(B{tr})=0,"",C{tr}/B{tr})'
        elif col == "L":
            f = f'=IF(N(K{tr})=0,"",C{tr}/K{tr})'
        elif col == "M":
            f = f'=IF(N(K{tr})=0,"",H{tr}/K{tr})'
        else:
            f = f'=IFERROR(AVERAGEIFS({J.rng("var")},{crit_base}),"")'
        kv_cell(ws, f"{col}{tr}", f, fmts[i - 2], bold=True, fillc=C_LIGHT)
    SV_TOTAL = tr
    staff = LOGS["STAFF & TASKS"]
    w_items = [f"='STAFF & TASKS'!A{staff.first + i}&\"\"" for i in range(10)]
    w_rows = block(tr + 3, "BY WORKER (lead worker on the job)", "worker", w_items)
    ct_items = [f"=SETTINGS!{LISTCOL['ctype']}{LIST_FIRST + i}&\"\"" for i in range(6)]
    ct_rows = block(w_rows[-1] + 3, "BY CUSTOMER TYPE", "ctype", ct_items)
    note_r = ct_rows[-1] + 2
    ws[f"A{note_r}"] = ("Notes: Gross profit = price − labor (with your extra labor %) − supplies − vehicle − payment fees. "
                        "It does not include overhead such as insurance, software or marketing — see MONTHLY for operating profit. "
                        "Amber margin = below your target in SETTINGS.")
    ws[f"A{note_r}"].font = font(8, italic=True, color=C_MUTED)
    # chart margin by service
    ch = BarChart(); ch.type = "bar"; ch.style = 10
    ch.title = "Gross profit by service"
    ch.y_axis.title = None; ch.x_axis.title = None
    data = Reference(ws, min_col=8, min_row=svc_rows[0] - 1, max_row=svc_rows[0] + len(PRICING_TYPES) - 1)
    cats = Reference(ws, min_col=1, min_row=svc_rows[0], max_row=svc_rows[0] + len(PRICING_TYPES) - 1)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None; ch.height = 7; ch.width = 16
    ch.series[0].graphicalProperties.solidFill = C_ACCENT
    ws.add_chart(ch, f"B{note_r + 2}")
    ws.freeze_panes = "B5"
    return svc_rows

# ---------------------------------------------------------------- MONTHLY
MONTHLY_ROWS = {}


def build_monthly(wb):
    ws = wb["MONTHLY"]
    title_block(ws, "MONTHLY PERFORMANCE", "Twelve-month view of the reporting year (set in SETTINGS). Money figures are management "
                "estimates from your logs — not accounting or tax statements.", 14, "money")
    ws.column_dimensions["A"].width = 38
    for i in range(2, 15):
        ws.column_dimensions[L(i)].width = 11
    J, LE, RV, EX, MI, QC, IS, RW, RF, CU = (LOGS[k] for k in ["JOBS", "LEADS", "REVENUE", "EXPENSES", "MILEAGE",
                                                               "QUALITY CONTROL", "ISSUES", "REVIEWS", "REFERRALS", "CUSTOMERS"])
    kv_cell(ws, "A4", '="Year: "&ReportYear', bold=True, color="FFFFFF", fillc=C_PRIMARY)
    for m in range(12):
        kv_cell(ws, f"{L(2+m)}4", f"=DATE(ReportYear,{m+1},1)", 'mmm', bold=True, color="FFFFFF", fillc=C_PRIMARY,
                align=Alignment(horizontal="center"))
    kv_cell(ws, "N4", "Year Total", bold=True, color="FFFFFF", fillc=C_PRIMARY, align=Alignment(horizontal="center"))
    M = lambda col: f"{col}$4"
    comp = f'{J.rng("status")},"Completed"'
    rows = [
        ("sec", "SALES"),
        ("leads", "New leads", lambda c: f'=COUNTIFS({LE.rng("month")},{M(c)})', FMT_INT, "sum"),
        ("quotes", "Quotes given (leads with a quote value)", lambda c: f'=COUNTIFS({LE.rng("month")},{M(c)},{LE.rng("quote")},">0")', FMT_INT, "sum"),
        ("booked", "Leads booked", lambda c: f'=COUNTIFS({LE.rng("month")},{M(c)},{LE.rng("outcome")},"Booked")', FMT_INT, "sum"),
        ("conv", "Lead-to-booked rate", lambda c: f'=IF({c}{{leads}}=0,"",{c}{{booked}}/{c}{{leads}})', FMT_PCT, "ratio:booked/leads"),
        ("avgq", "Average quote", lambda c: f'=IFERROR(AVERAGEIFS({LE.rng("quote")},{LE.rng("month")},{M(c)},{LE.rng("quote")},">0"),"")', FMT_MONEY0, "avgq"),
        ("sec", "OPERATIONS"),
        ("jobs", "Jobs completed", lambda c: f'=COUNTIFS({J.rng("month")},{M(c)},{comp})', FMT_INT, "sum"),
        ("hrs", "Labor hours", lambda c: f'=SUMIFS({J.rng("hrs")},{J.rng("month")},{M(c)},{comp})', FMT_NUM1, "sum"),
        ("newc", "New customers (first completed job)", lambda c: f'=COUNTIFS({CU.rng("first")},">="&{M(c)},{CU.rng("first")},"<="&EOMONTH({M(c)},0))', FMT_INT, "sum"),
        ("qc", "Average QC score", lambda c: f'=IFERROR(AVERAGEIFS({QC.rng("score")},{QC.rng("month")},{M(c)}),"")', FMT_PCT, "avgqc"),
        ("issues", "Issues reported", lambda c: f'=COUNTIFS({IS.rng("month")},{M(c)})', FMT_INT, "sum"),
        ("miles", "Miles logged", lambda c: f'=SUMIFS({MI.rng("used")},{MI.rng("month")},{M(c)})', FMT_INT, "sum"),
        ("sec", "MONEY  (estimates)"),
        ("rev", "Revenue (completed jobs, before tax)", lambda c: f'=SUMIFS({J.rng("price")},{J.rng("month")},{M(c)},{comp})', FMT_MONEY0, "sum"),
        ("labor", "Labor cost", lambda c: f'=SUMIFS({J.rng("labor")},{J.rng("month")},{M(c)},{comp})', FMT_MONEY0, "sum"),
        ("sup", "Supplies (per-job estimate)", lambda c: f'=SUMIFS({J.rng("sup")},{J.rng("month")},{M(c)},{comp})', FMT_MONEY0, "sum"),
        ("veh", "Vehicle cost (miles × your rate)", lambda c: f'=SUMIFS({J.rng("vehicle")},{J.rng("month")},{M(c)},{comp})', FMT_MONEY0, "sum"),
        ("fees", "Payment fees", lambda c: f'=SUMIFS({J.rng("fees")},{J.rng("month")},{M(c)},{comp})', FMT_MONEY0, "sum"),
        ("gp", "Job gross profit", lambda c: f'={c}{{rev}}-{c}{{labor}}-{c}{{sup}}-{c}{{veh}}-{c}{{fees}}', FMT_MONEY0, "sum"),
        ("gm", "Gross margin", lambda c: f'=IF({c}{{rev}}=0,"",{c}{{gp}}/{c}{{rev}})', FMT_PCT, "ratio:gp/rev"),
        ("oh", "Overhead expenses (categories marked Overhead)", lambda c: f'=SUMIFS({EX.rng("amount")},{EX.rng("month")},{M(c)},{EX.rng("oh")},"Yes")', FMT_MONEY0, "sum"),
        ("op", "ESTIMATED OPERATING PROFIT", lambda c: f'={c}{{gp}}-{c}{{oh}}', FMT_MONEY0, "sum"),
        ("opm", "Operating margin", lambda c: f'=IF({c}{{rev}}=0,"",{c}{{op}}/{c}{{rev}})', FMT_PCT, "ratio:op/rev"),
        ("avgjob", "Average job value", lambda c: f'=IF({c}{{jobs}}=0,"",{c}{{rev}}/{c}{{jobs}})', FMT_MONEY0, "ratio:rev/jobs"),
        ("revhr", "Revenue per labor hour", lambda c: f'=IF(N({c}{{hrs}})=0,"",{c}{{rev}}/{c}{{hrs}})', FMT_MONEY, "ratio:rev/hrs"),
        ("coll", "Payments collected (incl. tax)", lambda c: f'=SUMIFS({RV.rng("amount")},{RV.rng("month")},{M(c)})', FMT_MONEY0, "sum"),
        ("spend", "All spending logged in EXPENSES", lambda c: f'=SUMIFS({EX.rng("amount")},{EX.rng("month")},{M(c)})', FMT_MONEY0, "sum"),
        ("sec", "GROWTH"),
        ("revs", "Reviews received", lambda c: f'=COUNTIFS({RW.rng("recmonth")},{M(c)},{RW.rng("status")},"Received")', FMT_INT, "sum"),
        ("rating", "Average rating", lambda c: f'=IFERROR(AVERAGEIFS({RW.rng("rating")},{RW.rng("recmonth")},{M(c)},{RW.rng("status")},"Received"),"")', '0.0', "avgrating"),
        ("refs", "Referrals received", lambda c: f'=COUNTIFS({RF.rng("month")},{M(c)})', FMT_INT, "sum"),
        ("refb", "Referrals booked", lambda c: f'=COUNTIFS({RF.rng("month")},{M(c)},{RF.rng("status")},"Booked")', FMT_INT, "sum"),
    ]
    r = 5
    for item in rows:
        if item[0] == "sec":
            section(ws, f"A{r}", item[1], 14)
            r += 1
            continue
        MONTHLY_ROWS[item[0]] = r
        r += 1
    for item in rows:
        if item[0] == "sec":
            continue
        key, label, fn, fmt, total = item
        r = MONTHLY_ROWS[key]
        big = label.isupper()
        kv_cell(ws, f"A{r}", label, bold=big)
        for m in range(12):
            c = L(2 + m)
            f = fn(c).replace("{{", "{").replace("}}", "}").format(**MONTHLY_ROWS)
            kv_cell(ws, f"{c}{r}", f, fmt, bold=big, fillc=C_GREENBG if big else None)
        if total == "sum":
            tf = f"=SUM(B{r}:M{r})"
        elif total.startswith("ratio:"):
            a, b = total[6:].split("/")
            tf = f'=IF(N(N{MONTHLY_ROWS[b]})=0,"",N{MONTHLY_ROWS[a]}/N{MONTHLY_ROWS[b]})'
        elif total == "avgq":
            tf = f'=IFERROR(AVERAGEIFS({LE.rng("quote")},{LE.rng("month")},">="&B$4,{LE.rng("month")},"<="&M$4,{LE.rng("quote")},">0"),"")'
        elif total == "avgqc":
            tf = f'=IFERROR(AVERAGEIFS({QC.rng("score")},{QC.rng("month")},">="&B$4,{QC.rng("month")},"<="&M$4),"")'
        elif total == "avgrating":
            tf = f'=IFERROR(AVERAGEIFS({RW.rng("rating")},{RW.rng("recmonth")},">="&B$4,{RW.rng("recmonth")},"<="&M$4,{RW.rng("status")},"Received"),"")'
        kv_cell(ws, f"N{r}", tf, fmt, bold=True, fillc=C_LIGHT)
    cf_formula(ws, f"B{MONTHLY_ROWS['op']}:N{MONTHLY_ROWS['op']}", f"B{MONTHLY_ROWS['op']}<0", C_RED, C_REDBG)
    nr = r + 1
    ws[f"A{nr}"] = ("How to read this: Job gross profit uses your per-job cost estimates. Overhead = EXPENSES in categories marked "
                    "'Counted as Overhead? = Yes' in SETTINGS (staff pay, job supplies, fuel and merchant fees are excluded there because "
                    "they are already costed per job). Talk to a qualified accountant for tax or financial statements.")
    ws[f"A{nr}"].font = font(8, italic=True, color=C_MUTED)
    ws.merge_cells(f"A{nr}:N{nr+1}")
    ws[f"A{nr}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[nr].height = 30
    # chart
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.style = 10
    ch.title = "Revenue vs job gross profit vs operating profit"
    for key, color in (("rev", C_PRIMARY), ("gp", C_ACCENT), ("op", "E08E2B")):
        rr = MONTHLY_ROWS[key]
        ref = Reference(ws, min_col=1, max_col=13, min_row=rr, max_row=rr)
        ch.add_data(ref, from_rows=True, titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
    ch.set_categories(Reference(ws, min_col=2, max_col=13, min_row=4, max_row=4))
    ch.height = 8; ch.width = 26
    ch.y_axis.numFmt = '$#,##0'
    ch.y_axis.majorGridlines = None
    ws.add_chart(ch, f"A{nr + 3}")
    ws.freeze_panes = "B5"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- DASHBOARD
def build_dashboard(wb):
    ws = wb["DASHBOARD"]
    ws.sheet_properties.tabColor = TAB["start"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for i in range(2, 14):
        ws.column_dimensions[L(i)].width = 13.5
    ws.column_dimensions["N"].width = 2
    ws.merge_cells("B1:M1"); ws.merge_cells("B2:M2")
    for col in range(1, 15):
        ws.cell(1, col).fill = fill(C_PRIMARY)
    ws["B1"] = '="EXECUTIVE DASHBOARD  ·  "&UPPER(BizName)'
    ws["B1"].font = font(18, True, "FFFFFF"); ws["B1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 36
    ws["B2"] = '="Reporting month: "&TEXT(ReportMonth,"mmmm yyyy")&"   ·   As of "&TEXT(AsOf,"mmm d, yyyy")&"   ·   Change dates in SETTINGS"'
    ws["B2"].font = font(10, False, C_MUTED, italic=True)
    link(ws["B3"], "START HERE", "← Start Here")
    link(ws["D3"], "SETTINGS", "Settings →")
    J, LE, RV, EX, QC, IS, RW, RF, CU, RC, SU, FU = (LOGS[k] for k in ["JOBS", "LEADS", "REVENUE", "EXPENSES", "QUALITY CONTROL",
                                                                      "ISSUES", "REVIEWS", "REFERRALS", "CUSTOMERS", "RECURRING",
                                                                      "SUPPLIES", "FOLLOW-UPS"])
    MR = MONTHLY_ROWS
    mcol = 'INDEX(MONTHLY!$B${r}:$M${r},MONTH(ReportMonth))'
    prev = 'IF(MONTH(ReportMonth)=1,"",INDEX(MONTHLY!$B${r}:$M${r},MONTH(ReportMonth)-1))'

    def mval(key):
        return mcol.format(r=MR[key])

    def delta(key):
        p = prev.format(r=MR[key])
        return (f'=IFERROR(IF(OR({p}="",N({p})=0),"No prior-month comparison",TEXT({mval(key)}/{p}-1,"+0%;-0%;0%")&" vs last month"),"No prior-month comparison")')

    comp = f'{J.rng("status")},"Completed"'
    followups_due = (f'COUNTIF({LE.rng("fu")},"Overdue")+COUNTIF({LE.rng("fu")},"Due today")'
                     f'+COUNTIFS({FU.rng("status")},"Open",{FU.rng("due")},"<="&AsOf)')
    tiles = [
        ("MONEY  —  REPORTING MONTH", [
            ("Revenue", f"={mval('rev')}", FMT_MONEY0, delta("rev")),
            ("Job Gross Profit", f"={mval('gp')}", FMT_MONEY0, delta("gp")),
            ("Gross Margin", f"={mval('gm')}", FMT_PCT, '="Target: "&TEXT(TargetMargin,"0%")'),
            ("Est. Operating Profit", f"={mval('op')}", FMT_MONEY0, '="After overhead expenses"'),
            ("Payments Collected", f"={mval('coll')}", FMT_MONEY0, '="Incl. tax, by payment date"'),
            ("Unpaid Balances (all)", f'=SUM({J.rng("balance")})', FMT_MONEY0, f'=COUNTIF({J.rng("paystat")},"Unpaid")+COUNTIF({J.rng("paystat")},"Partial")&" jobs with a balance"'),
        ]),
        ("SALES PIPELINE", [
            ("New Leads", f"={mval('leads')}", FMT_INT, delta("leads")),
            ("Quotes Given", f"={mval('quotes')}", FMT_INT, '="Leads added this month with a quote"'),
            ("Leads Booked", f"={mval('booked')}", FMT_INT, '="From this month\'s leads"'),
            ("Conversion Rate", f"={mval('conv')}", FMT_PCT, '="Booked ÷ new leads"'),
            ("Average Quote", f"={mval('avgq')}", FMT_MONEY0, '="This month\'s quotes"'),
            ("Follow-Ups Due Now", "=" + followups_due, FMT_INT, '="Overdue + due today"'),
        ]),
        ("OPERATIONS", [
            ("Jobs Completed", f"={mval('jobs')}", FMT_INT, delta("jobs")),
            ("Average Job Value", f"={mval('avgjob')}", FMT_MONEY0, '="Revenue ÷ jobs"'),
            ("Revenue / Labor Hour", f"={mval('revhr')}", FMT_MONEY, '="All workers combined"'),
            ("Jobs Next 7 Days", f'=COUNTIFS({J.rng("date")},">="&AsOf,{J.rng("date")},"<="&AsOf+7,{J.rng("status")},"Scheduled")', FMT_INT, '="Scheduled in JOBS"'),
            ("Active Recurring Plans", f'=COUNTIF({RC.rng("status")},"Active")', FMT_INT,
             f'=SUMPRODUCT(({CU.rng("recur")}="Yes")*1)&" recurring clients"'),
            ("Recurring Monthly Value", f'=SUM({RC.rng("mval")})', FMT_MONEY0, '="Active plans, per month"'),
        ]),
        ("REPUTATION & QUALITY", [
            ("Reviews Received", f"={mval('revs')}", FMT_INT, '="This month"'),
            ("Average Rating", f'=IFERROR(AVERAGEIFS({RW.rng("rating")},{RW.rng("status")},"Received"),"")', '0.0', '="All-time, received reviews"'),
            ("Referrals", f"={mval('refs')}", FMT_INT, f'="{"{"}"&COUNTIF({RF.rng("status")},"Booked")&" booked all-time"'.replace('"{"&', '"')),
            ("QC Pass Rate", f'=IFERROR(COUNTIFS({QC.rng("month")},ReportMonth,{QC.rng("result")},"Pass")/COUNTIFS({QC.rng("month")},ReportMonth,{QC.rng("result")},"?*"),"")', FMT_PCT, '="Inspections this month"'),
            ("Open Issues", f'=COUNTIF({IS.rng("status")},"Open")+COUNTIF({IS.rng("status")},"In progress")', FMT_INT, '="From the issue log"'),
            ("Customers to Reactivate", f'=COUNTIF({CU.rng("action")},"Reactivate")', FMT_INT, '="No job in "&ReactDays&"+ days"'),
        ]),
    ]
    r = 4
    for sec_title, items in tiles:
        section(ws, f"B{r}", sec_title, 12)
        r += 1
        for i, (label, f, fmt, sub) in enumerate(items):
            c1, c2 = L(2 + i * 2), L(3 + i * 2)
            for rr in (r, r + 1, r + 2):
                ws.merge_cells(f"{c1}{rr}:{c2}{rr}")
                for cc in (c1, c2):
                    ws[f"{cc}{rr}"].fill = fill(C_LIGHT)
            ws[f"{c1}{r}"] = label.upper(); ws[f"{c1}{r}"].font = font(8, True, C_MUTED)
            ws[f"{c1}{r}"].alignment = Alignment(indent=1, vertical="bottom")
            ws[f"{c1}{r+1}"] = f; ws[f"{c1}{r+1}"].font = font(20, True, C_PRIMARY)
            ws[f"{c1}{r+1}"].number_format = fmt; ws[f"{c1}{r+1}"].alignment = Alignment(indent=1, horizontal="left")
            ws[f"{c1}{r+2}"] = sub; ws[f"{c1}{r+2}"].font = font(8, False, C_MUTED, italic=True)
            ws[f"{c1}{r+2}"].alignment = Alignment(indent=1, vertical="top")
            for rr in (r, r + 1, r + 2):
                ws[f"{c1}{rr}"].border = Border(left=Side(style="thick", color=C_ACCENT))
        ws.row_dimensions[r].height = 18
        ws.row_dimensions[r + 1].height = 30
        ws.row_dimensions[r + 2].height = 16
        r += 4
    # fix referrals subtitle formula (simple)
    # Action center
    ac = r
    section(ws, f"B{ac}", "ACTION CENTER  —  what needs attention today", 12)
    actions = [
        ("Leads overdue or due for follow-up", f'=COUNTIF({LE.rng("fu")},"Overdue")+COUNTIF({LE.rng("fu")},"Due today")', "FOLLOW-UPS"),
        ("Other follow-ups overdue / due (your log)", f'=COUNTIFS({FU.rng("status")},"Open",{FU.rng("due")},"<="&AsOf)', "FOLLOW-UPS"),
        ("Completed jobs not fully paid", f'=COUNTIF({J.rng("paystat")},"Unpaid")+COUNTIF({J.rng("paystat")},"Partial")', "FOLLOW-UPS"),
        ("Recurring visits that need booking", f'=SUM({RC.rng("needs")})', "SCHEDULE"),
        ("Booked leads missing from CUSTOMERS", f'=COUNTIF({LE.rng("check")},"Add to CUSTOMERS")', "LEADS"),
        ("Review requests needing a reminder or reply", f'=COUNTIF({RW.rng("chase")},"Send reminder")+COUNTIF({RW.rng("chase")},"Reply to review")', "REVIEWS"),
        ("Referral rewards owed", f'=COUNTIF({RF.rng("rstatus")},"Owed")', "REFERRALS"),
        ("Rework jobs still open", f'=COUNTIF({QC.rng("open")},"Rework open")', "QUALITY CONTROL"),
        ("Supplies at or below reorder level", f'=COUNTIF({SU.rng("flag")},"Reorder")', "SUPPLIES"),
        ("Overdue staff tasks", f'=COUNTIF({LOGS["TASKS"].rng("flag")},"Overdue")', "STAFF & TASKS"),
    ]
    for i, (label, f, target) in enumerate(actions):
        rr = ac + 1 + i
        ws.merge_cells(f"B{rr}:E{rr}")
        kv_cell(ws, f"B{rr}", label, border=False)
        kv_cell(ws, f"F{rr}", f, FMT_INT, bold=True, color=C_PRIMARY, border=False, align=Alignment(horizontal="center"))
        link(ws[f"G{rr}"], target, f"Open {target.title()} →")
        cf_formula(ws, f"F{rr}", f"F{rr}>0", C_RED, C_REDBG)
        for cc in "BCDEFG":
            ws[f"{cc}{rr}"].border = Border(bottom=Side(style="hair", color="C9D3D7"))
    # Pipeline table (right side of action center)
    pc = ac
    section(ws, f"I{pc}", "PIPELINE BY STAGE  (all leads, current stage)", 5)
    for i, h in enumerate(["Stage", "Leads", "Quote Value"]):
        kv_cell(ws, f"{L(9 + i*2)}{pc+1}", h, bold=True, size=9, color="FFFFFF", fillc=C_PRIMARY)
        ws.merge_cells(f"{L(9 + i*2)}{pc+1}:{L(10 + i*2)}{pc+1}") if i < 2 else None
    stages = [x for x in LISTS if x[0] == "stages"][0][2]
    for i, s in enumerate(stages):
        rr = pc + 2 + i
        ws.merge_cells(f"I{rr}:J{rr}"); ws.merge_cells(f"K{rr}:L{rr}")
        kv_cell(ws, f"I{rr}", s, size=9)
        kv_cell(ws, f"K{rr}", f'=COUNTIF({LE.rng("stage")},I{rr})', FMT_INT, size=9, align=Alignment(horizontal="center"))
        kv_cell(ws, f"M{rr}", f'=SUMIF({LE.rng("stage")},I{rr},{LE.rng("quote")})', FMT_MONEY0, size=9)
    # charts: pipeline + monthly
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Leads by stage"
    ch.add_data(Reference(ws, min_col=11, min_row=pc + 2, max_row=pc + 1 + len(stages)), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=9, min_row=pc + 2, max_row=pc + 1 + len(stages)))
    ch.legend = None; ch.height = 7.5; ch.width = 13
    ch.series[0].graphicalProperties.solidFill = C_BLUE
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.majorGridlines = None
    cr = ac + len(actions) + 3
    section(ws, f"B{cr}", "TRENDS", 12)
    ws.add_chart(ch, f"I{cr+1}")
    ms = wb["MONTHLY"]
    ch2 = BarChart(); ch2.type = "col"; ch2.style = 10; ch2.title = "Revenue & job gross profit by month"
    for key, color in (("rev", C_PRIMARY), ("gp", C_ACCENT)):
        rr = MR[key]
        ch2.add_data(Reference(ms, min_col=1, max_col=13, min_row=rr, max_row=rr), from_rows=True, titles_from_data=True)
        ch2.series[-1].graphicalProperties.solidFill = color
    ch2.set_categories(Reference(ms, min_col=2, max_col=13, min_row=4, max_row=4))
    ch2.height = 7.5; ch2.width = 17; ch2.y_axis.numFmt = '$#,##0'; ch2.y_axis.majorGridlines = None
    ch2.legend.position = "b"
    ws.add_chart(ch2, f"B{cr+1}")
    ws[f"B{cr+17}"] = ("All figures come from your logs and your assumptions. They are management estimates, not accounting, "
                       "tax or legal advice.")
    ws[f"B{cr+17}"].font = font(8, italic=True, color=C_MUTED)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 1
    # replace the tricky referrals subtitle with a clean formula
    for row in ws.iter_rows(min_row=4, max_row=ac):
        for cell in row:
            if isinstance(cell.value, str) and "booked all-time" in cell.value:
                cell.value = f'=COUNTIF({RF.rng("status")},"Booked")&" booked all-time"'

# ---------------------------------------------------------------- START HERE
MODULES = [
    ("GET STARTED", [("START HERE", "This page: setup steps, map and color legend."),
                     ("DASHBOARD", "Executive dashboard: money, pipeline, operations, reputation and today's action list."),
                     ("SETTINGS", "Business details, cost assumptions, targets, dates and all dropdown lists.")]),
    ("WIN THE CUSTOMER", [("LEADS", "Lead CRM and sales pipeline: 8 stages, automatic follow-up status and outcome."),
                          ("QUOTE BUILDER", "Price one job: price, labor cost, direct costs, gross profit, margin and break-even."),
                          ("PRICING", "Your pricing assumptions by cleaning type, condition, frequency and add-ons + rate card."),
                          ("CUSTOMERS", "Customer list with lifetime revenue, balance due and a suggested next action."),
                          ("PROPERTIES", "Property details (size, rooms, access, pets) that feed automatic time estimates.")]),
    ("DO THE WORK", [("JOBS", "Every job: schedule, worker, hours, price, payment status and estimated profit."),
                     ("RECURRING", "Recurring plans: next visit due, booking status and recurring revenue."),
                     ("SCHEDULE", "Weekly calendar built from JOBS + recurring visits that need booking."),
                     ("STAFF & TASKS", "Team roster with pay rates and monthly output, plus a task assignment list."),
                     ("SUPPLIES", "Supply inventory with reorder alerts."),
                     ("QUALITY CONTROL", "Inspection log with pass/fail scores, rework tracking and an issue log."),
                     ("QC CHECKLISTS", "Six printable checklists: standard, deep, move, Airbnb, commercial, final inspection.")]),
    ("GET PAID & KNOW YOUR NUMBERS", [("REVENUE", "Payments received, linked to jobs, with processing-fee estimates."),
                                      ("EXPENSES", "Business expenses with overhead flag."),
                                      ("MILEAGE", "Trip log that feeds vehicle cost into each job."),
                                      ("PROFITABILITY", "Profit by service, worker and customer type, for any period."),
                                      ("MONTHLY", "Twelve-month performance: sales, operations, money and growth.")]),
    ("GROW", [("FOLLOW-UPS", "Automatic list of leads to chase and unpaid jobs, plus your follow-up log."),
              ("REVIEWS", "Review requests, ratings and reply tracking."),
              ("REFERRALS", "Referrals, rewards owed and revenue from referrals.")]),
]


def build_start(wb, demo):
    ws = wb["START HERE"]
    ws.sheet_properties.tabColor = TAB["start"]
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", [3, 26, 72, 3, 30, 30]):
        ws.column_dimensions[col].width = w
    for col in range(1, 7):
        ws.cell(1, col).fill = fill(C_PRIMARY); ws.cell(2, col).fill = fill(C_PRIMARY)
    ws["B1"] = "CLEANING BUSINESS AI GROWTH OS"
    ws["B1"].font = font(22, True, "FFFFFF"); ws.row_dimensions[1].height = 42
    ws["B2"] = "Run your cleaning business from lead to payment in one connected system.   ·   by OperatorGrid"
    ws["B2"].font = font(11, False, "D8F0EC", italic=True); ws.row_dimensions[2].height = 22
    ws["B3"] = ("DEMO FILE — every name, address and number is fictional sample data. Explore it, then use the CLEAN file for your business."
                if demo else "CLEAN FILE — ready for your business. Follow the 5 setup steps below.")
    ws["B3"].font = font(10, True, C_AMBER if demo else C_GREEN)
    r = 5
    section(ws, f"B{r}", "SET UP IN 5 STEPS  (about 30 minutes)", 2); r += 1
    steps = [
        ("1. SETTINGS", "Enter your business details. Replace the example labor, supply, vehicle and fee assumptions with your own. Leave the as-of date BLANK."),
        ("2. STAFF & TASKS", "Add each worker's name and pay rate (the owner too, if you clean)."),
        ("3. PRICING", "Adjust billing rates, time assumptions and add-ons until the rate card matches how you want to price. Test a few quotes."),
        ("4. CUSTOMERS & PROPERTIES", "Add current customers and their properties (square feet, bedrooms and bathrooms power automatic time estimates)."),
        ("5. RECURRING & JOBS", "Add recurring plans, then your upcoming jobs. From then on: log payments in REVENUE and check the DASHBOARD daily."),
    ]
    for a, b in steps:
        kv_cell(ws, f"B{r}", a, bold=True, color=C_PRIMARY, border=False)
        kv_cell(ws, f"C{r}", b, border=False, align=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    section(ws, f"B{r}", "DAILY FLOW:  Lead → Quote → Book → Schedule → Clean → Inspect → Get paid → Review → Referral → Repeat", 2)
    r += 2
    for grp, mods in MODULES:
        kv_cell(ws, f"B{r}", grp, bold=True, size=10, color=C_ACCENT, border=False)
        r += 1
        for name, desc in mods:
            link(ws[f"B{r}"], name, name + "  →")
            ws[f"B{r}"].font = font(10, True, C_BLUE, underline="single")
            kv_cell(ws, f"C{r}", desc, border=False, color=C_TEXT)
            for cc in "BC":
                ws[f"{cc}{r}"].border = Border(bottom=Side(style="hair", color="C9D3D7"))
            r += 1
        r += 1
    # legend on right
    lr = 5
    section(ws, f"E{lr}", "COLOR LEGEND", 2)
    legend = [("Type here", "FFFFFF", C_TEXT), ("Key setting / input", C_INPUT, C_TEXT), ("Automatic — don't type", C_AUTOFILL, "33434A"),
              ("Column header you fill", C_PRIMARY, "FFFFFF"), ("Automatic column header ⚙", C_AUTOHDR, "FFFFFF"),
              ("Needs attention", C_REDBG, C_RED), ("Below target", C_AMBERBG, C_AMBER), ("Good / done", C_GREENBG, C_GREEN)]
    for i, (t, bg, fg) in enumerate(legend):
        kv_cell(ws, f"E{lr+1+i}", t, fillc=bg, color=fg, bold=True)
    gr = lr + len(legend) + 3
    section(ws, f"E{gr}", "GOLDEN RULES", 2)
    rules = ["Never type in gray cells or ⚙ columns.", "Keep customer names, property labels and staff names unique.",
             "Use dropdowns — they keep the links working.", "Filter freely; sort a sheet only by selecting the whole table.",
             "Add rows by typing in the next empty row (1,500 job rows included).", "Back up weekly (File → Save a copy).",
             "Numbers are estimates from your assumptions — not tax, legal or accounting advice."]
    for i, t in enumerate(rules):
        c = kv_cell(ws, f"E{gr+1+i}", "•  " + t, border=False, size=9, align=Alignment(wrap_text=True, vertical="top"))
        ws.merge_cells(f"E{gr+1+i}:F{gr+1+i}")
        ws.row_dimensions[gr + 1 + i].height = max(ws.row_dimensions[gr + 1 + i].height or 15, 26)
    hr_ = gr + len(rules) + 2
    section(ws, f"E{hr_}", "INCLUDED WITH THIS SYSTEM", 2)
    inc = ["Quick Start Guide (PDF)", "Full User Guide (PDF)", "AI Workflow Library — 18 workflows (PDF)",
           "Client Forms — 14 editable templates", "Marketing Kit — print + social templates"]
    for i, t in enumerate(inc):
        kv_cell(ws, f"E{hr_+1+i}", "✓  " + t, border=False, size=9)
    ws.freeze_panes = "A4"

# ---------------------------------------------------------------- simple log pages with extras
def add_log_cf():
    pass


def build(out_path, demo):
    wb = Workbook()
    wb.remove(wb.active)
    for name, tab in SHEETS:
        wb.create_sheet(name)
    data = demo_data.generate() if demo else {}
    g = lambda k: data.get(k, [])

    build_settings(wb, demo)
    build_pricing(wb)

    # Log sheets
    ws = wb["LEADS"]; sp = LOGS["LEADS"]
    setup_log_sheet(ws, sp, "LEAD CRM & SALES PIPELINE", "One row per lead. Move the Stage as the lead progresses. "
                    "Follow-up status, outcome and 'add to customers' reminders are automatic.", "sales", g("leads"))
    for s, fg, bg in [("New Lead", C_BLUE, C_BLUEBG), ("Quote Sent", C_AMBER, C_AMBERBG), ("Follow-Up", C_AMBER, C_AMBERBG),
                      ("Booked", C_GREEN, C_GREENBG), ("Recurring Client", C_GREEN, C_GREENBG), ("Lost/Declined", C_MUTED, "ECEFF1")]:
        cf_cell(ws, sp, "stage", s, fg, bg)
    cf_cell(ws, sp, "fu", "Overdue", C_RED, C_REDBG); cf_cell(ws, sp, "fu", "Due today", C_AMBER, C_AMBERBG)
    cf_cell(ws, sp, "check", "Add to CUSTOMERS", C_RED, C_REDBG)

    ws = wb["CUSTOMERS"]; sp = LOGS["CUSTOMERS"]
    setup_log_sheet(ws, sp, "CUSTOMERS", "Your customer list. History, lifetime revenue, balance due and the suggested next action "
                    "fill in automatically from JOBS, RECURRING, REVIEWS and REFERRALS.", "sales", g("customers"))
    cf_cell(ws, sp, "action", "Collect balance", C_RED, C_REDBG); cf_cell(ws, sp, "action", "Reactivate", C_AMBER, C_AMBERBG)
    cf_cell(ws, sp, "action", "Offer recurring plan", C_BLUE, C_BLUEBG)

    ws = wb["PROPERTIES"]; sp = LOGS["PROPERTIES"]
    setup_log_sheet(ws, sp, "PROPERTIES", "One row per property. Square footage, bedrooms and bathrooms power the automatic "
                    "time estimate in JOBS. Keep access notes here, not in public places.", "sales", g("properties"))

    build_quote(wb, demo)

    ws = wb["JOBS"]; sp = LOGS["JOBS"]
    setup_log_sheet(ws, sp, "JOBS", "Every job, one row. Fill the white columns; payment status, costs and estimated profit are automatic. "
                    "Link recurring visits with a Plan ID. Log payments in REVENUE and trips in MILEAGE.", "ops", g("jobs"), freeze="E5")
    for s, fg, bg in [("Completed", C_GREEN, C_GREENBG), ("Scheduled", C_BLUE, C_BLUEBG), ("Cancelled", C_MUTED, "ECEFF1"),
                      ("No-Show", C_RED, C_REDBG), ("Rescheduled", C_AMBER, C_AMBERBG)]:
        cf_cell(ws, sp, "status", s, fg, bg)
    cf_cell(ws, sp, "paystat", "Unpaid", C_RED, C_REDBG); cf_cell(ws, sp, "paystat", "Partial", C_AMBER, C_AMBERBG)
    cf_cell(ws, sp, "paystat", "Paid", C_GREEN, C_GREENBG)
    c = sp.letters["margin"]
    cf_formula(ws, f"{c}{sp.first}:{c}{sp.last}", f'AND({c}{sp.first}<>"",{c}{sp.first}<0)', C_RED, C_REDBG)
    cf_formula(ws, f"{c}{sp.first}:{c}{sp.last}", f'AND({c}{sp.first}<>"",{c}{sp.first}<TargetMargin)', C_AMBER, C_AMBERBG)
    for hk in ("stime", "key", "urank"):
        ws.column_dimensions[sp.letters[hk]].hidden = False

    ws = wb["RECURRING"]; sp = LOGS["RECURRING"]
    setup_log_sheet(ws, sp, "RECURRING JOBS", "One row per recurring plan. When you book a visit in JOBS, choose this Plan ID. "
                    "Next due date, booking status and recurring revenue update automatically.", "ops", g("recurring"))
    cf_cell(ws, sp, "book", "Overdue — book now", C_RED, C_REDBG); cf_cell(ws, sp, "book", "Due this week — book", C_AMBER, C_AMBERBG)
    cf_cell(ws, sp, "book", "Booked", C_GREEN, C_GREENBG)

    build_schedule(wb)

    # STAFF & TASKS (two tables)
    ws = wb["STAFF & TASKS"]; sp = LOGS["STAFF & TASKS"]; tk = LOGS["TASKS"]
    title_block(ws, "STAFF & TASK ASSIGNMENTS", "Top: team roster (names feed every Worker dropdown; pay rates feed job labor cost). "
                "Bottom: task assignments. Hiring, pay and classification decisions are yours — get professional advice where needed.", 13, "ops")
    write_log(ws, sp, g("staff"))
    ws["A28"] = "TASK ASSIGNMENTS"; ws["A28"].font = font(11, True, C_PRIMARY)
    ws["A29"] = "Give each task an owner and due date. Use AI Workflow 16 to turn notes into clear task instructions."
    ws["A29"].font = font(8, italic=True, color=C_MUTED)
    write_log(ws, tk, g("tasks"))
    ws.column_dimensions["A"].width = 34
    cf_cell(ws, tk, "flag", "Overdue", C_RED, C_REDBG); cf_cell(ws, tk, "flag", "✓ Done", C_GREEN, C_GREENBG)
    ws.auto_filter.ref = f"A{tk.header_row}:{L(len(tk.cols))}{tk.last}"
    ws.freeze_panes = "B5"

    ws = wb["REVENUE"]; sp = LOGS["REVENUE"]
    setup_log_sheet(ws, sp, "REVENUE  —  PAYMENTS RECEIVED", "Log every payment against its Job ID. JOBS then shows Paid / Partial / Unpaid "
                    "automatically. Tips are tracked separately and are not counted as job revenue.", "money", g("payments"))
    cf_cell(ws, sp, "customer", "Job ID not found", C_RED, C_REDBG)

    ws = wb["EXPENSES"]; sp = LOGS["EXPENSES"]
    setup_log_sheet(ws, sp, "EXPENSES", "Log business spending. Categories marked 'Overhead' in SETTINGS reduce operating profit on MONTHLY. "
                    "Keep receipts — ask a tax professional what you can deduct.", "money", g("expenses"))

    build_profitability(wb)

    ws = wb["MILEAGE"]; sp = LOGS["MILEAGE"]
    setup_log_sheet(ws, sp, "MILEAGE LOG", "Log business trips. Use odometer readings or type miles directly. Trips linked to a Job ID "
                    "add vehicle cost to that job's profit. Ask a tax professional about mileage record requirements where you live.", "money", g("mileage"))

    ws = wb["SUPPLIES"]; sp = LOGS["SUPPLIES"]
    setup_log_sheet(ws, sp, "SUPPLIES & INVENTORY", "Track what you have and when to reorder. Always follow product labels and safety data sheets.",
                    "ops", g("supplies"), freeze="B5")
    cf_cell(ws, sp, "flag", "Reorder", C_RED, C_REDBG); cf_cell(ws, sp, "flag", "OK", C_GREEN, C_GREENBG)

    build_followups(wb, g("followups"))

    ws = wb["REVIEWS"]; sp = LOGS["REVIEWS"]
    setup_log_sheet(ws, sp, "REVIEWS", "Track every review request. Ask within a day of a great clean (AI Workflow 07), remind once, "
                    "and reply to every review (AI Workflow 08).", "growth", g("reviews"))
    cf_cell(ws, sp, "chase", "Send reminder", C_AMBER, C_AMBERBG); cf_cell(ws, sp, "chase", "Reply to review", C_RED, C_REDBG)

    ws = wb["REFERRALS"]; sp = LOGS["REFERRALS"]
    setup_log_sheet(ws, sp, "REFERRALS", "Track who sends you business and make sure every thank-you is delivered.", "growth", g("referrals"))
    cf_cell(ws, sp, "rstatus", "Owed", C_RED, C_REDBG); cf_cell(ws, sp, "status", "Booked", C_GREEN, C_GREENBG)

    ws = wb["QUALITY CONTROL"]; sp = LOGS["QUALITY CONTROL"]; isp = LOGS["ISSUES"]
    setup_log_sheet(ws, sp, "QUALITY CONTROL", "Top: inspection log (score = items passed ÷ items checked). "
                    "Bottom (row 311+): issue & rework log for complaints and problems.", "ops", g("qc"))
    ws[f"A{isp.header_row - 3}"] = "ISSUE & REWORK LOG"; ws[f"A{isp.header_row - 3}"].font = font(11, True, C_PRIMARY)
    ws[f"A{isp.header_row - 2}"] = "Log every complaint or problem, what you did about it and when it was resolved. Use AI Workflow 17 to draft the reply."
    ws[f"A{isp.header_row - 2}"].font = font(8, italic=True, color=C_MUTED)
    write_log(ws, isp, g("issues"))
    link(ws["D3"], "QUALITY CONTROL", "Jump to Issue Log ↓", f"A{isp.header_row}")
    cf_cell(ws, sp, "result", "Fail", C_RED, C_REDBG); cf_cell(ws, sp, "result", "Pass", C_GREEN, C_GREENBG)
    cf_cell(ws, sp, "open", "Rework open", C_RED, C_REDBG)
    cf_cell(ws, isp, "status", "Open", C_RED, C_REDBG); cf_cell(ws, isp, "status", "Resolved", C_GREEN, C_GREENBG)

    build_checklists(wb)
    build_monthly(wb)
    build_dashboard(wb)
    build_start(wb, demo)

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        if ws.title not in ("START HERE", "DASHBOARD"):
            ws.page_setup.fitToHeight = 0
        else:
            ws.page_setup.fitToHeight = 1
        ws.page_setup.orientation = "landscape" if ws.title != "START HERE" else "portrait"
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
    if os.environ.get("SCREENSHOT"):
        areas = {"LEADS": "A1:P80", "JOBS": "A1:P24", "PROFITABILITY": "A1:I13", "RECURRING": "A1:R17",
                 "QUOTE BUILDER": "A1:H41", "SCHEDULE": "A7:F21", "CUSTOMERS": "A1:S24", "MONTHLY": "A1:N40"}
        for c in "DEHMNO":
            wb["LEADS"].column_dimensions[c].hidden = True
        for r in range(14, 19):
            wb["SCHEDULE"].row_dimensions[r].hidden = True
        for c in "GH":
            wb["SCHEDULE"].column_dimensions[c].hidden = True
        for r in range(5, 63):
            wb["LEADS"].row_dimensions[r].hidden = True
        wb["SCHEDULE"]["B4"] = dt.date(int(os.environ["SHOTWEEK"][:4]), int(os.environ["SHOTWEEK"][5:7]), int(os.environ["SHOTWEEK"][8:])) if os.environ.get("SHOTWEEK") else None
        for ws in wb.worksheets:
            ws.print_area = areas.get(ws.title, ws.print_area or "A1:A1") if ws.title in areas or ws.title in ("DASHBOARD", "START HERE") else "A1:A1"
            if ws.title in ("DASHBOARD", "START HERE"):
                ws.print_area = None
    wb.properties.creator = "OperatorGrid"
    wb.properties.lastModifiedBy = "OperatorGrid"
    wb.properties.title = "Cleaning Business AI Growth OS"
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1]
    build(os.path.join(out, "CLEANING-BUSINESS-AI-GROWTH-OS-DEMO.xlsx"), True)
    build(os.path.join(out, "CLEANING-BUSINESS-AI-GROWTH-OS-CLEAN.xlsx"), False)
    print("built")
